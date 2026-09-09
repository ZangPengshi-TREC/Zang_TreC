#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""贈答品振データ → TRIAL振替伝票  マッピング表 / GAP分析書 / 変換ルール定義書"""

from openpyxl import Workbook
from openpyxl.styles import Font, Fill, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

OUT_DIR = "/Users/treqd/Desktop/Cursor/西友PMI/贈答品振"

NAVY = "0D2B4E"
BLUE = "1A5C9E"
TEAL = "0F766E"
ORANGE = "C2410C"
RED = "B91C1C"
GREEN = "166534"
GRAY = "4B5563"
WHITE = "FFFFFF"
BG_HEAD = "1A5C9E"
BG_ALT = "F1F5F9"
BG_WARN = "FEF3C7"
BG_GAP = "FEE2E2"
BG_OK = "DCFCE7"
BG_AUTO = "E0E7FF"
BG_NOTE = "EEF2FF"

thin = Border(
    left=Side(style="thin", color="CBD5E1"),
    right=Side(style="thin", color="CBD5E1"),
    top=Side(style="thin", color="CBD5E1"),
    bottom=Side(style="thin", color="CBD5E1"),
)
wrap = Alignment(wrap_text=True, vertical="center")
center = Alignment(wrap_text=True, vertical="center", horizontal="center")


def font(size=10, bold=False, color=NAVY, name="Meiryo UI"):
    return Font(name=name, size=size, bold=bold, color=color)


def fill(hex_color):
    return PatternFill("solid", fgColor=hex_color)


def style_header_row(ws, row, cols, bg=BG_HEAD):
    for c in range(1, cols + 1):
        cell = ws.cell(row, c)
        cell.font = font(10, True, WHITE)
        cell.fill = fill(bg)
        cell.alignment = center
        cell.border = thin


def style_body(ws, start, end, cols, row_fill=None):
    for r in range(start, end + 1):
        bg = row_fill(r) if row_fill else (BG_ALT if r % 2 == 0 else WHITE)
        for c in range(1, cols + 1):
            cell = ws.cell(r, c)
            cell.font = font(9)
            cell.alignment = wrap
            cell.border = thin
            cell.fill = fill(bg)


def set_widths(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def merge_title(ws, cols, title, subtitle=None):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=cols)
    ws.cell(1, 1, title).font = font(16, True, WHITE)
    ws.cell(1, 1).fill = fill(NAVY)
    ws.cell(1, 1).alignment = Alignment(vertical="center", horizontal="left", indent=1)
    ws.row_dimensions[1].height = 28
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=cols)
    ws.cell(2, 1, subtitle or "").font = font(9, False, "334155")
    ws.cell(2, 1).fill = fill(BG_NOTE)
    ws.cell(2, 1).alignment = wrap
    ws.row_dimensions[2].height = 48


def judge_fill(val):
    s = str(val)
    if "△△" in s or "✕" in s:
        return BG_GAP
    if "△" in s:
        return BG_WARN
    if "自動" in s:
        return BG_AUTO
    if "○" in s or "不要" in s:
        return BG_OK
    return WHITE


# ─────────────────────────────────────────────
# 1. マッピング表
# ─────────────────────────────────────────────
def build_mapping():
    wb = Workbook()
    ws = wb.active
    ws.title = "マッピング表"

    merge_title(
        ws, 11,
        "贈答品振データ  項目マッピング表",
        "出典：プログラム関連図（西友ギフトシステム）／詳細設計（贈答システム連動02-2）贈答品振データ作成\n"
        "西友ギフト D_Accounting.dat（PKG_BSP_MOVE_CREATE） → TRIAL 振替伝票IF。仕入明細ではない。サンプル D_Accounting_20260729.dat（1,001件）検証済み。",
    )

    headers = [
        "No", "必須", "TRIAL項目名", "TRIAL項目ID", "TRIAL型",
        "西友ギフト項目", "ギフト型/桁", "判定", "変換・設定案", "サンプル実測", "備考",
    ]
    ws.append([])  # row3 spacer conceptually; we'll write header at row 4
    for c, h in enumerate(headers, 1):
        ws.cell(4, c, h)
    style_header_row(ws, 4, 11)

    rows = [
        [1, "○", "振替日", "date", "日付(8)",
         "作成日", "X(8) YYYYMMDD", "○",
         "作成日をそのまま設定", "全件 20260729",
         "設計上はバッチ本日日付+1。実測はバッチ更新日20260728の翌日。会計取込日に近く、実物流完了日ではない。今回は元データの日付を使用する方針。"],
        [2, "△", "伝票ID", "id", "int64(19)",
         "該当なし", "—", "自動",
         "TRIAL側自動採番、または連携管理番号", "—",
         "同一PKが複数ある場合（赤黒等）の識別用。初回連携は空欄可。"],
        [3, "—", "伝票履歴先頭ID", "id_1st", "int64(19)",
         "該当なし", "—", "自動",
         "初回連携時は空欄", "—", "TRIAL仕様：空欄でよい。"],
        [4, "○", "伝票区分", "bill_kindid", "int16(5)",
         "伝種", "X(4) 固定『640 』", "△",
         "ギフト固定値640をTRIAL振替伝票区分へ変換", "全件『640 』",
         "TRIAL振替サンプル全2,461件は伝票区分=30。640→30でよいか要確認。"],
        [5, "○", "伝票番号", "code", "char(10)",
         "伝票№", "X(7) 頭ZERO埋め", "○",
         "7桁をそのまま設定（前ゼロ保持）。TRIAL側は10桁のため左ゼロ埋めまたは右寄せ要確認",
         "1,001件すべて一意。例:2800796",
         "サンプルでは1伝票=1明細。複数明細時は同一伝票番号で行番号を振る。"],
        [6, "○", "振替元店舗コード", "out_org_id", "int32(10)",
         "センター事業所コード", "X(5)", "△",
         "センター事業所コードを振替元店舗コードへ。店舗マスタ統合が必要",
         "06412=850件 / 06414=151件",
         "品振は在庫センター→店舗の社内在庫移動。コード体系がTRIAL店舗コード（サンプルは117,200,714等）と異なる。"],
        [7, "○", "振替元部門コード", "out_itgrp_id", "int32(10)",
         "該当なし", "—", "△△",
         "案A: 商品コード経由でTRIAL商品マスタから部門を取得\n案B: 取引先コード内デパ2桁→TRIALディビジョンCD",
         "ギフト側に部門項目なし",
         "必須項目。方針未確定。"],
        [8, "○", "振替先店舗コード", "in_org_id", "int32(10)",
         "事業所コード", "X(5) ＝店舗コード", "△",
         "事業所コード（店舗）を振替先店舗コードへ。店舗マスタ統合が必要",
         "59店舗。例:02108, 00330",
         "センターコード(3桁, 921/971)はTRIAL項目に直接対応なし。振替元の補助情報として保持要否は未定。"],
        [9, "○", "振替先部門コード", "in_itgrp_id", "int32(10)",
         "該当なし", "—", "△△",
         "振替元部門と同じ導出案", "ギフト側に部門項目なし",
         "サンプルTRIALでは振替元・振替先部門は同一値。ギフト品振も同一部門想定か要確認。"],
        [10, "—", "振替伝票ID管理区分", "itemid_typeid", "int16(5)",
         "データ種（候補）", "X(1) 固定『8』", "△",
         "1:発注振替伝票 2:振替出庫伝票 0:その他。ギフトデータ種8との対応は未確定",
         "全件 データ種=8 / レコード種=2",
         "TRIALサンプルは1が2,366件、0が92件、2が3件。No.10が1か2の場合、出荷番号が必須。"],
        [11, "条件", "振替伝票（出荷）番号", "o_transbill_id", "int64(19)",
         "伝票№（候補）", "X(7)", "△",
         "ギフト品振伝票番号を設定可能か確認", "伝票№と同一候補",
         "itemid_typeid=1 or 2 の場合必須。"],
        [12, "条件", "仕入伝票（親伝）番号", "dlvbill_id", "int64(19)",
         "ファイル上なし", "—", "○",
         "伝票番号と一致、または空欄", "—",
         "総量発注有の場合必須。ギフト品振では空欄候補。"],
        [13, "条件", "取引先コード", "vendor_id", "int32(10)",
         "取引先コード", "X(9)", "△",
         "9桁をTRIAL取引先コードへ変換（マスタ統合）",
         "39取引先。上位:048934950(288), 029900953(116), 035173959(108)",
         "発注振替伝票（itemid_typeid=1）の場合に使用。"],
        [14, "—", "伝票日付", "ref_slip_srch_date", "日付(8)",
         "作成日", "X(8)", "○", "振替日と同値", "全件 20260729", "—"],
        [15, "—", "入庫日", "in_date", "日付(8)",
         "作成日", "X(8)", "○", "振替日と同値（案）", "全件 20260729",
         "TRIALサンプルでは入庫日=振替日+1のケースあり。同日でよいか確認済み方針は「作成日を使用」。"],
        [16, "—", "計上日", "count_date", "日付(8)",
         "作成日", "X(8)", "○", "振替日と同日", "全件 20260729", "—"],
        [17, "—", "便", "mail", "int8(2)",
         "該当なし", "—", "自動",
         "未設定時はTRIALルールで1固定", "ギフトに便項目なし",
         "TRIALサンプル全件 便=1。"],
        [18, "—", "納品経路区分", "dlvroute_typeid", "int16(5)",
         "該当なし", "—", "自動", "空欄可のため空欄", "—", "TRIAL仕様：空欄でよい。"],
        [19, "—", "通過在庫区分", "tcdc_typeid", "int16(5)",
         "該当なし", "—", "自動", "空欄", "—", "—"],
        [20, "—", "納品形式区分", "dlvway_typeid", "int16(5)",
         "該当なし", "—", "自動", "空欄", "—", "—"],
        [21, "—", "物流タイプ", "btype", "int16(5)",
         "該当なし", "—", "△", "固定値候補", "ギフトに項目なし",
         "TRIALサンプル全件 0。ギフト非生鮮固定で0とするか要確認。"],
        [22, "—", "温度帯区分", "thermal_zone_typeid", "int16(5)",
         "該当なし", "—", "△", "非生鮮固定値候補", "ギフトに項目なし",
         "TRIALサンプルは1が2,189件、2が269件。ギフトは非生鮮中心だが、2016年改訂で生鮮フラグ抽出条件は削除済み。"],
        [23, "○", "原価金額（伝票計）", "o_am", "int64(19)",
         "原価単価×数量", "原価:X(8) 6v2 / 数量:X(4)", "○",
         "明細原価金額を伝票単位で合計", "数量1〜30、原単価13.32〜103.00円",
         "原価単価は格納値/100が実額（例:00001420→14.20円）。"],
        [24, "○", "売価金額（伝票計）", "s_am", "int64(19)",
         "売価単価×数量", "売価:X(6) 整数円", "○",
         "明細売価金額を伝票単位で合計", "売単価1,800〜14,500円",
         "売価単価は整数円（例:002100→2,100円）。1000倍しない。"],
        [25, "—", "削除フラグ", "f_del", "int8(1)",
         "該当なし", "—", "自動", "新規時0固定", "—", "—"],
        ["26-30", "—", "登録日/時分秒/者/機能/操作", "ins_*", "各種",
         "バッチ更新日・時間・PGMID", "日付8 / 時刻9 / PGMID 30", "自動",
         "TRIAL側自動設定を基本。ギフトのバッチ更新情報は監査用に破棄またはログ",
         "バッチ更新日=20260728 / PGMID=PKG_BSP_MOVE_CREATE",
         "売場=000固定、普特区分=1固定はTRIALに対応項目なし。"],
        [31, "○", "行番号", "line_no", "int8(3)",
         "該当なし", "—", "○",
         "同一伝票内で商品単位に1からの連番", "本サンプルは1伝票1明細のため全件1",
         "設計上は品振伝票NOで集計。複数明細発生時に連番。"],
        [32, "○", "数量（1000倍）", "qy", "int64(19)",
         "数量（バラ）", "X(4) 頭ZERO埋め", "○",
         "現行数量 × 1,000", "min1 / max30 / 合計1,786 / 0件なし",
         "例: 0001 → 1000。"],
        [33, "○", "原単価（1000倍）", "cost", "int64(19)",
         "原価単価", "X(8) 6v2（×100格納）", "○",
         "実額円 × 1,000。実額 = 格納値 / 100",
         "例:00001420 → 14.20円 → TRIAL 14200",
         "ギフト設計: ABS(原価単価×100) を8桁ZERO埋め。"],
        [34, "○", "原価金額", "trans_am", "int64(19)",
         "原価単価×数量", "算出", "○",
         "実額原単価 × 数量。円未満は四捨五入",
         "例:14.20×1=14.20 → 14円",
         "TRIALサンプルは金額が実額（1000倍しない）。"],
        [35, "○", "売価金額", "sales_am", "int64(19)",
         "売価単価×数量", "算出", "○",
         "売単価（実額円）× 数量",
         "例:2100×1=2100", "—"],
        [36, "○", "売単価", "price", "int32(10)",
         "売価単価", "X(6) 整数円", "○",
         "実値のまま設定（1000倍しない）", "例:002100 → 2100",
         "TRIALサンプルも売単価は実値。"],
        [37, "○", "商品コード", "item_id", "int32(10)",
         "アイテム№", "X(9) 頭0付き9桁", "△△",
         "西友アイテムNo(実質8桁) → スキャンコード/JANへ変換し、TRIAL商品マスタへ事前登録したコードを設定",
         "484商品。全件9桁（先頭0+8桁）。例:033081100",
         "TRIALサンプルの商品コードは13桁JAN。int32最大2,147,483,647のため13桁JANは型オーバーフロー。商品マスタ統合が前提。"],
        [38, "—", "勘定科目コード", "accitem_id", "int16(5)",
         "該当なし", "—", "○", "空白", "—", "—"],
        [39, "—", "振替先商品コード", "in_item_id", "int32(10)",
         "該当なし", "—", "○", "空白", "—", "ダミー項目。"],
        [40, "—", "理由CD", "reason_typeid", "int16",
         "該当なし", "—", "○", "空白", "—", "店内振替理由。ギフト品振は空欄候補。"],
        [41, "—", "振替-仕入紐付（伝票番号）", "dlvbill_id_1st", "int64(19)",
         "ファイル上なし", "—", "○", "空白", "—", "預託用。対象外。"],
        [42, "—", "振替-仕入紐付（行番号）", "dlvitem_line_no", "int8(2)",
         "ファイル上なし", "—", "○", "空白", "—", "預託用。対象外。"],
    ]

    for i, row in enumerate(rows):
        r = 5 + i
        for c, v in enumerate(row, 1):
            ws.cell(r, c, v)
        ws.row_dimensions[r].height = 48
        for c in range(1, 12):
            ws.cell(r, c).font = font(9)
            ws.cell(r, c).alignment = wrap
            ws.cell(r, c).border = thin
            ws.cell(r, c).fill = fill(BG_ALT if i % 2 else WHITE)
        ws.cell(r, 8).fill = fill(judge_fill(row[7]))
        ws.cell(r, 8).alignment = center
        ws.cell(r, 1).alignment = center
        ws.cell(r, 2).alignment = center

    # ギフト余剰
    r0 = 5 + len(rows) + 1
    ws.merge_cells(start_row=r0, start_column=1, end_row=r0, end_column=11)
    ws.cell(r0, 1, "西友ギフト側のみ存在する項目（TRIAL振替伝票に直接対応なし）").font = font(11, True, WHITE)
    ws.cell(r0, 1).fill = fill(TEAL)
    extra = [
        ["—", "—", "—", "—", "—", "レコード種", "X(1) 固定『2』", "不要", "破棄", "全件2", "レコード識別用。"],
        ["—", "—", "—", "—", "—", "センターコード", "X(3)", "△", "振替元の補助。店舗マッピングに使う可能性", "921=850件 / 971=151件", "センター事業所コードとセット（06412-921, 06414-971）。"],
        ["—", "—", "—", "—", "—", "売場", "X(3) 固定『000』", "不要", "破棄", "全件000", "—"],
        ["—", "—", "—", "—", "—", "普特区分", "X(1) 固定『1』", "不要", "破棄", "全件1", "—"],
        ["—", "—", "—", "—", "—", "バッチ更新日", "X(8)", "自動", "TRIAL登録日の代替候補、基本は自動", "全件20260728", "作成日の前日。"],
        ["—", "—", "—", "—", "—", "バッチ更新時間", "X(9) HHMISSFF3", "不要", "破棄", "23:09:20台", "—"],
        ["—", "—", "—", "—", "—", "バッチPGMID", "X(30)", "不要", "破棄", "PKG_BSP_MOVE_CREATE", "—"],
    ]
    for i, row in enumerate(extra):
        r = r0 + 1 + i
        for c, v in enumerate(row, 1):
            ws.cell(r, c, v)
            ws.cell(r, c).font = font(9)
            ws.cell(r, c).alignment = wrap
            ws.cell(r, c).border = thin
            ws.cell(r, c).fill = fill(BG_ALT if i % 2 else WHITE)
        ws.row_dimensions[r].height = 32

    set_widths(ws, [8, 8, 22, 18, 14, 20, 22, 8, 36, 28, 42])
    ws.freeze_panes = "A5"
    ws.auto_filter.ref = f"A4:K{4+len(rows)}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.print_title_rows = "1:4"

    # 概要シート
    ws2 = wb.create_sheet("対象IF概要", 0)
    merge_title(
        ws2, 4,
        "贈答品振データ  対象IF概要",
        "プログラム関連図「◆非生鮮品振情報データ出力◆」／詳細設計 贈答品振データ作成",
    )
    overview = [
        ["項目", "内容"],
        ["出典資料①", "プログラム関連図（西友ギフトシステム）.xlsx  プロセス：贈答品振システム連動"],
        ["出典資料②", "詳細設計（贈答システム連動02-2）贈答品振データ作成.xls  PKG_BSP_MOVE_CREATE.PROC_MOVE_CREATE"],
        ["プログラム", "PKG_BSP_MOVE_CREATE.PROC_MOVE_CREATE"],
        ["出力ファイル", "D_Accounting.dat（固定長121byte、CRLF）"],
        ["TRIAL取込先", "振替伝票IF（標準外部IFレイアウト）"],
        ["業務定義", "在庫センターを振替元、店舗を振替先とする社内在庫移動"],
        ["抽出条件（現行）", "贈答品振WORK全件。2016年ギフトEOL対応で生鮮フラグ条件は削除済み"],
        ["集計単位", "品振伝票NO。数量はSUM、他項目はMAX"],
        ["結合マスタ", "店舗マスタ MST_SHOP / エリア商品マスタ MST_AREA_GOODS / 本日日付 MST_DATE"],
        ["検証サンプル", "D_Accounting_20260729.dat  1,001件  レコード長121byte  全件解析OK"],
        ["サンプル固定値", "伝種=『640 』 / レコード種=2 / データ種=8 / 売場=000 / 普特区分=1"],
        ["サンプル日付", "作成日=20260729、バッチ更新日=20260728（設計どおり +1日）"],
        ["サンプルキー", "伝票№は全件ユニーク（1伝票1明細）。センター2拠点、店舗59、商品484、取引先39"],
        ["対応しないもの", "TRIAL仕入明細（ALLSIRE_MEISAI）／仕入伝票IF。仕入は GiftStock.dat 側"],
    ]
    for r, row in enumerate(overview, 4):
        for c, v in enumerate(row, 1):
            ws2.cell(r, c, v)
            ws2.cell(r, c).font = font(10, bold=(r == 4 or c == 1))
            ws2.cell(r, c).alignment = wrap
            ws2.cell(r, c).border = thin
            if r == 4:
                ws2.cell(r, c).fill = fill(BG_HEAD)
                ws2.cell(r, c).font = font(10, True, WHITE)
            elif c == 1:
                ws2.cell(r, c).fill = fill(BG_NOTE)
        ws2.row_dimensions[r].height = 28
    ws2.merge_cells("A4:B4")
    # fix: overview has 2 cols, title is 4 cols - that's ok
    set_widths(ws2, [22, 88, 12, 12])
    ws2.row_dimensions[4].height = 22
    for r in range(4, 4 + len(overview)):
        ws2.merge_cells(start_row=r, start_column=2, end_row=r, end_column=4)

    path = f"{OUT_DIR}/贈答品振_①_項目マッピング表.xlsx"
    wb.save(path)
    print("wrote", path)


# ─────────────────────────────────────────────
# 2. GAP分析書
# ─────────────────────────────────────────────
def build_gap():
    wb = Workbook()
    ws = wb.active
    ws.title = "GAP分析書"
    merge_title(
        ws, 8,
        "贈答品振データ  GAP分析書",
        "出典：プログラム関連図／詳細設計（贈答品振データ作成）。D_Accounting.dat → TRIAL振替伝票IF。結論は出さず未解決点を列挙。関連QAは④QA一覧を参照。",
    )
    headers = ["GAP No", "優先度", "カテゴリ", "対象項目", "現状（分かっていること）", "問題点（未解決）", "影響", "関連QA"]
    for c, h in enumerate(headers, 1):
        ws.cell(4, c, h)
    style_header_row(ws, 4, 8)

    gaps = [
        ["G01", "★最高", "対象IFの取り違え防止", "IF全体",
         "プログラム関連図の「非生鮮品振情報データ出力」が本IF。社内在庫移動（センター→店舗）。仕入は GiftStock.dat。",
         "仕入明細（ALLSIRE_MEISAI）とは対応先が異なる。振替伝票IFを正とする前提の確認が必要。",
         "マッピング全体", "QA-01"],
        ["G02", "★最高", "商品コード", "item_id ← アイテム№",
         "詳細設計は MST_AREA_GOODS.アイテムNO（9桁）。サンプル484種は先頭0+8桁。TRIALサンプルは13桁JAN。",
         "①アイテムNo→JAN変換未確定。②TRIAL商品マスタ事前登録が必要。③item_id が int32 のため13桁JANは型オーバーフロー。",
         "全明細（必須）", "QA-02"],
        ["G03", "★最高", "部門コード欠落", "out_itgrp_id / in_itgrp_id",
         "項目編集表に部門項目なし。TRIALでは必須。導出案は商品マスタ経由、または取引先コード内デパ2桁。",
         "導出元が未確定。振替元と振替先を同一部門にしてよいかも未確定。",
         "伝票ヘッダ必須", "QA-03"],
        ["G04", "★高", "店舗コード体系", "out_org_id / in_org_id",
         "詳細設計：振替元＝センター事業所コード、振替先＝店舗コード（事業所コード）。サンプルはセンター2拠点・店舗59。",
         "TRIAL店舗コードとの対応表がない。未登録コードの扱いも未定義。",
         "伝票ヘッダ必須", "QA-04"],
        ["G05", "★高", "伝票区分変換", "bill_kindid ← 伝種",
         "項目編集表は伝種『640 』固定。TRIAL振替サンプルは全件 30。",
         "640→30 でよいか未確認。",
         "PK構成要素", "QA-05"],
        ["G06", "★中", "管理区分と出荷番号", "itemid_typeid / o_transbill_id",
         "データ種は固定『8』、レコード種は固定『2』。TRIALは 1=発注振替 / 2=振替出庫 / 0=その他。1または2なら出荷番号必須。",
         "データ種8の落とし先、品振伝票番号の出荷番号流用可否が未確認。",
         "条件付必須", "QA-06 / QA-07"],
        ["G07", "★中", "取引先コード", "vendor_id",
         "詳細設計は MAX(WRK.取引先コード) 9桁。サンプル39取引先。発注振替の場合に使用。",
         "TRIAL取引先マスタとの変換が未整備。",
         "条件付", "QA-08"],
        ["G08", "★中", "日付の意味", "date / in_date / count_date",
         "作成日＝バッチ本日日付+1（詳細設計どおりサンプルも一致）。会計取込日に近く実物流日ではない。今回は元データ日付を使う方針。",
         "入庫日を同日にするか+1にするか未確定。TRIALサンプルには入庫日=振替日+1の例がある。",
         "日付3項目", "QA-09"],
        ["G09", "★中", "物流・温度帯", "btype / thermal_zone_typeid",
         "ギフトに項目なし。2016年改訂で生鮮フラグ抽出条件は削除（全件出力）。プログラム関連図の処理名はなお「非生鮮品振」。",
         "固定値にするか空欄か、非生鮮前提でよいか未確定。",
         "ヘッダ任意", "QA-10"],
        ["G10", "★低", "金額換算", "qy / cost / trans_am / price / sales_am",
         "数量・原単価は×1000、売単価は実額、金額は単価×数量（円未満四捨五入）。原価は6v2（格納/100）。サンプルに0・負値なし。",
         "端数処理・税込税抜がTRIAL定義と一致するか未突き合わせ。",
         "明細金額", "QA-11"],
        ["G11", "★低", "ギフト余剰項目", "売場/普特/バッチ情報/センターコード",
         "売場000・普特1は固定。センターコードはセンター事業所と対で出現（921/971）。",
         "センターコードをマッピングキーに含めるかは未定。他は破棄想定。",
         "破棄/補助", "QA-12"],
        ["G12", "★低", "明細粒度", "line_no / 伝票番号",
         "詳細設計の集計単位は品振伝票NO。本サンプルは1伝票1明細。",
         "本番で同一伝票複数明細が出るか未確認。",
         "明細PK", "QA-13"],
    ]
    for i, row in enumerate(gaps):
        r = 5 + i
        for c, v in enumerate(row, 1):
            ws.cell(r, c, v)
            ws.cell(r, c).font = font(9)
            ws.cell(r, c).alignment = wrap
            ws.cell(r, c).border = thin
            ws.cell(r, c).fill = fill(BG_ALT if i % 2 else WHITE)
        pri = row[1]
        ws.cell(r, 2).fill = fill(BG_GAP if "最高" in pri else BG_WARN if "高" in pri else BG_AUTO if "中" in pri else BG_OK)
        ws.cell(r, 1).alignment = center
        ws.cell(r, 2).alignment = center
        ws.cell(r, 8).alignment = center
        ws.row_dimensions[r].height = 72

    set_widths(ws, [10, 10, 18, 28, 40, 44, 14, 16])
    ws.freeze_panes = "A5"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0

    path = f"{OUT_DIR}/贈答品振_②_GAP分析書.xlsx"
    wb.save(path)
    print("wrote", path)


# ─────────────────────────────────────────────
# 3. 変換ルール定義書
# ─────────────────────────────────────────────
def build_rules():
    wb = Workbook()
    ws = wb.active
    ws.title = "変換ルール定義書"
    merge_title(
        ws, 8,
        "贈答品振データ  変換ルール定義書",
        "出典：詳細設計 項目編集表1／モジュール仕様。確定ルールではない。未確定箇所は関連QAを参照。",
    )
    headers = [
        "No", "ギフト項目", "ギフト形式", "TRIAL項目", "TRIAL型",
        "変換ロジック（案）", "具体例（サンプル）", "注意事項・関連QA",
    ]
    for c, h in enumerate(headers, 1):
        ws.cell(4, c, h)
    style_header_row(ws, 4, 8)

    rules = [
        ["R01", "作成日", "X(8) YYYYMMDD", "date / ref_slip_srch_date / in_date / count_date",
         "日付8",
         "YYYYMMDDをそのまま4項目へ設定（入庫日も同日案）",
         "20260729 → 20260729",
         "実物流日ではない。入庫日はQA-09。"],
        ["R02", "伝種", "X(4) 『640 』固定", "bill_kindid",
         "int16",
         "暫定: 30 を設定（TRIAL振替サンプル全件が30）。640との対応は要確認",
         "『640 』 → 30（仮）",
         "未確定。QA-05。"],
        ["R03", "伝票№", "X(7) 数字ZERO埋め", "code",
         "char(10)",
         "7桁文字列を設定。10桁に合わせる場合は左ゼロ埋め",
         "2800796 → 2800796 または 0002800796",
         "前ゼロ保持方針。桁合わせはTRIAL取込仕様に従う。"],
        ["R04", "センター事業所コード", "X(5)", "out_org_id",
         "int32",
         "ORG_MAP[センター事業所コード] → TRIAL店舗コード。未整備",
         "06412 → ??? / 06414 → ???",
         "マッピングテーブル必須。QA-04。"],
        ["R05", "事業所コード", "X(5)", "in_org_id",
         "int32",
         "ORG_MAP[事業所コード] → TRIAL店舗コード。未整備",
         "02108 → ???",
         "振替先。センターコード(921/971)はキー補助。"],
        ["R06", "（導出）部門", "ファイルになし", "out_itgrp_id / in_itgrp_id",
         "int32",
         "案A: アイテムNo → 商品マスタ → TRIAL部門\n案B: 取引先コードのデパ2桁 → ディビジョンCD\n元・先は同一値とする案",
         "未実装",
         "必須。方針未確定。QA-03。"],
        ["R07", "データ種", "X(1) 『8』固定", "itemid_typeid",
         "int16",
         "暫定未設定。候補: 1（発注振替）",
         "8 → 1（仮）",
         "1または2なら o_transbill_id 必須。QA-06。"],
        ["R08", "伝票№", "X(7)", "o_transbill_id",
         "int64",
         "itemid_typeid が1または2のとき、伝票№を数値化して設定。それ以外は空欄",
         "2800796 → 2800796",
         "流用可否は未確認。QA-07。"],
        ["R09", "取引先コード", "X(9)", "vendor_id",
         "int32",
         "VENDOR_MAP[9桁] → TRIAL取引先。数値化のみでは不足する可能性",
         "035173959 → ???",
         "マスタ統合待ち。QA-08。"],
        ["R10", "数量（バラ）", "X(4) ZERO埋め", "qy",
         "int64",
         "qy = INT(数量) * 1000",
         "0001 → 1000 / 0003 → 3000",
         "負値・0は本サンプルになし。ABS済み前提。"],
        ["R11", "原価単価", "X(8) 6v2（×100格納）", "cost",
         "int64",
         "実額円 = INT(格納値) / 100\ncost = ROUND(実額円 * 1000)",
         "00001420 → 14.20円 → 14200",
         "格納値は原価×100。"],
        ["R12", "原価単価×数量", "算出", "trans_am / o_am",
         "int64",
         "trans_am = ROUND(実額原単価 × 数量)\no_am = 伝票内 trans_am の合計（本サンプルは1明細のため同一）",
         "14.20 × 1 → 14",
         "円未満四捨五入。税抜前提かは未確認。QA-11。"],
        ["R13", "売価単価", "X(6) 整数円", "price",
         "int32",
         "price = INT(売価単価)  ※1000倍しない",
         "002100 → 2100",
         "TRIALサンプルも実値。"],
        ["R14", "売価単価×数量", "算出", "sales_am / s_am",
         "int64",
         "sales_am = 売単価 × 数量\ns_am = 伝票内合計",
         "2100 × 1 → 2100",
         "—"],
        ["R15", "アイテム№", "X(9)", "item_id",
         "int32",
         "ITEM_MAP[アイテムNo] → TRIAL商品コード（JAN等）\n先頭0を除いた8桁でマスタ検索する案",
         "033081100 → （JAN 13桁？）",
         "事前マスタ登録必須。int32制約あり。QA-02。"],
        ["R16", "—", "—", "line_no",
         "int8",
         "同一伝票番号内で1からの連番",
         "本サンプルは全て1",
         "—"],
        ["R17", "—", "—", "mail",
         "int8",
         "未設定のため 1 固定",
         "→ 1",
         "TRIAL仕様。"],
        ["R18", "—", "—", "f_del",
         "int8",
         "0 固定",
         "→ 0",
         "新規連携。"],
        ["R19", "—", "—", "dlvroute / tcdc / dlvway / accitem / in_item / reason / 紐付2項目",
         "任意",
         "空欄",
         "—",
         "TRIAL仕様で空欄可、または預託対象外。"],
        ["R20", "レコード種/売場/普特/バッチ時刻/PGMID", "固定または監査", "（対応なし）",
         "—",
         "受信後破棄",
         "2 / 000 / 1 / PKG_BSP_MOVE_CREATE",
         "センターコードはR04の補助に残す可能性あり。QA-12。"],
        ["R21", "物流タイプ/温度帯", "ファイルになし", "btype / thermal_zone_typeid",
         "int16",
         "暫定: btype=0、温度帯は未設定または非生鮮固定",
         "TRIALサンプル btype=0",
         "未確定。QA-10。"],
    ]
    for i, row in enumerate(rules):
        r = 5 + i
        for c, v in enumerate(row, 1):
            ws.cell(r, c, v)
            ws.cell(r, c).font = font(9)
            ws.cell(r, c).alignment = wrap
            ws.cell(r, c).border = thin
            ws.cell(r, c).fill = fill(BG_ALT if i % 2 else WHITE)
        ws.cell(r, 1).alignment = center
        ws.row_dimensions[r].height = 56

    set_widths(ws, [8, 22, 24, 36, 12, 48, 32, 36])
    ws.freeze_panes = "A5"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0

    # レイアウトシート
    ws3 = wb.create_sheet("ギフト固定長レイアウト")
    merge_title(
        ws3, 7,
        "D_Accounting.dat 固定長レイアウト（121byte）",
        "詳細設計 項目編集表1 ＋ サンプル全1,001件で桁位置を検証済み。",
    )
    h2 = ["No", "項目", "開始", "終了", "桁", "編集元・固定値", "サンプル先頭レコード"]
    for c, h in enumerate(h2, 1):
        ws3.cell(4, c, h)
    style_header_row(ws3, 4, 7, TEAL)
    layout = [
        [1, "伝種", 1, 4, 4, "'640 ' 固定", "640 "],
        [2, "レコード種", 5, 5, 1, "'2' 固定", "2"],
        [3, "データ種", 6, 6, 1, "'8' 固定", "8"],
        [4, "伝票№", 7, 13, 7, "LPAD(品振伝票No,7,'0')", "2800796"],
        [5, "センター事業所コード", 14, 18, 5, "MAX(WRK.センター事業所コード)", "06412"],
        [6, "センターコード", 19, 21, 3, "MAX(WRK.センターコード)", "921"],
        [7, "事業所コード", 22, 26, 5, "MAX(WRK.店舗コード)", "02108"],
        [8, "売場", 27, 29, 3, "'000' 固定", "000"],
        [9, "アイテム№", 30, 38, 9, "MAX(MST.アイテムNO)", "033081100"],
        [10, "数量（バラ）", 39, 42, 4, "ABS(SUM(数量)) ZERO埋め", "0001"],
        [11, "原価単価", 43, 50, 8, "ABS(原価単価×100) 6v2 ZERO埋め", "00001420"],
        [12, "売価単価", 51, 56, 6, "ABS(総額売価単価) ZERO埋め", "002100"],
        [13, "普特区分", 57, 57, 1, "'1' 固定", "1"],
        [14, "取引先コード", 58, 66, 9, "MAX(WRK.取引先コード)", "035173959"],
        [15, "作成日", 67, 74, 8, "バッチ本日日付+1", "20260729"],
        [16, "バッチ更新日", 75, 82, 8, "SYSDATE", "20260728"],
        [17, "バッチ更新時間", 83, 91, 9, "SYSTIMESTAMP HHMISSFF3", "230920445"],
        [18, "バッチPGMID", 92, 121, 30, "RPAD(パッケージ名,30)", "PKG_BSP_MOVE_CREATE"],
    ]
    for i, row in enumerate(layout):
        r = 5 + i
        for c, v in enumerate(row, 1):
            ws3.cell(r, c, v)
            ws3.cell(r, c).font = font(9)
            ws3.cell(r, c).alignment = wrap if c >= 6 else center
            ws3.cell(r, c).border = thin
            ws3.cell(r, c).fill = fill(BG_ALT if i % 2 else WHITE)
        ws3.row_dimensions[r].height = 22
    set_widths(ws3, [6, 22, 8, 8, 6, 40, 28])

    path = f"{OUT_DIR}/贈答品振_③_変換ルール定義書.xlsx"
    wb.save(path)
    print("wrote", path)


# ─────────────────────────────────────────────
# 4. QA一覧
# ─────────────────────────────────────────────
def build_qa():
    wb = Workbook()
    ws = wb.active
    ws.title = "QA一覧"
    merge_title(
        ws, 6,
        "贈答品振データ  QA一覧",
        "出典：プログラム関連図（西友ギフトシステム）／詳細設計（贈答システム連動02-2）贈答品振データ作成。\n"
        "回答欄は空欄。結論は出さず、確認が必要な論点のみ記載。",
    )
    headers = ["ID", "優先度", "関連GAP", "課題・質問概要", "詳細説明", "回答"]
    for c, h in enumerate(headers, 1):
        ws.cell(4, c, h)
    style_header_row(ws, 4, 6)

    qas = [
        ["QA-01", "★最高", "G01",
         "取込先はTRIAL振替伝票IFでよいか",
         "プログラム関連図の「非生鮮品振情報データ出力」（PKG_BSP_MOVE_CREATE / D_Accounting.dat）は、在庫センター→店舗の社内移動である。\n"
         "調査資料でも D_Accounting.dat → TRIAL振替伝票IF と整理されている。\n"
         "仕入は GiftStock.dat → 仕入伝票IF。仕入明細（ALLSIRE_MEISAI）とは別物という理解でよいか。"],
        ["QA-02", "★最高", "G02",
         "アイテムNoをTRIAL商品コードへどう変換するか",
         "詳細設計の取得元は MST_AREA_GOODS.アイテムNO（9桁）。サンプルは先頭0+8桁（例:033081100）。\n"
         "TRIAL振替サンプルの商品コードは13桁JAN。基本商品マスタに西友アイテムNoとスキャンコードがある。\n"
         "確認したいこと：\n"
         "① 変換キーはアイテムNoかスキャンコードか\n"
         "② TRIAL商品マスタへの事前登録は必須か\n"
         "③ TRIAL item_id は int32。13桁JANはオーバーフローするが、どう格納するか"],
        ["QA-03", "★最高", "G03",
         "振替元・振替先の部門コードを何から導出するか",
         "項目編集表に部門項目はない。TRIALでは out_itgrp_id / in_itgrp_id が必須。\n"
         "案A：アイテムNo → TRIAL商品マスタ → 部門コード\n"
         "案B：取引先コード内のデパ2桁 → TRIALディビジョンCD\n"
         "振替元と振替先は同一部門でよいか。"],
        ["QA-04", "★高", "G04",
         "センター事業所コード・店舗コードとTRIAL店舗コードの対応表",
         "詳細設計：センター事業所コード＝振替元、事業所コード（店舗）＝振替先。\n"
         "サンプル：振替元 06412/06414、振替先59店舗。TRIALサンプルの店舗コードは117,200,714等。\n"
         "対応表の提供元、未登録コード受信時の扱い（エラー/スキップ）を確認したい。"],
        ["QA-05", "★高", "G05",
         "伝種『640 』はTRIAL伝票区分30でよいか",
         "項目編集表は伝種『640 』固定。TRIAL振替サンプル2,461件はすべて bill_kindid=30。\n"
         "640→30 の変換でよいか。他区分が必要か。"],
        ["QA-06", "★中", "G06",
         "データ種『8』は itemid_typeid のどれに当たるか",
         "TRIAL定義：1=発注振替伝票、2=振替出庫伝票、0=その他。\n"
         "ギフトはデータ種『8』固定、レコード種『2』固定。\n"
         "1または2を設定すると、振替伝票（出荷）番号が必須になる。"],
        ["QA-07", "★中", "G06",
         "品振伝票番号を出荷番号に流用してよいか",
         "詳細設計の伝票№は LPAD(WRK.品振伝票No,7,'0')。\n"
         "QA-06 で 1 または 2 とした場合、o_transbill_id にこの番号を設定してよいか。"],
        ["QA-08", "★中", "G07",
         "取引先コード9桁のTRIAL変換ルール",
         "MAX(WRK.取引先コード)。サンプル39取引先（例:035173959）。\n"
         "発注振替（itemid_typeid=1）の場合に vendor_id が使われる。マスタ対応表は誰が持つか。"],
        ["QA-09", "★中", "G08",
         "入庫日は作成日と同日か、+1日か",
         "作成日＝バッチ本日日付+1（詳細設計）。会計取込日に近く、実物流完了日ではない。\n"
         "今回のIFは元データの日付を使う方針。\n"
         "TRIALサンプルには入庫日＝振替日+1の例がある。同日でよいか。"],
        ["QA-10", "★中", "G09",
         "物流タイプ・温度帯の固定値",
         "ギフトファイルに項目なし。プログラム関連図の処理名は「非生鮮品振情報データ出力」。\n"
         "一方、変更履歴書（2016）で生鮮フラグ抽出条件は削除されている。\n"
         "btype / thermal_zone_typeid を固定値にするか空欄にするか。非生鮮前提でよいか。"],
        ["QA-11", "★低", "G10",
         "金額の端数・税の扱い",
         "案：数量・原単価は×1000、売単価は実額、金額＝単価×数量（円未満四捨五入）。\n"
         "原価単価は6v2（格納値/100）。税込/税抜、四捨五入位置がTRIALと一致するか。"],
        ["QA-12", "★低", "G11",
         "センターコード3桁をマッピングキーに含めるか",
         "センターコードはセンター事業所コードと対で出現（06412-921、06414-971）。\n"
         "TRIAL振替伝票に直接項目はない。振替元特定の補助キーとするか、破棄するか。"],
        ["QA-13", "★低", "G12",
         "同一品振伝票に複数明細は本番であるか",
         "詳細設計の集計単位は品振伝票NO。サンプル D_Accounting_20260729.dat は1,001件すべて伝票№ユニーク。\n"
         "本番で複数明細が出るなら line_no 連番が必要。実運用の粒度を確認したい。"],
    ]

    for i, row in enumerate(qas):
        r = 5 + i
        vals = row + [""]  # 回答空欄
        for c, v in enumerate(vals, 1):
            ws.cell(r, c, v)
            ws.cell(r, c).font = font(9)
            ws.cell(r, c).alignment = wrap
            ws.cell(r, c).border = thin
            ws.cell(r, c).fill = fill(WHITE if c == 6 else (BG_ALT if i % 2 else WHITE))
        pri = row[1]
        ws.cell(r, 2).fill = fill(
            BG_GAP if "最高" in pri else BG_WARN if "高" in pri else BG_AUTO if "中" in pri else BG_OK
        )
        ws.cell(r, 1).alignment = center
        ws.cell(r, 2).alignment = center
        ws.cell(r, 3).alignment = center
        ws.row_dimensions[r].height = 88

    set_widths(ws, [10, 10, 12, 36, 72, 28])
    ws.freeze_panes = "A5"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.print_title_rows = "1:4"

    path = f"{OUT_DIR}/贈答品振_④_QA一覧.xlsx"
    wb.save(path)
    print("wrote", path)


if __name__ == "__main__":
    build_mapping()
    build_gap()
    build_rules()
    build_qa()
