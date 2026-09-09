#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
受払明細（【倉庫CD】uke.txt）訂正版
  TRIAL振替伝票 / 仕入伝票 → sinops uke.txt
  ①項目マッピング表  ②GAP分析書  ③変更ルール定義書
"""
import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

OUT_DIR = "/Users/treqd/Desktop/Cursor/西友PMI/受払明細_訂正版"
os.makedirs(OUT_DIR, exist_ok=True)

NAVY = "0D2B4E"
BLUE = "1A5C9E"
TEAL = "0F766E"
WHITE = "FFFFFF"
BG_HEAD = "1A5C9E"
BG_ALT = "F1F5F9"
BG_WARN = "FEF3C7"
BG_GAP = "FEE2E2"
BG_OK = "DCFCE7"
BG_AUTO = "E0E7FF"
BG_NOTE = "EEF2FF"
BG_BANNER = "FFF7ED"

thin = Border(
    left=Side(style="thin", color="CBD5E1"),
    right=Side(style="thin", color="CBD5E1"),
    top=Side(style="thin", color="CBD5E1"),
    bottom=Side(style="thin", color="CBD5E1"),
)
wrap = Alignment(wrap_text=True, vertical="center")
center = Alignment(wrap_text=True, vertical="center", horizontal="center")

BANNER = (
    "訂正：旧版は ALLSIRE_MEISAI（全仕入明細）を単一ソースとし、"
    "uke.txt サンプルの伝票区分 IN(20,30,50)/(10,11,12,31) で出庫/入庫を分けていた。"
    "正は TRIAL 標準外部IFの伝票種別を参照する。"
    "出庫＝振替伝票（bill_kindid=30 かつ btype=0）、"
    "入庫＝仕入伝票（bill_kindid=10 かつ btype=0）。"
    "廃棄・その他区分は本資料の対象外。結論は出さず未解決は GAP に残す。"
)

SRC = (
    "sinops：【倉庫CD】uke.txt（10項目）。"
    "TRIAL：標準外部IFレイアウト 振替伝票 / 仕入伝票。"
)


def font(size=10, bold=False, color=NAVY, name="Meiryo UI"):
    return Font(name=name, size=size, bold=bold, color=color)


def fill(hex_color):
    return PatternFill("solid", fgColor=hex_color)


def style_header_row(ws, row, cols, bg=BG_HEAD):
    for c in range(1, cols + 1):
        cell = ws.cell(row, c)
        cell.font = font(10, True, WHITE)
        cell.fill = fill(bg)
        cell.alignment = center
        cell.border = thin


def set_widths(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def merge_title(ws, cols, title, subtitle=None, height=56):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=cols)
    ws.cell(1, 1, title).font = font(16, True, WHITE)
    ws.cell(1, 1).fill = fill(NAVY)
    ws.cell(1, 1).alignment = Alignment(vertical="center", horizontal="left", indent=1)
    ws.row_dimensions[1].height = 28
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=cols)
    ws.cell(2, 1, subtitle or "").font = font(9, False, "334155")
    ws.cell(2, 1).fill = fill(BG_BANNER)
    ws.cell(2, 1).alignment = wrap
    ws.row_dimensions[2].height = height


def judge_fill(val):
    s = str(val)
    if "△△" in s or "✕" in s:
        return BG_GAP
    if "△" in s:
        return BG_WARN
    if "自動" in s or "―" == s or s.startswith("―"):
        return BG_AUTO
    if "○" in s or "不要" in s:
        return BG_OK
    return WHITE


def write_kv_table(ws, start_row, rows, cols=2):
    for i, row in enumerate(rows):
        r = start_row + i
        for c, v in enumerate(row, 1):
            ws.cell(r, c, v)
            ws.cell(r, c).font = font(10, bold=(i == 0 or c == 1))
            ws.cell(r, c).alignment = wrap
            ws.cell(r, c).border = thin
            if i == 0:
                ws.cell(r, c).fill = fill(BG_HEAD)
                ws.cell(r, c).font = font(10, True, WHITE)
            elif c == 1:
                ws.cell(r, c).fill = fill(BG_NOTE)
        ws.row_dimensions[r].height = 28
        if cols > 2:
            ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=cols)
    return start_row + len(rows)


def write_grid(ws, start_row, headers, rows, judge_col=None):
    for c, h in enumerate(headers, 1):
        ws.cell(start_row, c, h)
    style_header_row(ws, start_row, len(headers))
    for i, row in enumerate(rows):
        r = start_row + 1 + i
        for c, v in enumerate(row, 1):
            ws.cell(r, c, v)
            ws.cell(r, c).font = font(9)
            ws.cell(r, c).alignment = wrap
            ws.cell(r, c).border = thin
            ws.cell(r, c).fill = fill(BG_ALT if i % 2 else WHITE)
        if judge_col:
            ws.cell(r, judge_col).fill = fill(judge_fill(row[judge_col - 1]))
            ws.cell(r, judge_col).alignment = center
        ws.cell(r, 1).alignment = center
        ws.row_dimensions[r].height = 52
    return start_row + 1 + len(rows)


# ─────────────────────────────────────────────
# ① マッピング表
# ─────────────────────────────────────────────
def build_mapping():
    wb = Workbook()

    # --- 概要 ---
    ws = wb.active
    ws.title = "対象IF概要"
    merge_title(ws, 4, "受払明細IF  対象IF概要（訂正版）", BANNER, 72)
    overview = [
        ["項目", "内容"],
        ["出力先", "sinops 倉庫系 受払明細  【倉庫CD】uke.txt（カンマ区切り・10項目）"],
        ["方向", "TRIAL → sinops"],
        ["出庫の参照先", "TRIAL 標準外部IF 振替伝票。抽出：伝票区分 bill_kindid = 30 かつ 物流タイプ btype = 0"],
        ["入庫の参照先", "TRIAL 標準外部IF 仕入伝票。抽出：伝票区分 bill_kindid = 10 かつ 物流タイプ btype = 0"],
        ["使わないソース", "ALLSIRE_MEISAI（全仕入明細）を単一ソースにしない。伝票区分 IN リストでの逆引きもしない"],
        ["本資料の対象外", "廃棄、伝票区分 11/12/19/20/31/40/41/50、物流タイプ ≠ 0"],
        ["sinops未使用項目", "デポ / 注文№ / 入荷数 / 伝票№ / 得意先コード → 固定値 0（定義どおり）"],
        ["数量の単位", "TRIAL qy は1000倍格納。sinops 入庫数・売上数は実数。変更ルール：qy ÷ 1000"],
        ["符号", "正値のまま。旧資料の ×(-1) は使わない（ukeサンプル負値0件は参考情報）"],
        ["出典①", "標準外部IFレイアウト_取込.xlsx  シート「振替伝票」「仕入伝票」"],
        ["出典②", "sinops uke.txt 項目定義（処理日/デポ/品番/注文№/発行日/入荷数/入庫数/伝票№/売上数/得意先コード）"],
        ["旧資料", "受払明細_①_項目マッピング表(【倉庫CD】uke.txt)、_発注仕入IF分析_① v7 の ALLSIRE ベース整理"],
        ["ギフトとの関係", "贈答品振も振替伝票（サンプル bill_kindid=30）へ入る。倉庫受払出庫と同一抽出になるかは未整理"],
    ]
    write_kv_table(ws, 4, overview, cols=4)
    set_widths(ws, [22, 88, 12, 12])
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0

    # --- 抽出条件 ---
    ws2 = wb.create_sheet("抽出条件")
    merge_title(
        ws2, 8,
        "受払明細IF  抽出条件（訂正版）",
        "フィルタは TRIAL 伝票側。uke.txt サンプルに出た区分を条件にしない。",
    )
    h2 = ["パターン", "sinops出力", "TRIAL伝票", "伝票区分", "物流タイプ", "削除", "倉庫CDの候補", "備考"]
    r2 = [
        ["出庫", "売上数(No.9) に数量。入庫数=0",
         "振替伝票", "30", "0", "f_del ≠ 1",
         "振替元店舗コード out_org_id（倉庫が出側の場合）",
         "ファイル名の【倉庫CD】と out_org_id / in_org_id の対応は未確定（GAP-04）"],
        ["入庫", "入庫数(No.7) に数量。売上数=0",
         "仕入伝票", "10", "0", "f_del ≠ 1",
         "納品先店舗・センターコード to_org_id、または店舗コード org_id",
         "倉庫入庫が to_org_id なのか org_id なのか未確定（GAP-04）"],
        ["同一キー合算", "1行に入庫数+売上数、または2行",
         "両伝票", "30 と 10", "0", "—",
         "同一倉庫CD＋品番＋処理日",
         "旧ukeサンプルは1行に両方入る例あり。ソース分離後も1行にまとめるかは未確定（GAP-08）"],
        ["対象外", "出力しない",
         "—", "11/12/19/20/31/40/41/50 等", "≠0 含む", "—",
         "—",
         "廃棄（旧19）も含め本訂正では割り当てない"],
    ]
    write_grid(ws2, 4, h2, r2)
    set_widths(ws2, [14, 28, 14, 14, 12, 12, 36, 42])
    ws2.freeze_panes = "A5"

    # --- 出庫 mapping ---
    ws3 = wb.create_sheet("マッピング_出庫")
    merge_title(
        ws3, 12,
        "受払明細 出庫  項目マッピング表",
        "TRIAL 振替伝票（bill_kindid=30 かつ btype=0）→ sinops uke.txt。"
        "数量は売上数へ。入庫数は0。",
        64,
    )
    hm = [
        "No", "必須", "sinops項目", "sinops型/桁", "FMT",
        "TRIAL項目名", "TRIAL項目ID", "TRIAL型", "判定", "変換・設定案", "旧ALLSIRE対応", "備考",
    ]
    out_rows = [
        [1, "○", "処理日", "整数8", "在庫変動日 yyyymmdd",
         "振替日", "date", "dbTsHead_t(8) 必須", "△",
         "案：date を YYYYMMDD で設定",
         "ALLSIRE_MEISAI.処理日",
         "発行日との差分は旧サンプルで43%。振替日=変動日でよいか未確認（GAP-06）"],
        [2, "○", "デポ", "文字100", "未使用",
         "—", "—", "—", "―",
         '固定値 "0"', "固定0", "sinops定義どおり"],
        [3, "○", "品番", "文字100", "商品コード",
         "商品コード", "item_id", "int32_t", "△△",
         "item_id を文字列化して品番へ",
         "ALLSIRE_MEISAI.JAN",
         "振替の item_id は int32。13桁JANは型オーバーフローの可能性（GAP-07）。仕入側 JAN は char"],
        [4, "", "注文№", "文字100", "未使用",
         "—", "—", "—", "―", '固定値 "0"', "固定0", "—"],
        [5, "○", "発行日", "整数8", "変動日/出力日 yyyymmdd",
         "振替日 / 入庫日 / 計上日 / 伝票日付",
         "date / in_date / count_date / ref_slip_srch_date",
         "日付8", "△",
         "未確定。候補は date と同値、または in_date / count_date",
         "ALLSIRE_MEISAI.納品日",
         "sinops定義※2は変動日またはデータ出力日。どの日付を発行日にするか未確認（GAP-06）"],
        [6, "○", "入荷数", "整数8", "未使用",
         "—", "—", "—", "―", '固定値 "0"', "固定0", "出庫行では入荷数も入庫数も0"],
        [7, "○", "入庫数", "数値8.1", "出庫行では未使用",
         "—", "—", "—", "―",
         "0 固定（出庫行）", "出庫シートは0",
         "入庫と1行合算する場合のみ値あり（GAP-08）"],
        [8, "", "伝票№", "文字100", "未使用",
         "伝票番号", "code", "char(10)", "―",
         'sinops未使用のため "0"。code は連携しない',
         "固定0（ALLSIRE.伝票NOも送付不要）",
         "TRIAL code は残るが uke には出さない"],
        [9, "○", "売上数", "数値9.2", "売上数量（出庫数量）",
         "数量（1000倍）", "qy", "int64_t", "△",
         "売上数 = qy ÷ 1000（正値）",
         "ALLSIRE_MEISAI.納品数（出庫系）を正値",
         "1000で割り切れない場合の小数は sinops 8/9桁制約と未突合（GAP-05）。×(-1)しない"],
        [10, "", "得意先コード", "整数100", "未使用",
         "取引先コード", "vendor_id", "int32_t", "―",
         '固定値 "0"', "固定0", "vendor_id は uke に出さない"],
    ]
    write_grid(ws3, 4, hm, out_rows, judge_col=9)
    set_widths(ws3, [6, 6, 14, 12, 18, 28, 28, 16, 8, 32, 28, 40])
    ws3.freeze_panes = "A5"
    ws3.auto_filter.ref = f"A4:L{4 + len(out_rows)}"

    # 振替側フィルタ行
    r0 = 6 + len(out_rows)
    ws3.merge_cells(start_row=r0, start_column=1, end_row=r0, end_column=12)
    ws3.cell(r0, 1, "抽出キー（出庫）— マッピング対象行の条件。sinops項目ではない").font = font(11, True, WHITE)
    ws3.cell(r0, 1).fill = fill(TEAL)
    extra = [
        ["F1", "○", "（抽出）", "—", "伝票区分=30",
         "伝票区分", "bill_kindid", "int16_t", "○",
         "30 のみ", "旧：IN(20,30,50) 等", "振替伝票対応の区分。他区分は出庫に入れない"],
        ["F2", "○", "（抽出）", "—", "物流タイプ=0",
         "物流タイプ", "btype", "int16_t", "○",
         "0 のみ", "旧：フィルタなし", "0の業務意味は本資料では決めない（GAP-12）"],
        ["F3", "○", "（抽出）", "—", "削除以外",
         "削除フラグ", "f_del", "int8_t", "○",
         "1 は出力しない", "—", "—"],
        ["F4", "△", "ファイル名", "—", "【倉庫CD】",
         "振替元店舗コード", "out_org_id", "int32_t", "△",
         "倉庫CDへ変換してファイル分割する案",
         "旧：サンプルファイル名 007452",
         "in_org_id（振替先）を倉庫にするケースの有無は未確認"],
    ]
    write_grid(ws3, r0 + 1, hm, extra, judge_col=9)

    # --- 入庫 mapping ---
    ws4 = wb.create_sheet("マッピング_入庫")
    merge_title(
        ws4, 12,
        "受払明細 入庫  項目マッピング表",
        "TRIAL 仕入伝票（bill_kindid=10 かつ btype=0）→ sinops uke.txt。"
        "数量は入庫数へ。売上数は0。",
        64,
    )
    in_rows = [
        [1, "○", "処理日", "整数8", "在庫変動日 yyyymmdd",
         "納品日", "dlv_ymd", "dbDate_t(8) 必須", "△",
         "案：dlv_ymd を YYYYMMDD で設定",
         "入庫シート：ALLSIRE.納品日 ／ 纏め：処理日",
         "計上日 count_date との使い分け未確認（GAP-06）"],
        [2, "○", "デポ", "文字100", "未使用",
         "—", "—", "—", "―", '固定値 "0"', "固定0", "—"],
        [3, "○", "品番", "文字100", "商品コード",
         "JAN", "item_id", "char", "△",
         "item_id をトリムして品番へ",
         "ALLSIRE_MEISAI.JAN",
         "仕入レイアウト上の名称は JAN。桁・社内コード混在は未確認"],
        [4, "", "注文№", "文字100", "未使用",
         "発注伝票番号", "odrbill_id", "int64_t", "―",
         '固定値 "0"', "固定0", "sinops未使用"],
        [5, "○", "発行日", "整数8", "変動日/出力日 yyyymmdd",
         "納品日 / 計上日", "dlv_ymd / count_date", "日付8", "△",
         "案：処理日と同値（dlv_ymd）。count_date 案もあり",
         "ALLSIRE.納品日（入庫シートは処理日と同値）",
         "GAP-06"],
        [6, "○", "入荷数", "整数8", "未使用",
         "—", "—", "—", "―", '固定値 "0"', "固定0", "入荷実績IF（nyuka.txt）とは別"],
        [7, "○", "入庫数", "数値8.1", "入庫数量",
         "数量（1000倍）", "qy", "int64_t", "△",
         "入庫数 = qy ÷ 1000（正値）",
         "ALLSIRE_MEISAI.納品数（入庫系）を正値",
         "GAP-05。×(-1)しない"],
        [8, "", "伝票№", "文字100", "未使用",
         "伝票番号", "code", "char(10)", "―",
         '固定値 "0"', "固定0", "—"],
        [9, "○", "売上数", "数値9.2", "入庫行では未使用",
         "—", "—", "—", "―",
         "0 固定（入庫行）", "入庫シートは0",
         "合算時のみ値あり（GAP-08）"],
        [10, "", "得意先コード", "整数100", "未使用",
         "仕入先コード", "vendor_id", "char(10)", "―",
         '固定値 "0"', "固定0", "—"],
    ]
    write_grid(ws4, 4, hm, in_rows, judge_col=9)
    set_widths(ws4, [6, 6, 14, 12, 18, 28, 28, 16, 8, 32, 28, 40])
    ws4.freeze_panes = "A5"
    ws4.auto_filter.ref = f"A4:L{4 + len(in_rows)}"

    r1 = 6 + len(in_rows)
    ws4.merge_cells(start_row=r1, start_column=1, end_row=r1, end_column=12)
    ws4.cell(r1, 1, "抽出キー（入庫）— マッピング対象行の条件。sinops項目ではない").font = font(11, True, WHITE)
    ws4.cell(r1, 1).fill = fill(TEAL)
    extra_in = [
        ["F1", "○", "（抽出）", "—", "伝票区分=10",
         "伝票区分", "bill_kindid", "int16_t", "○",
         "10 のみ（仕入・店舗発注に相当）",
         "旧：IN(10,11,12,31)",
         "11 バイヤー / 12 FAX / 31 振入庫 は入れない"],
        ["F2", "○", "（抽出）", "—", "物流タイプ=0",
         "物流タイプ", "btype", "int16_t", "○",
         "0 のみ", "旧：フィルタなし", "GAP-12"],
        ["F3", "○", "（抽出）", "—", "削除以外",
         "削除フラグ", "f_del", "int8_t", "○",
         "1 は出力しない", "—", "—"],
        ["F4", "△", "ファイル名", "—", "【倉庫CD】",
         "納品先店舗・センター / 店舗コード",
         "to_org_id / org_id", "char(10)", "△",
         "倉庫入庫のキーにする項目が未確定",
         "旧：サンプル 007452",
         "GAP-04"],
    ]
    write_grid(ws4, r1 + 1, hm, extra_in, judge_col=9)

    # --- 旧差分 ---
    ws5 = wb.create_sheet("旧整理との差分")
    merge_title(ws5, 5, "旧整理との差分", "对照用。旧は破棄せず、参照先が違うことを示す。")
    hd = ["箇所", "旧（ALLSIRE / ukeサンプル）", "訂正後", "判定", "残件"]
    diffs = [
        ["データソース", "ALLSIRE_MEISAI 単一",
         "出庫＝振替伝票、入庫＝仕入伝票", "訂正", "—"],
        ["出庫フィルタ", "伝票区分 IN (20,30,40,41,50) または (20,30,50)",
         "bill_kindid=30 かつ btype=0", "訂正", "20/40/41/50 は対象外のまま未整理"],
        ["入庫フィルタ", "伝票区分 IN (10,11,12,31)",
         "bill_kindid=10 かつ btype=0", "訂正", "11/12/31 は対象外のまま未整理"],
        ["廃棄", "伝票区分=19、サンプル未検出",
         "本資料の対象外", "範囲外", "廃棄IFが別なのか未整理"],
        ["数量項目", "納品数をそのまま",
         "qy ÷ 1000", "変更", "端数・sinops桁"],
        ["符号", "旧資料×(-1) → サンプルで正値確定",
         "正値のまま（維持）", "維持", "—"],
        ["固定値5項目", "0",
         "0（維持）", "維持", "—"],
        ["非排他1行", "同一品番で入庫数+売上数 32行",
         "ソースが別伝票のため再定義が必要", "要確認", "GAP-08"],
    ]
    write_grid(ws5, 4, hd, diffs)
    set_widths(ws5, [16, 44, 40, 10, 32])

    ws.page_setup.orientation = "landscape"
    for s in (ws2, ws3, ws4, ws5):
        s.page_setup.orientation = "landscape"
        s.page_setup.fitToPage = True
        s.page_setup.fitToWidth = 1
        s.page_setup.fitToHeight = 0
        s.print_title_rows = "1:4"

    path = f"{OUT_DIR}/受払明細_①_項目マッピング表.xlsx"
    wb.save(path)
    print("wrote", path)


# ─────────────────────────────────────────────
# ② GAP
# ─────────────────────────────────────────────
def build_gap():
    wb = Workbook()
    ws = wb.active
    ws.title = "GAP分析書"
    merge_title(
        ws, 8,
        "受払明細IF  GAP分析書（訂正版）",
        BANNER + " 関連は③変更ルール。結論は出さない。",
        72,
    )
    headers = ["GAP No", "優先度", "カテゴリ", "対象", "現状（分かっていること）", "問題点（未解決）", "影響", "関連ルール"]
    for c, h in enumerate(headers, 1):
        ws.cell(4, c, h)
    style_header_row(ws, 4, 8)

    gaps = [
        ["G01", "★最高", "参照先の取り違え", "IF全体",
         "旧版は ALLSIRE_MEISAI と uke サンプルの伝票区分バケットで出庫/入庫を分けていた。",
         "正は振替伝票30・物流0（出庫）と仕入伝票10・物流0（入庫）。旧Excel・口播稿は参照しない。",
         "抽出全体", "R01"],
        ["G02", "★最高", "出庫抽出条件", "bill_kindid / btype",
         "出庫は振替伝票の伝票区分30、物流タイプ0と指示されている。",
         "itemid_typeid（発注振替/振替出庫/その他）でさらに絞るか未確定。ギフト品振（同じく30）との切り分けも未確定。",
         "出庫件数", "R01 / R12"],
        ["G03", "★最高", "入庫抽出条件", "bill_kindid / btype",
         "入庫は仕入伝票の伝票区分10、物流タイプ0と指示されている。",
         "10以外の仕入（11/12）や振入庫31を本当に除外してよいか、業務確認が残る。本資料では入れない。",
         "入庫件数", "R01"],
        ["G04", "★高", "倉庫CD", "ファイル名【倉庫CD】",
         "uke は倉庫単位ファイル。振替は out_org_id / in_org_id、仕入は org_id / to_org_id。",
         "倉庫コードにどれを使うか、店舗コード→倉庫CD変換表があるか未確定。",
         "ファイル分割", "R02"],
        ["G05", "★高", "数量スケール", "売上数 / 入庫数 ← qy",
         "TRIAL qy は1000倍。sinops は実数レンジ。旧ALLSIRE納品数は実数想定だった。",
         "÷1000でよいか、端数、負値・0件の扱いが未突合。",
         "数量", "R04"],
        ["G06", "★高", "日付", "処理日 / 発行日",
         "出庫候補：振替日 date、入庫日 in_date、計上日 count_date。入庫候補：納品日 dlv_ymd、計上日 count_date。",
         "処理日＝在庫変動日、発行日＝出力日、の割り当てが未確定。旧サンプルは43%で不一致。",
         "需要予測日付", "R03"],
        ["G07", "★高", "商品コード型", "品番 ← item_id",
         "振替 item_id は int32。仕入 item_id はレイアウト上 JAN/char。sinops 品番は文字100。",
         "13桁JANを振替 int32 に載せている場合オーバーフロー。文字列化規則が未確定。",
         "品番", "R05"],
        ["G08", "★中", "行の粒度", "同一倉庫+品番+処理日",
         "旧ukeは同一行に入庫数と売上数が両方入る例（1.4%）があった。今回ソースは別伝票。",
         "1行に合算するか、出庫行と入庫行を分けるか未確定。",
         "レコード件数", "R06"],
        ["G09", "★中", "対象外区分", "11/12/19/20/31/40/41/50、btype≠0",
         "旧バケットに含まれていた区分がある。廃棄19はサンプル未検出だった。",
         "本IFに出さない前提でよいか、別IFがあるか未整理。",
         "カバレッジ", "R01"],
        ["G10", "★中", "削除伝票", "f_del",
         "両伝票に削除フラグがある。",
         "1を除外する案。取消・赤黒の扱い（正負相殺か出力しないか）は未確認。",
         "件数", "R01"],
        ["G11", "★中", "ギフト品振との重複", "振替伝票 30",
         "贈答品振 D_Accounting.dat も振替伝票30へ入る整理がある。",
         "倉庫uke出庫の抽出にギフト分が混ざるか、除外キーがあるか未整理。",
         "出庫件数", "R12"],
        ["G12", "★中", "物流タイプ0", "btype",
         "抽出条件に btype=0 が指定されている。レイアウト項目は物流タイプ。",
         "0の区分定義（店直等）は本資料では断定しない。0以外を出さない根拠の文書確認が残る。",
         "抽出", "R01"],
        ["G13", "★低", "固定値項目", "デポ等5項目",
         "sinops定義は未使用・固定0。旧サンプル全行0。",
         "維持でよいか（再確認程度）。",
         "レイアウト", "R07"],
    ]
    for i, row in enumerate(gaps):
        r = 5 + i
        for c, v in enumerate(row, 1):
            ws.cell(r, c, v)
            ws.cell(r, c).font = font(9)
            ws.cell(r, c).alignment = wrap
            ws.cell(r, c).border = thin
            ws.cell(r, c).fill = fill(BG_ALT if i % 2 else WHITE)
        pri = row[1]
        ws.cell(r, 2).fill = fill(
            BG_GAP if "最高" in pri else BG_WARN if "高" in pri else BG_AUTO if "中" in pri else BG_OK
        )
        ws.cell(r, 1).alignment = center
        ws.cell(r, 2).alignment = center
        ws.cell(r, 8).alignment = center
        ws.row_dimensions[r].height = 68

    set_widths(ws, [10, 10, 18, 22, 42, 44, 14, 14])
    ws.freeze_panes = "A5"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0

    path = f"{OUT_DIR}/受払明細_②_GAP分析書.xlsx"
    wb.save(path)
    print("wrote", path)


# ─────────────────────────────────────────────
# ③ 変更ルール
# ─────────────────────────────────────────────
def build_rules():
    wb = Workbook()
    ws = wb.active
    ws.title = "変更ルール定義書"
    merge_title(
        ws, 8,
        "受払明細IF  変更ルール定義書（訂正版）",
        "確定ルールではない。抽出・変換の案。未確定は GAP を参照。",
        56,
    )
    headers = [
        "No", "区分", "対象", "変更前（旧整理）",
        "変更後（案）", "具体例", "注意・GAP", "適用",
    ]
    for c, h in enumerate(headers, 1):
        ws.cell(4, c, h)
    style_header_row(ws, 4, 8)

    rules = [
        ["R01", "抽出", "出庫/入庫のソースとフィルタ",
         "ALLSIRE_MEISAI。出庫 IN(20,30,50)、入庫 IN(10,11,12,31)",
         "出庫：振替伝票 WHERE bill_kindid=30 AND btype=0 AND f_del<>1\n"
         "入庫：仕入伝票 WHERE bill_kindid=10 AND btype=0 AND f_del<>1",
         "30+0 → uke 売上数行\n10+0 → uke 入庫数行",
         "G01 G02 G03 G09 G10 G12", "必須"],
        ["R02", "ファイル", "【倉庫CD】uke.txt の分割キー",
         "サンプル名 007452uke.txt（倉庫コード文字列）",
         "案A 出庫：out_org_id を倉庫CDへ変換\n案B 入庫：to_org_id を倉庫CDへ変換\n未変換コードは出力しない/エラーは未定",
         "out_org_id=??? → 007452uke.txt",
         "G04", "必須・未確定"],
        ["R03", "日付", "処理日・発行日",
         "出庫：処理日=ALLSIRE.処理日、発行日=納品日\n入庫：両方納品日",
         "出庫処理日：振替日 date（案）\n出庫発行日：date 同値 または in_date / count_date（未定）\n"
         "入庫処理日：納品日 dlv_ymd（案）\n入庫発行日：dlv_ymd 同値（案）",
         "20260525 / 20260524 のようなズレをどう出すか未定",
         "G06", "必須・未確定"],
        ["R04", "数量", "売上数・入庫数",
         "ALLSIRE.納品数を正値のまま",
         "出庫 売上数 = qy / 1000、入庫数=0\n入庫 入庫数 = qy / 1000、売上数=0\n符号は正。×(-1)しない",
         "qy=3000 → 3",
         "G05", "必須"],
        ["R05", "品番", "品番 ← item_id",
         "ALLSIRE.JAN 13桁をそのまま",
         "文字列化して TRIM。振替 int32 は桁制約あり",
         "JAN 13桁を int32 に載せられない",
         "G07", "必須・未確定"],
        ["R06", "粒度", "1行のキー",
         "サンプル上、同一品番で入庫数と売上数が同居しうる",
         "案A：出庫行と入庫行を別レコード\n案B：倉庫CD+品番+処理日で合算し1行に両方セット",
         "案Aなら非排他問題は行分割で消える",
         "G08", "未確定"],
        ["R07", "固定値", "デポ/注文№/入荷数/伝票№/得意先コード",
         "0",
         "0 のまま",
         "全項目 0",
         "G13", "維持"],
        ["R08", "出力しない", "TRIAL余剰",
         "ALLSIREの伝票NO等は送付不要と整理済み",
         "code / vendor_id / 金額 / 部門 等は uke に出さない",
         "—",
         "—", "維持"],
        ["R09", "単位換算以外の計算", "金額",
         "uke に金額項目なし",
         "o_am / cost / price は連携しない",
         "—",
         "—", "対象外"],
        ["R10", "文字形式", "日付・品番",
         "yyyymmdd / JAN",
         "日付は8桁数字（ゼロ埋め）。品番は左スペース除去",
         "20260525",
         "—", "案"],
        ["R11", "0数量", "qy=0",
         "旧サンプル 77行で入庫数=売上数=0",
         "抽出後 qy=0 を出すか捨てるか未定。発行日 00000000 はしない（未確認）",
         "—",
         "旧GAP-6相当", "未確定"],
        ["R12", "他IFとの境界", "ギフト品振",
         "ギフトは振替伝票へ（別資料）",
         "uke出庫抽出からギフトを除外する条件は未定義。本ルールでは振替30+物流0をそのまま対象とする案",
         "—",
         "G11", "未確定"],
    ]
    for i, row in enumerate(rules):
        r = 5 + i
        for c, v in enumerate(row, 1):
            ws.cell(r, c, v)
            ws.cell(r, c).font = font(9)
            ws.cell(r, c).alignment = wrap
            ws.cell(r, c).border = thin
            ws.cell(r, c).fill = fill(BG_ALT if i % 2 else WHITE)
        ws.cell(r, 1).alignment = center
        ws.cell(r, 8).alignment = center
        ws.row_dimensions[r].height = 72

    set_widths(ws, [8, 12, 24, 36, 48, 28, 18, 14])
    ws.freeze_panes = "A5"

    # 処理フロー
    ws2 = wb.create_sheet("処理フロー（案）")
    merge_title(ws2, 4, "受払明細 生成フロー（案）", "実装手順のメモ。未確定ステップは GAP 待ち。")
    flow = [
        ["Step", "処理", "入力", "出力/条件"],
        ["1", "振替伝票を読む", "標準外部IF 振替伝票", "全明細"],
        ["2", "出庫抽出", "Step1", "bill_kindid=30 AND btype=0 AND f_del<>1"],
        ["3", "仕入伝票を読む", "標準外部IF 仕入伝票", "全明細"],
        ["4", "入庫抽出", "Step3", "bill_kindid=10 AND btype=0 AND f_del<>1"],
        ["5", "倉庫CD変換", "out_org_id / to_org_id 等", "未確定（G04）"],
        ["6", "項目変換", "date/dlv_ymd, item_id, qy", "処理日・発行日・品番・数量（R03–R05）"],
        ["7", "行の結合", "出庫行＋入庫行", "別行 or キー合算（G08）"],
        ["8", "固定値セット", "—", "デポ等=0"],
        ["9", "ファイル出力", "倉庫CDごと", "【倉庫CD】uke.txt 10項目 CSV"],
    ]
    write_grid(ws2, 4, flow[0], flow[1:])
    set_widths(ws2, [10, 18, 36, 48])

    ws3 = wb.create_sheet("ukeレイアウト")
    merge_title(ws3, 6, "sinops 【倉庫CD】uke.txt 10項目", "出力側。入力は振替/仕入。")
    lay = [
        ["No", "項目", "必須", "型", "出庫での値", "入庫での値"],
        ["1", "処理日", "○", "整数8", "振替日（案）", "納品日（案）"],
        ["2", "デポ", "○", "文字", "0", "0"],
        ["3", "品番", "○", "文字100", "item_id", "item_id（JAN）"],
        ["4", "注文№", "", "文字", "0", "0"],
        ["5", "発行日", "○", "整数8", "未確定", "納品日同値（案）"],
        ["6", "入荷数", "○", "整数", "0", "0"],
        ["7", "入庫数", "○", "数値", "0（合算時は入庫qy/1000）", "qy/1000"],
        ["8", "伝票№", "", "文字", "0", "0"],
        ["9", "売上数", "○", "数値", "qy/1000", "0（合算時は出庫qy/1000）"],
        ["10", "得意先コード", "", "整数", "0", "0"],
    ]
    write_grid(ws3, 4, lay[0], lay[1:])
    set_widths(ws3, [8, 16, 8, 12, 36, 36])

    for s in (ws, ws2, ws3):
        s.page_setup.orientation = "landscape"
        s.page_setup.fitToPage = True
        s.page_setup.fitToWidth = 1
        s.page_setup.fitToHeight = 0

    path = f"{OUT_DIR}/受払明細_③_変更ルール定義書.xlsx"
    wb.save(path)
    print("wrote", path)


if __name__ == "__main__":
    build_mapping()
    build_gap()
    build_rules()
