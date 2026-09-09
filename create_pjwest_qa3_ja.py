#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Create the Japanese WMS + automatic replenishment QA summary."""

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from create_pjwest_qa3 import (
    BLUE,
    DARK_BLUE,
    INK,
    MUTED,
    add_bullet,
    add_callout,
    add_header_footer,
    add_number,
    add_table,
    add_text,
    configure_styles,
)


OUT = "/Users/treqd/Desktop/Cursor/西友PMI/PJWest_WMS_Sinops_QA_ReSummary_20260901_QA3_JA.docx"


def prevent_row_split(table):
    """Keep compact decision tables readable when a page break occurs."""
    for row in table.rows:
        tr_pr = row._tr.get_or_add_trPr()
        if tr_pr.find(qn("w:cantSplit")) is None:
            tr_pr.append(OxmlElement("w:cantSplit"))


def build():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.85)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.35)
    add_header_footer(section)
    configure_styles(doc)

    add_text(doc, "PROJECT BRIEF", size=11, color=BLUE, bold=True, before=4, after=8)
    add_text(doc, "WMS・自動補充QA課題再整理サマリ", size=24, color=RGBColor(0, 0, 0), bold=True, after=4)
    add_text(doc, "最新課題管理表（2026-08-28更新）に基づく設計・開発判断", size=13, color=MUTED, after=12)
    add_text(doc, "日付: 2026-09-01", size=10.5, color=MUTED, after=2)
    add_text(doc, "対象: WMS（Hitluster）／自動補充（Sinops）", size=10.5, color=MUTED, after=2)
    add_text(doc, "目的: QA回答状況、残課題、システム間依存および今後の設計・開発着手条件を明確化", size=10.5, color=MUTED, after=10)
    add_callout(doc, "総合判断：QA回答は全体として収束しているが、「完了」は実装仕様の完全確定を意味しない。WMSとSinopsは、残課題および外部依存を分けて段階的に設計・開発へ進める。")

    doc.add_heading("1. QA状況の概要", level=1)
    add_text(doc, "最新管理表のステータスはQA管理状況を示すものであり、設計・開発の完了度と同義ではない。WMSとSinopsには「完了」後も継続管理が必要な課題が残っている。", after=6)
    add_table(
        doc,
        ["システム", "QA総数", "現在の状況", "設計・開発判断"],
        [
            ["WMS（Hitluster）", "63", "完了53；対応中4；中止6", "詳細設計は拡大可能。全体開発は重要方針項目の確定が前提"],
            ["自動補充（Sinops）", "80", "完了75；保留2；対応中1；空白1；中止1", "GCS直接IFは段階推進可能。Master/DWHと棚割は条件付き推進"],
        ],
        [1.35, 0.8, 2.25, 2.1],
    )
    add_text(doc, "注：同一課題の追加質問もQA行として記録している。「中止」は代替案または正式な対象外範囲を設計資料に明記する。", size=9.5, color=MUTED, italic=True, after=8)

    doc.add_heading("2. 課題1：WMS（Hitluster）", level=1)
    add_callout(doc, "WMS：QA 53/63件が完了、4件が対応中、6件が中止。主なリスクは商品Masterの物流属性、在庫の全量／差分処理、原価方針および外部システムの責任分界に集中している。", fill="FFF6E5", bar="A67C00")
    doc.add_heading("2.1 確認済みで設計入力にできる内容", level=2)
    add_bullet(doc, "商品階層は「部門＞大分類＞中分類＞小分類＞細分類」とし、基本商品Masterをデータソースとする。商品サンプルおよび一部分類資料は取得済み。")
    add_bullet(doc, "店舗Master、取引先Master、商品分類Masterのサンプルを取得済み。荷主コードは固定値0000、企業コードも固定値設定が可能。")
    add_bullet(doc, "TC店舗入荷でも店舗振替データが必要。入力連番は予定と実績の突合に使用。西友に預託在庫はなく、横持ちはセンター間の物流転送として扱う。")
    add_bullet(doc, "在庫一覧は日次の全件データ、在庫調整は発生時の差分データ。既存レコードがない負数調整は新規登録する。")
    add_bullet(doc, "倉庫コードは原則として西友の6桁を維持。更新時刻はデータ受信時刻。調整コードは基幹側で差異理由を識別し、WMSでは管理しない。")

    doc.add_heading("2.2 現在、優先してクローズすべきWMS課題", level=2)
    add_table(
        doc,
        ["QA", "課題", "状況", "影響"],
        [
            ["No.188", "商品Masterの発注単位数設定ルール", "対応中", "商品MasterのMappingおよび数量計算に影響"],
            ["No.189", "商品Masterのケース入数設定ルール", "対応中", "Case/BallおよびWMS保守範囲の確定が必要"],
            ["No.196", "センター返品・廃棄の実登録システム", "対応中", "WMSのIF範囲と業務フローに影響"],
            ["No.198", "DC入荷予定の出荷単位数定義", "対応中", "TRIAL運用の整理と数量定義の統一が必要"],
        ],
        [0.8, 2.55, 0.85, 2.3],
    )
    add_text(doc, "No.188とNo.189は商品Master詳細設計の凍結前にクローズ。No.196とNo.198は業務、MD基幹、TC/WMSで責任システムを確認する。", size=9.5, color=MUTED, after=8)

    doc.add_heading("2.3 完了だが継続管理が必要なWMS残課題", level=2)
    add_bullet(doc, "物流コントローラ依存：DC Location、Lot管理、パレット関連項目、商品分類階層、TC総量入荷の集約変換。")
    add_bullet(doc, "Shinise依存：便区分の保守、および一部の商品物流属性の補完。")
    add_bullet(doc, "マスタ依存：出荷先コード、店舗／センターコードを現行体系で継続するかをマスタ統合PJ方針と整合。")
    add_bullet(doc, "原価・在庫：移動平均単価の算出主体、WMSが提供する基礎項目、全量／差分の適用順序、排他、リトライ、全量欠損時のフォールバック。")
    add_bullet(doc, "ITFコード：1商品コードに複数ITFコードが存在する場合の同期方式。")

    doc.add_heading("3. 課題2：自動補充（Sinops）", level=1)
    add_callout(doc, "自動補充：QA 75/80件が完了。ただし、11本のMaster IFは「GCS→西友DWH→Sinops」、14本は「GCS→Sinops」。DWH追加経路と棚割経路が主要な設計変数。")
    doc.add_heading("3.1 経路および対象範囲の前提", level=2)
    add_bullet(doc, "25本のIFはすべて最初にGCSへアップロードし、14本はGCSからSinopsへ直接連携する。")
    add_bullet(doc, "11本のMaster IFはGCSの後に西友DWHへ取り込み、DWHで処理した上でSinopsへ転送する。DWHの受信、変換、スケジュール、再送、監視、照合を個別に設計する。")
    add_bullet(doc, "11本のMasterは「1本の業務IF＋1つのDWH追加経路」として見積もり、GCSアップロードとDWH転送を2本の完全IFとして重複計上しない。")

    doc.add_heading("3.2 確認済みで優先推進できる内容", level=2)
    add_bullet(doc, "販売実績、来客数、発注勧告、発注実績、休日設定、受払明細などは、主要項目、タイミング、全量／差分、数量変換の方向が概ね明確。")
    add_bullet(doc, "発注勧告は発注日で当日分／翌日以降分を区分。発注数量はSinopsの発注バラ数と発注単位をもとにTRIAL側で固定ルールにより変換する。")
    add_bullet(doc, "入荷実績・入荷予定の便区分、在庫修正などは方向性があり、一部の実装責任はShiniseへ移管。")
    add_bullet(doc, "新商品在庫データはTRIAL基幹で生成し、自動発注対象Masterを整備する。QA完了後も実装タスクとして追跡する。")

    doc.add_heading("3.3 現在、優先してクローズすべき自動補充課題", level=2)
    add_table(
        doc,
        ["QA", "課題", "状況", "影響"],
        [
            ["No.9", "商品Master統合および店内コード対応", "未回答／空白", "11本のMaster MappingとDWHデータ契約に直接影響"],
            ["No.56/57", "仕入先Masterと倉庫の関係、データ量および分割方式", "保留", "DWHでの倉庫分割、性能、ファイル出力責任に影響"],
            ["No.74", "棚割明細の最終連携経路", "対応中", "Sinops直連携、商談／棚割DB、DWH経路が未統一"],
        ],
        [0.95, 2.35, 0.9, 2.3],
    )

    doc.add_heading("3.4 完了だが継続管理が必要な自動補充残課題", level=2)
    add_bullet(doc, "Shinise依存：入荷実績／入荷予定の便区分、在庫修正処理。")
    add_bullet(doc, "Case／バラ変換Master：保守主体と店舗差異を商品単位データへ集約するルール。")
    add_bullet(doc, "棚割および新商品在庫：DWH参照方式、棚割データソース、TRIAL基幹生成ロジック、自動発注対象Masterの担当者・期限・受入成果物。")
    add_bullet(doc, "仕入先と倉庫の関係：調達Masterから推定可能だが、データ量が大きく、DWH側／Sinops側のどちらで倉庫分割するかを確定。")

    doc.add_heading("4. 領域横断の依存関係", level=1)
    add_text(doc, "以下はWMS、自動補充、BOに共通して影響するため、単一IFのMapping表だけでは解決できない。", after=5)
    add_bullet(doc, "区分管理Master、商品／店舗／仕入先コード体系および桁数制約。")
    add_bullet(doc, "GCSフォルダ、授受ファイル、アーカイブ、再送、監視、照合ルール。")
    add_bullet(doc, "物流コントローラとShiniseによる項目補完、便区分、DC Location、Case/Pack、センター物流変換。")
    add_bullet(doc, "MD基幹および業務側による移動平均単価、仕入伝票、税区分、返品／廃棄業務の最終方針。")
    add_callout(doc, "管理原則：QA回答状況と実装残課題は分けて管理する。残課題には少なくとも担当者、期限、設計成果物、依存システム、受入条件を記録する。")

    doc.add_heading("5. 設計・開発着手の判断", level=1)
    decision_table = add_table(
        doc,
        ["対象", "着手判断", "開始条件"],
        [
            ["WMS基礎Masterおよび確認済みIF", "GO", "確認済みサンプル、コード、形式を使用。残課題は後続タスクとして登録"],
            ["WMS在庫、原価、物流コントローラ関連IF", "条件付きGO", "全量／差分、異常、原価、外部システム責任を先に確定"],
            ["Sinops GCS直接IF（14本）", "段階GO", "項目、タイミング、GCS共通フレームワークを確認"],
            ["Sinops Master DWH IF（11本）", "条件付きGO", "No.9、No.56/57をクローズ、または仮定と手戻り範囲を明記"],
            ["棚割および新商品在庫IF", "保留／条件付き", "最終経路、DWHソース、自動発注対象Master方針を統一"],
        ],
        [1.75, 1.15, 3.6],
    )
    prevent_row_split(decision_table)

    doc.add_heading("6. 今後のアクション", level=1)
    add_number(doc, "WMS No.188、No.189、No.196、No.198をクローズし、商品Master、返品／廃棄、DC入荷予定の設計資料へ反映する。")
    add_number(doc, "自動補充No.9、No.56/57、No.74をクローズ。直ちにクローズできない場合は、設計仮定、担当者、期限、変更時の手戻り範囲を文書化する。")
    add_number(doc, "共通データ契約（Key、項目ソース、コード表、デフォルト値、タイミング、全量／差分、再送、冪等性、照合、異常復旧）を先に凍結する。")
    add_number(doc, "11本のMaster DWH追加経路を独立ワークパッケージとして管理し、DWH作業をGCSアップロードまたはSinops IF工数に隠さない。")
    add_number(doc, "中止QAには中止理由と代替課題番号を追記し、完了だが残課題ありの記録は追跡可能な実装タスクへ変換する。")
    add_number(doc, "結合テスト、総合テスト、UAT、移行リハーサル、切替、安定化支援は本フェーズの見積対象外とし、別途見積もる。")

    doc.add_heading("7. 参照資料", level=1)
    add_bullet(doc, "情報システムPMI_MD基幹統合PJ_TRE_課題管理表.xlsx（最新QA管理表、表内更新日2026-08-28）")
    add_bullet(doc, "PJWest_WMS_Sinops_QA_ReSummary_20260901_QA3.docx（中文版の構成・表現を参照）")
    add_bullet(doc, "03. WMS (Hitluster)-20260825T065143Z-1-001.zip")
    add_bullet(doc, "01. 自動補充（Sinops）-20260825T065143Z-1-001.zip")
    add_bullet(doc, "PJWest_情報システム_MD基幹統合PJ _週次定例_20260826.pptx")

    doc.core_properties.title = "WMS・自動補充QA課題再整理サマリ"
    doc.core_properties.subject = "最新QA状況、残課題および設計・開発着手判断"
    doc.core_properties.author = "PMI MD Integration Project"
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
