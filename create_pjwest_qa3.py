#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Create the latest WMS + automatic replenishment QA summary."""

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = "/Users/treqd/Desktop/Cursor/西友PMI/PJWest_WMS_Sinops_QA_ReSummary_20260901_QA3.docx"

BLUE = RGBColor(0x2E, 0x74, 0xB5)
DARK_BLUE = RGBColor(0x1F, 0x4D, 0x78)
INK = RGBColor(0x1F, 0x2D, 0x3D)
MUTED = RGBColor(0x70, 0x7D, 0x8D)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)


def set_run_font(run, name="Arial Unicode MS", size=10.5, color=INK, bold=False, italic=False):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for key in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(key), name)
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.bold = bold
    run.font.italic = italic


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_cell_border(cell, color="D9E0E8", size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        el = borders.find(tag)
        if el is None:
            el = OxmlElement(f"w:{edge}")
            borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), size)
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)


def set_table_geometry(table, widths_in):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    total_dxa = int(round(sum(widths_in) * 1440))
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total_dxa))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    grid = tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_in:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(int(round(width * 1440))))
        grid.append(col)
    for row in table.rows:
        for i, cell in enumerate(row.cells):
            cell.width = Inches(widths_in[i])
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(int(round(widths_in[i] * 1440))))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            set_cell_border(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_field(paragraph, field):
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = field
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_begin, instr, fld_sep, text, fld_end])
    set_run_font(run, name="Arial", size=9, color=MUTED)


def configure_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = "Arial Unicode MS"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial Unicode MS")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = INK
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.08

    for name, size, color, before, after in (
        ("Heading 1", 16, BLUE, 14, 7),
        ("Heading 2", 13, BLUE, 10, 5),
        ("Heading 3", 11.5, DARK_BLUE, 8, 4),
    ):
        st = doc.styles[name]
        st.font.name = "Arial Unicode MS"
        st._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial Unicode MS")
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = color
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)
        st.paragraph_format.keep_with_next = True

    for name in ("List Bullet", "List Number"):
        st = doc.styles[name]
        st.font.name = "Arial Unicode MS"
        st._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial Unicode MS")
        st.font.size = Pt(10.5)
        st.font.color.rgb = INK
        st.paragraph_format.left_indent = Inches(0.28)
        st.paragraph_format.first_line_indent = Inches(-0.18)
        st.paragraph_format.space_after = Pt(3)
        st.paragraph_format.line_spacing = 1.08


def add_text(doc, text, size=10.5, color=INK, bold=False, italic=False, align=None, before=0, after=6):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    r = p.add_run(text)
    set_run_font(r, size=size, color=color, bold=bold, italic=italic)
    return p


def add_rich_text(doc, parts, after=6, before=0):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    for text, kwargs in parts:
        r = p.add_run(text)
        set_run_font(r, **kwargs)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(text)
    set_run_font(r, size=10.5, color=INK)
    return p


def add_number(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(text)
    set_run_font(r, size=10.5, color=INK)
    return p


def add_callout(doc, text, fill="E8EEF5", bar="2E74B5"):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [6.5])
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    left = borders.find(qn("w:left")) if borders is not None else None
    if left is None:
        if borders is None:
            borders = OxmlElement("w:tcBorders")
            tc_pr.append(borders)
        left = OxmlElement("w:left")
        borders.append(left)
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "22")
    left.set(qn("w:space"), "0")
    left.set(qn("w:color"), bar)
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    set_run_font(r, size=10.5, color=DARK_BLUE, bold=True)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_table(doc, headers, rows, widths, header_fill="F2F4F7", font_size=9.5):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        set_cell_shading(cell, header_fill)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(h)
        set_run_font(r, size=font_size, color=DARK_BLUE, bold=True)
    for row_data in rows:
        row = table.add_row()
        for i, value in enumerate(row_data):
            cell = row.cells[i]
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.03
            r = p.add_run(str(value))
            set_run_font(r, size=font_size, color=INK)
    set_table_geometry(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_header_footer(section):
    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    hp.paragraph_format.space_after = Pt(0)
    for idx, label in enumerate(("MD", "PJ", "QA")):
        if idx:
            sep = hp.add_run("    ")
            set_run_font(sep, name="Arial", size=9, color=MUTED)
        r = hp.add_run(label)
        set_run_font(r, name="Arial", size=9, color=MUTED, bold=True)

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    fp.paragraph_format.space_before = Pt(0)
    fp.paragraph_format.space_after = Pt(0)
    r = fp.add_run("Page ")
    set_run_font(r, name="Arial", size=9, color=MUTED)
    add_field(fp, "PAGE")


def build():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.85)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.35)
    add_header_footer(section)
    configure_styles(doc)

    add_text(doc, "PROJECT BRIEF", size=11, color=BLUE, bold=True, before=4, after=8)
    add_text(doc, "WMS・自动补充QA课题再整理总结", size=24, color=RGBColor(0, 0, 0), bold=True, after=4)
    add_text(doc, "基于最新課題管理表（2026-08-28更新）整理的设计与开发判断", size=13, color=MUTED, after=12)
    add_text(doc, "日期: 2026-09-01", size=10.5, color=MUTED, after=2)
    add_text(doc, "对象: WMS（Hitluster）／自动补充（Sinops）", size=10.5, color=MUTED, after=2)
    add_text(doc, "目的: 明确QA关闭状态、残课题、跨系统依赖及后续设计开发入口", size=10.5, color=MUTED, after=10)
    add_callout(doc, "总体判断：QA回答整体已大幅收敛，但“完了”不等于实现规格全部闭环；WMS与Sinops均应按残课题和外部依赖分批进入设计开发。")

    doc.add_heading("1. QA状态概览", level=1)
    add_text(doc, "最新管理表的状态数字反映QA管理状态，不直接等同于设计开发完成度。WMS与Sinops均存在“完了但仍有残课题”的记录。", after=6)
    add_table(
        doc,
        ["系统", "QA总数", "当前状态", "设计开发判断"],
        [
            ["WMS（Hitluster）", "63", "完了53；対応中4；中止6", "可扩大详细设计；整体开发需先关闭关键方针项"],
            ["自动补充（Sinops）", "80", "完了75；保留2；対応中1；空白1；中止1", "直接GCS接口可分批推进；Master/DWH及棚割需条件式推进"],
        ],
        [1.35, 0.8, 2.25, 2.1],
    )
    add_text(doc, "注：同一课题的追加提问按QA行记录；“中止”需在设计资料中明确替代方案或正式排除范围。", size=9.5, color=MUTED, italic=True, after=8)

    doc.add_heading("2. 课题一：WMS（Hitluster）", level=1)
    add_callout(doc, "WMS：53/63条QA已完了，4条仍対応中，6条中止；主要风险集中在商品Master物流属性、库存全量/差分处理、成本方针及外部系统责任边界。", fill="FFF6E5", bar="A67C00")
    doc.add_heading("2.1 已确认并可作为设计输入的内容", level=2)
    add_bullet(doc, "商品层级采用“部門＞大分類＞中分類＞小分類＞細分類”，来源为基本商品Master；商品样例及部分分类资料已提供。")
    add_bullet(doc, "店Master、取引先Master、商品分類Master的样例已取得；荷主代码按固定值0000处理，企业代码可固定值设定。")
    add_bullet(doc, "TC店別入荷也需要生成店舗振替数据；输入连番用于计划与实绩的匹配；西友无预托库存，横持属于中心间物流转送。")
    add_bullet(doc, "在庫一覧为日次全件数据，在庫調整为发生时发送的差分数据；负数调整在无既存记录时需新规登记。")
    add_bullet(doc, "仓库代码原则上保持西友六码；更新时间暂按数据接收时间；调整代码由基干侧用于差异原因识别，WMS侧不管理。")

    doc.add_heading("2.2 当前应优先关闭的WMS课题", level=2)
    add_table(
        doc,
        ["QA", "课题", "状态", "影响"],
        [
            ["No.188", "商品Master的发注单位数设置规则", "対応中", "影响商品Master Mapping及数量计算"],
            ["No.189", "商品Master的ケース入数设置规则", "対応中", "需明确Case/Ball及WMS维护范围"],
            ["No.196", "中心返品、廃棄的实际登记系统", "対応中", "影响WMS接口范围和业务流程"],
            ["No.198", "DC入荷予定的出荷单位数定义", "対応中", "TRIAL运用仍需整理并统一数量定义"],
        ],
        [0.8, 2.55, 0.85, 2.3],
    )
    add_text(doc, "No.188和No.189应在商品Master详细设计冻结前关闭；No.196和No.198需要由业务、MD基干、TC/WMS共同确认责任系统。", size=9.5, color=MUTED, after=8)

    doc.add_heading("2.3 完了但仍需跟踪的WMS残课题", level=2)
    add_bullet(doc, "物流控制器依赖：DC Location、Lot管理、托盘相关字段、商品分类层级以及TC总量入荷汇总转换。")
    add_bullet(doc, "Shinise依赖：便区分维护，以及部分商品物流属性的补充。")
    add_bullet(doc, "主数据依赖：出荷先代码、店铺/中心代码是否沿用现行体系，需与Master统合PJ方针一致。")
    add_bullet(doc, "成本与库存：移动平均单价由哪一侧计算、WMS需提供哪些基础字段、全量与差分的应用顺序、并发控制、重试及全量缺失回退。")
    add_bullet(doc, "ITF代码：一个商品代码对应多个ITF代码的同步方式仍需形成可执行规则。")

    doc.add_heading("3. 课题二：自动补充（Sinops）", level=1)
    add_callout(doc, "自动补充：QA完了75/80，但11本Master经“GCS→西友DWH→Sinops”转送，14本走“GCS→Sinops”；DWH追加链路及棚割路径仍是主要设计变量。")
    doc.add_heading("3.1 链路和范围前提", level=2)
    add_bullet(doc, "25本IF均先上传GCS；14本直接从GCS向Sinops传送。")
    add_bullet(doc, "11本Master在GCS之后进入西友DWH，再由DWH处理并转送Sinops；DWH接收、转换、调度、重送、监控和对账需单独设计。")
    add_bullet(doc, "11本Master按“一个业务IF＋一段DWH追加链路”估算，避免把GCS上传和DWH转送重复计为两个完整IF。")

    doc.add_heading("3.2 已确认并可优先推进的内容", level=2)
    add_bullet(doc, "销售实绩、来客数、发注勧告、发注实绩、休日设置及受払明细等接口的主要字段、时间点、全量/差分和数量换算规则已基本明确。")
    add_bullet(doc, "发注勧告的当日/翌日以发注日区分；发注数量按Sinops发注バラ数与发注单位换算，TRIAL侧使用固定换算规则。")
    add_bullet(doc, "入荷实绩和入荷预定的便区分、库存修正等事项已有方向，部分实施责任转交Shinise。")
    add_bullet(doc, "新商品库存数据明确需要从TRIAL基干生成，并建立自动发注对象Master；该事项虽为完了，仍应作为实施任务跟踪。")

    doc.add_heading("3.3 当前应优先关闭的自动补充课题", level=2)
    add_table(
        doc,
        ["QA", "课题", "状态", "影响"],
        [
            ["No.9", "商品Master统合及店内代码对应", "未回答／空白", "直接影响11本Master Mapping及DWH数据契约"],
            ["No.56/57", "仕入先Master与仓库关系、数据量及拆分方式", "保留", "影响DWH仓库拆分、性能及文件输出责任"],
            ["No.74", "棚割明细的最终连携路径", "対応中", "Sinops直连、商谈/棚割DB及DWH路径尚未统一"],
        ],
        [0.95, 2.35, 0.9, 2.3],
    )

    doc.add_heading("3.4 完了但仍需跟踪的自动补充残课题", level=2)
    add_bullet(doc, "Shinise依赖：入荷实绩/预定的便区分、库存修正处理。")
    add_bullet(doc, "Case/散装转换Master：由哪一侧维护、如何把店铺差异归并为商品级数据仍需形成规则。")
    add_bullet(doc, "棚割及新商品库存：DWH参照方式、棚割来源、TRIAL基干生成逻辑及自动发注对象Master建设需要明确负责人、期限和验收物。")
    add_bullet(doc, "供应商与仓库关系：可以从调达Master推导，但数据量较大，需确认是否由DWH或Sinops侧进行仓库分割。")

    doc.add_heading("4. 跨领域依赖", level=1)
    add_text(doc, "以下事项同时影响WMS、自动补充和BO，不能只在单个IF的Mapping表中处理：", after=5)
    add_bullet(doc, "区分管理Master、商品/店铺/供应商代码体系及长度限制。")
    add_bullet(doc, "GCS目录、授受文件、归档、重送、监控和对账规则。")
    add_bullet(doc, "物流控制器与Shinise的字段补充、便区分、DC Location、Case/Pack以及中心物流转换。")
    add_bullet(doc, "MD基干与业务侧对移动平均单价、仕入传票、税区分、返品/廃棄业务的最终方针。")
    add_callout(doc, "管理原则：QA回答状态与实施残课题必须分开管理；残课题至少补充负责人、期限、设计产物、依赖系统和验收条件。")

    doc.add_heading("5. 设计开发入口判断", level=1)
    add_table(
        doc,
        ["对象", "入口判断", "进入条件"],
        [
            ["WMS基础Master及已确认IF", "GO", "沿用已确认样例、代码和格式；残课题登记为后续任务"],
            ["WMS库存、成本及物流控制器相关IF", "条件式GO", "先冻结全量/差分、异常、成本及外部系统责任"],
            ["Sinops直接GCS接口（14本）", "分批GO", "字段、时间点和GCS共通框架确认"],
            ["Sinops Master DWH接口（11本）", "条件式GO", "No.9、No.56/57关闭或明确假设和返工边界"],
            ["棚割及新商品库存相关接口", "暂缓/条件式", "统一最终路径，明确DWH来源和自动发注对象Master方案"],
        ],
        [1.75, 1.15, 3.6],
    )

    doc.add_heading("6. 后续行动建议", level=1)
    add_number(doc, "关闭WMS No.188、No.189、No.196、No.198，并将结论回填商品Master、返品/廃棄及DC入荷予定的设计资料。")
    add_number(doc, "关闭自动补充No.9、No.56/57、No.74；若无法立即关闭，须书面固定设计假设、责任人、期限和变更时的返工范围。")
    add_number(doc, "先冻结共通数据契约：Key、字段来源、代码表、默认值、时间点、全量/差分、重送、幂等、对账和异常恢复。")
    add_number(doc, "将11本Master的DWH追加链路作为独立工作包管理，不把DWH工作隐藏在GCS上传或Sinops接口工数中。")
    add_number(doc, "中止QA补充中止原因及替代课题编号；完了但带残课题的记录转为可追踪实施任务。")
    add_number(doc, "结合测试、综合测试、UAT、移行演练、上线切换及稳定化支持不纳入本阶段估算，另行估算。")

    doc.add_heading("7. 参照资料", level=1)
    add_bullet(doc, "情報システムPMI_MD基幹統合PJ_TRE_課題管理表.xlsx（最新QA管理表，表内更新日2026-08-28）")
    add_bullet(doc, "PJWest_WMS_Sinops_QA_ReSummary_20260826_QA2.docx（本资料的结构和表现形式参考）")
    add_bullet(doc, "03. WMS (Hitluster)-20260825T065143Z-1-001.zip")
    add_bullet(doc, "01. 自动补充（Sinops）-20260825T065143Z-1-001.zip")
    add_bullet(doc, "PJWest_情報システム_MD基幹統合PJ _週次定例_20260826.pptx")

    doc.core_properties.title = "WMS・自动补充QA课题再整理总结"
    doc.core_properties.subject = "最新QA状态、残课题及设计开发判断"
    doc.core_properties.author = "PMI MD Integration Project"
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
