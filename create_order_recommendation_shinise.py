from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


OUT = Path(__file__).parent / "発注勧告_Shinise直接データ版"
OUT.mkdir(exist_ok=True)

BLUE = "1F4E78"
LIGHT_BLUE = "D9EAF7"
LIGHT_YELLOW = "FFF2CC"
LIGHT_RED = "FCE4D6"
LIGHT_GREEN = "E2F0D9"
GRAY = "E7E6E6"
WHITE = "FFFFFF"
THIN = Side(style="thin", color="808080")

SCHEMA = [
    ["id", "bigserial", "物理列", "DB自動採番"],
    ["data", "jsonb", "物理列", "発注勧告の業務項目をJSONで格納"],
    ["registered_by", "varchar(8)", "物理列", "既存サンプル値は継続使用不可。新規採番が必要"],
    ["registered_at", "timestamptz", "物理列", "サンプル：タイムゾーン付き登録日時"],
    ["action_type", "int2", "物理列", "サンプル値：0。コード定義は未受領"],
]

DAILY = [
    ["id", "DB列", "bigserial", "—", "—", "○", "DBで自動採番し、IFから設定しない", ""],
    ["data.seq", "JSON", "数値", "対応項目なし", "—", "△", "取込単位で一意となる連番を生成", "G02"],
    ["data.order_date", "JSON", "yyyy-mm-dd", "No.2 発注日", "yyyymmdd", "○",
     "yyyymmdd → yyyy-mm-dd", ""],
    ["data.order_type", "JSON", "数値", "日配判定", "—", "○",
     "日配対象は1を設定", "G01"],
    ["data.store_code", "JSON", "数値", "No.1 店舗コード", "文字・実態6桁", "△",
     "数値化して設定。先頭ゼロは除去される", "G08"],
    ["data.product_code", "JSON", "文字", "No.3 商品コード", "文字・実態13桁", "○",
     "文字列のまま設定し、先頭ゼロを保持", "G08"],
    ["data.delivery_date", "JSON", "yyyy-mm-dd", "No.4 納品日", "yyyymmdd", "○",
     "日配対象のみ送信。yyyymmdd → yyyy-mm-dd", ""],
    ["data.order_quantity", "JSON", "整数", "No.11 発注バラ数", "整数8桁", "○",
     "発注単位調整後のバラ総数を直接設定。No.11＝No.12×No.15。×1000しない", "G07"],
    ["data.order_closing_number", "JSON", "数値", "対応項目なし", "—", "×",
     "発注締番号。No.5 便区分は使用しない。発注締めマスタまたは処理回から取得", "G11"],
    ["registered_by", "DB列", "varchar(8)", "対応項目なし", "—", "△",
     "既存サンプル値は使用しない。Shinise側で新規採番した8桁コードを設定", "G04"],
    ["registered_at", "DB列", "timestamptz", "対応項目なし", "—", "○",
     "DB登録時刻をタイムゾーン付きで設定", "G06"],
    ["action_type", "DB列", "int2", "対応項目なし", "—", "△",
     "サンプルは0。追加・更新・削除のコード定義を確認", "G03"],
]

NON_DAILY = [
    ["id", "DB列", "bigserial", "—", "—", "○", "DBで自動採番し、IFから設定しない", ""],
    ["data.seq", "JSON", "数値", "対応項目なし", "—", "△", "取込単位で一意となる連番を生成", "G02"],
    ["data.order_date", "JSON", "yyyy-mm-dd", "No.2 発注日", "yyyymmdd", "○",
     "yyyymmdd → yyyy-mm-dd", ""],
    ["data.order_type", "JSON", "数値", "非日配判定", "—", "○",
     "非日配対象は0を設定", "G01"],
    ["data.store_code", "JSON", "数値", "No.1 店舗コード", "文字・実態6桁", "△",
     "数値化して設定。先頭ゼロは除去される", "G08"],
    ["data.product_code", "JSON", "文字", "No.3 商品コード", "文字・実態13桁", "○",
     "文字列のまま設定し、先頭ゼロを保持", "G08"],
    ["data.delivery_date", "JSON", "送信しない", "No.4 納品日", "yyyymmdd", "—",
     "非日配対象ではJSONキー自体を出力しない。nullや空文字も送らない", ""],
    ["data.order_quantity", "JSON", "整数", "No.11 発注バラ数", "整数8桁", "○",
     "発注単位調整後のバラ総数を直接設定。No.11＝No.12×No.15。×1000しない", "G07"],
    ["data.order_closing_number", "JSON", "数値", "対応項目なし", "—", "×",
     "発注締番号。No.5 便区分は使用しない。発注締めマスタまたは処理回から取得", "G11"],
    ["registered_by", "DB列", "varchar(8)", "対応項目なし", "—", "△",
     "既存サンプル値は使用しない。Shinise側で新規採番した8桁コードを設定", "G04"],
    ["registered_at", "DB列", "timestamptz", "対応項目なし", "—", "○",
     "DB登録時刻をタイムゾーン付きで設定", "G06"],
    ["action_type", "DB列", "int2", "対応項目なし", "—", "△",
     "サンプルは0。追加・更新・削除のコード定義を確認", "G03"],
]

WAREHOUSE_DAILY = [
    ["id", "DB列", "bigserial", "—", "—", "○", "DBで自動採番し、IFから設定しない", ""],
    ["data.seq", "JSON", "数値", "対応項目なし", "—", "△", "取込単位で一意となる連番を生成", "G02"],
    ["data.order_date", "JSON", "yyyy-mm-dd", "No.1 発行日", "yyyymmdd", "○",
     "yyyymmdd → yyyy-mm-dd", ""],
    ["data.order_type", "JSON", "数値", "日配判定", "—", "○",
     "日配対象は1を設定", "G01"],
    ["data.store_code", "JSON", "数値", "ファイル名の倉庫コード", "実態6桁", "○",
     "倉庫コードを数値化して設定。例：007452 → 7452", "G08 G09"],
    ["data.product_code", "JSON", "文字", "No.7 品番", "文字・実態13桁", "○",
     "文字列のまま設定し、先頭ゼロを保持", "G08"],
    ["data.delivery_date", "JSON", "yyyy-mm-dd", "No.14 納品日", "yyyymmdd", "○",
     "日配対象のみ送信。yyyymmdd → yyyy-mm-dd", ""],
    ["data.order_quantity", "JSON", "整数", "No.11 バラ換算数量", "整数8桁", "○",
     "発注単位調整後のバラ総数を直接設定。No.11＝No.9×No.12。×1000しない", "G07"],
    ["data.order_closing_number", "JSON", "数値", "対応項目なし", "—", "×",
     "発注締番号。発注締めマスタまたは処理回から取得", "G11"],
    ["registered_by", "DB列", "varchar(8)", "対応項目なし", "—", "△",
     "既存サンプル値は使用しない。Shinise側で新規採番した8桁コードを設定", "G04"],
    ["registered_at", "DB列", "timestamptz", "対応項目なし", "—", "○",
     "DB登録時刻をタイムゾーン付きで設定", "G06"],
    ["action_type", "DB列", "int2", "対応項目なし", "—", "△",
     "サンプルは0。追加・更新・削除のコード定義を確認", "G03"],
]

WAREHOUSE_NON_DAILY = [
    ["id", "DB列", "bigserial", "—", "—", "○", "DBで自動採番し、IFから設定しない", ""],
    ["data.seq", "JSON", "数値", "対応項目なし", "—", "△", "取込単位で一意となる連番を生成", "G02"],
    ["data.order_date", "JSON", "yyyy-mm-dd", "No.1 発行日", "yyyymmdd", "○",
     "yyyymmdd → yyyy-mm-dd", ""],
    ["data.order_type", "JSON", "数値", "非日配判定", "—", "○",
     "非日配対象は0を設定", "G01"],
    ["data.store_code", "JSON", "数値", "ファイル名の倉庫コード", "実態6桁", "○",
     "倉庫コードを数値化して設定。例：007452 → 7452", "G08 G09"],
    ["data.product_code", "JSON", "文字", "No.7 品番", "文字・実態13桁", "○",
     "文字列のまま設定し、先頭ゼロを保持", "G08"],
    ["data.delivery_date", "JSON", "送信しない", "No.14 納品日", "yyyymmdd", "—",
     "非日配対象ではJSONキー自体を出力しない。nullや空文字も送らない", ""],
    ["data.order_quantity", "JSON", "整数", "No.11 バラ換算数量", "整数8桁", "○",
     "発注単位調整後のバラ総数を直接設定。No.11＝No.9×No.12。×1000しない", "G07"],
    ["data.order_closing_number", "JSON", "数値", "対応項目なし", "—", "×",
     "発注締番号。発注締めマスタまたは処理回から取得", "G11"],
    ["registered_by", "DB列", "varchar(8)", "対応項目なし", "—", "△",
     "既存サンプル値は使用しない。Shinise側で新規採番した8桁コードを設定", "G04"],
    ["registered_at", "DB列", "timestamptz", "対応項目なし", "—", "○",
     "DB登録時刻をタイムゾーン付きで設定", "G06"],
    ["action_type", "DB列", "int2", "対応項目なし", "—", "△",
     "サンプルは0。追加・更新・削除のコード定義を確認", "G03"],
]

GAPS = [
    ["G01", "日配／非日配判定元", "★最高",
     "サンプルではorder_type=1にdelivery_dateがあり、order_type=0にはない",
     "どのマスタ・項目で日配対象を判定するかを確定。日配=1、非日配=0としてよいか確認", "未確認"],
    ["G02", "seq採番規則", "★高",
     "sinopsファイルにseqはなく、サンプルは数値を保持",
     "全体一意、ファイル内連番、取込バッチ単位のいずれかを確認", "未確認"],
    ["G03", "action_typeコード", "★最高",
     "物理型int2、サンプルは全件0だが意味の定義がない",
     "追加・更新・削除のコード一覧と、再送時の設定値を確認", "未確認"],
    ["G04", "registered_by新規採番", "★最高",
     "varchar(8)。既存サンプル値は継続使用できない",
     "Shinise側で本IF専用コードを新規採番し、8桁の採番値・利用環境・管理責任者を確定", "採番待ち"],
    ["G05", "重複・再送制御", "★最高",
     "物理PKはbigserialのidのみ。同じ業務データを再送すると重複登録される可能性がある",
     "業務キー、UPSERT条件、再送時の冪等性を確定", "未確認"],
    ["G06", "登録日時・タイムゾーン", "★中",
     "registered_atはtimestamptz、サンプルは+0800",
     "本番の基準タイムゾーン、DB自動設定かIF明示設定かを確認", "未確認"],
    ["G07", "推奨発注数量", "—",
     "発注単位調整後のバラ総数で、発注単位の整数倍",
     "店舗系はNo.11 発注バラ数、倉庫系はNo.11 バラ換算数量を直接設定。×1000しない", "確認済"],
    ["G08", "店舗・商品コード型", "★中",
     "store_codeは数値、product_codeは文字。sinops店舗コードはゼロ埋め文字",
     "店舗コードは数値化、商品コードは文字列保持。コード同一性を確認", "未確認"],
    ["G09", "倉庫系store_code", "—",
     "ka/subkaは行内に店舗コードを持たない",
     "ファイル名の倉庫コードをstore_codeに設定する", "確認済"],
    ["G10", "0数量の扱い", "★高",
     "サンプルにorder_quantity=0が存在する",
     "0を有効な推奨数更新として登録するか、取消・削除扱いかを確認", "未確認"],
    ["G11", "発注締番号の取得元", "★最高",
     "order_closing_numberは発注締番号。sinops発注勧告ファイルに直接対応項目はなく、No.5 便区分とは別項目",
     "発注締めマスタ、処理時刻、処理回等の正式な取得元と採番規則を確認", "未確認"],
]

RULES = [
    ["R01", "分岐", "日配対象", "order_type=1", "`delivery_date`をJSONに含める", "G01"],
    ["R02", "分岐", "非日配対象", "order_type=0", "`delivery_date`キーを出力しない。nullも送らない", "確定"],
    ["R03", "日付", "order_date", "店舗No.2 発注日／倉庫No.1 発行日", "yyyymmdd → yyyy-mm-dd", "確定"],
    ["R04", "日付", "delivery_date", "店舗No.4／倉庫No.14 納品日",
     "日配時のみyyyymmdd → yyyy-mm-dd", "確定"],
    ["R05", "数量", "order_quantity", "店舗No.11 発注バラ数／倉庫No.11 バラ換算数量",
     "発注単位調整後のバラ総数を直接設定。×1000しない", "確定"],
    ["R06", "コード", "store_code", "店舗No.1／ファイル名の倉庫コード",
     "文字列を数値化。倉庫系は倉庫コードを設定", "確定（G08 G09）"],
    ["R07", "コード", "product_code", "店舗No.3 商品コード／倉庫No.7 品番",
     "文字列のまま設定し先頭ゼロを保持", "確定"],
    ["R08", "コード", "order_closing_number", "対応項目なし",
     "発注締番号を正式なマスタまたは処理回から取得。No.5 便区分は使用しない", "G11"],
    ["R09", "採番", "seq", "対応元なし", "確定した採番単位で生成", "G02"],
    ["R10", "監査", "registered_by", "対応元なし",
     "既存サンプル値は使用せず、Shinise側で新規採番した8桁コードを設定", "G04"],
    ["R11", "監査", "registered_at", "対応元なし", "DB登録時刻をtimestamptzで設定", "G06"],
    ["R12", "操作", "action_type", "対応元なし", "コード定義に従って設定。サンプル値0を無条件使用しない", "G03"],
    ["R13", "採番", "id", "対応元なし", "DBのbigserialで自動採番", "確定"],
]


def base_style(ws, widths, verdict_col=None, priority_col=None, status_col=None):
    ws.freeze_panes = "A5"
    ws.auto_filter.ref = ws.dimensions
    for index, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(index)].width = width
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
    for cell in ws[1]:
        cell.fill = PatternFill("solid", fgColor=BLUE)
        cell.font = Font(color=WHITE, bold=True, size=14)
    for cell in ws[4]:
        cell.fill = PatternFill("solid", fgColor=LIGHT_BLUE)
        cell.font = Font(bold=True)
    for row_num in range(5, ws.max_row + 1):
        fill = WHITE
        if verdict_col:
            verdict = str(ws.cell(row_num, verdict_col).value or "")
            fill = (
                LIGHT_GREEN if verdict == "○"
                else LIGHT_YELLOW if verdict == "△"
                else LIGHT_RED if verdict == "×"
                else GRAY
            )
        if priority_col:
            priority = str(ws.cell(row_num, priority_col).value or "")
            fill = LIGHT_RED if priority == "★最高" else LIGHT_YELLOW if priority == "★高" else GRAY
        if status_col:
            status = str(ws.cell(row_num, status_col).value or "")
            fill = LIGHT_GREEN if status.startswith("確定") else LIGHT_YELLOW
        for cell in ws[row_num]:
            cell.fill = PatternFill("solid", fgColor=fill)


def add_intro(ws, title, subtitle, note, columns):
    ws.append([title])
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=columns)
    ws.append([subtitle])
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=columns)
    ws.append([note])
    ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=columns)


def build_mapping():
    wb = Workbook()
    schema = wb.active
    schema.title = "Shinise物理構造"
    add_intro(schema, "Shinise auto_orders 物理構造", "テーブル：auto_orders",
              "出典：ユーザー提供の構造図・サンプルデータ（2026-09-04）", 4)
    schema.append(["列名", "型", "区分", "説明"])
    for row in SCHEMA:
        schema.append(row)
    base_style(schema, [24, 20, 14, 55])

    headers = ["対象項目", "格納先", "型/形式", "sinops対応元", "sinops型/形式", "判定", "変換ルール", "関連GAP"]
    for title, rows, source, note in [
        ("店舗系_日配_type1", DAILY, "kankoku.txt / subkankoku.txt",
         "日配対象：delivery_dateを送信"),
        ("店舗系_非日配_type0", NON_DAILY, "kankoku.txt / subkankoku.txt",
         "非日配対象：delivery_dateキーを送信しない"),
        ("倉庫系_日配_type1", WAREHOUSE_DAILY, "【倉庫CD】ka.txt / subka.txt",
         "store_code＝ファイル名の倉庫コード／日配はdelivery_dateを送信"),
        ("倉庫系_非日配_type0", WAREHOUSE_NON_DAILY, "【倉庫CD】ka.txt / subka.txt",
         "store_code＝ファイル名の倉庫コード／非日配はdelivery_dateキーを送信しない"),
    ]:
        ws = wb.create_sheet(title)
        add_intro(ws, f"発注勧告 項目マッピング — {title}", f"対象：{source}",
                  note, len(headers))
        ws.append(headers)
        for row in rows:
            ws.append(row)
        base_style(ws, [30, 13, 18, 28, 20, 9, 60, 13], verdict_col=6)
    wb.save(OUT / "発注勧告_①_項目マッピング表.xlsx")


def build_gap():
    wb = Workbook()
    ws = wb.active
    ws.title = "GAP分析一覧"
    headers = ["No.", "GAPカテゴリ", "優先度", "現状・差異", "確認・対応方針", "ステータス"]
    add_intro(ws, "GAP分析書 — 発注勧告IF（Shinise auto_orders直接データ版）",
              "対象：kankoku.txt / subkankoku.txt / 【倉庫CD】ka.txt / subka.txt",
              "日配はdelivery_dateあり、非日配は同キーを送信しない。倉庫系store_codeは倉庫コード。", len(headers))
    ws.append(headers)
    for row in GAPS:
        ws.append(row)
    base_style(ws, [9, 27, 12, 62, 65, 15], priority_col=3)
    wb.save(OUT / "発注勧告_②_GAP分析書.xlsx")


def build_rules():
    wb = Workbook()
    ws = wb.active
    ws.title = "変更ルール"
    headers = ["Rule ID", "分類", "対象", "入力元/条件", "設定ルール", "状態/関連GAP"]
    add_intro(ws, "変更ルール定義書 — 発注勧告IF（Shinise auto_orders直接データ版）",
              "対象：店舗系kankoku/subkankoku＋倉庫系ka/subka",
              "JSONキーの有無を含め、日配／非日配を分岐する。", len(headers))
    ws.append(headers)
    for row in RULES:
        ws.append(row)
    base_style(ws, [12, 13, 28, 30, 68, 18], status_col=6)
    wb.save(OUT / "発注勧告_③_変更ルール定義書.xlsx")


def verify():
    files = [
        "発注勧告_①_項目マッピング表.xlsx",
        "発注勧告_②_GAP分析書.xlsx",
        "発注勧告_③_変更ルール定義書.xlsx",
    ]
    for name in files:
        wb = load_workbook(OUT / name, data_only=True, read_only=True)
        assert all(wb[sheet].max_row >= 9 for sheet in wb.sheetnames)
    wb = load_workbook(OUT / files[0], data_only=True, read_only=True)
    assert len(wb.sheetnames) == 5
    daily_rows = list(wb["店舗系_日配_type1"].iter_rows(values_only=True))
    non_daily_rows = list(wb["店舗系_非日配_type0"].iter_rows(values_only=True))
    warehouse_daily = list(wb["倉庫系_日配_type1"].iter_rows(values_only=True))
    warehouse_non_daily = list(wb["倉庫系_非日配_type0"].iter_rows(values_only=True))
    assert next(row for row in daily_rows if row[0] == "data.delivery_date")[6].startswith("日配対象のみ送信")
    assert "JSONキー自体を出力しない" in next(
        row for row in non_daily_rows if row[0] == "data.delivery_date"
    )[6]
    assert next(row for row in daily_rows if row[0] == "data.order_quantity")[3] == "No.11 発注バラ数"
    closing = next(row for row in daily_rows if row[0] == "data.order_closing_number")
    assert closing[3] == "対応項目なし"
    assert "No.5 便区分は使用しない" in closing[6]
    registered_by = next(row for row in daily_rows if row[0] == "registered_by")
    assert "新規採番" in registered_by[6]
    assert "90000001" not in registered_by[6]
    assert next(row for row in warehouse_daily if row[0] == "data.store_code")[3] == "ファイル名の倉庫コード"
    assert next(row for row in warehouse_daily if row[0] == "data.order_quantity")[3] == "No.11 バラ換算数量"
    assert "JSONキー自体を出力しない" in next(
        row for row in warehouse_non_daily if row[0] == "data.delivery_date"
    )[6]


if __name__ == "__main__":
    build_mapping()
    build_gap()
    build_rules()
    verify()
    print(f"created: {OUT}")
