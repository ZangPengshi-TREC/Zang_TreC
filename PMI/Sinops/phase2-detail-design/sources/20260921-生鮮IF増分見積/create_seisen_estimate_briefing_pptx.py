#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""生鮮IF増分・標準案見積の説明PPT（リスク着眼点を厚めに）。"""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt, Emu

OUT = Path(__file__).resolve().parent / "【説明】生鮮IF増分_標準案見積_リスク着眼_20260921.pptx"

# Palette
NAVY = RGBColor(0x1F, 0x4E, 0x78)
BLUE = RGBColor(0x2E, 0x75, 0xB6)
LIGHT = RGBColor(0xD6, 0xE3, 0xF0)
ORANGE = RGBColor(0xC6, 0x59, 0x11)
ORANGE_BG = RGBColor(0xFC, 0xE4, 0xD6)
GREEN = RGBColor(0x54, 0x89, 0x35)
GREEN_BG = RGBColor(0xE2, 0xF0, 0xD9)
GRAY = RGBColor(0x59, 0x59, 0x59)
DARK = RGBColor(0x2F, 0x2F, 0x2F)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
RED_SOFT = RGBColor(0xC0, 0x00, 0x00)


def set_run(run, size=14, bold=False, color=DARK, name="Yu Gothic"):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = name


def add_text(tf, text, size=14, bold=False, color=DARK, align=PP_ALIGN.LEFT, clear=True):
    if clear:
        tf.clear()
        p = tf.paragraphs[0]
        p.text = ""
    else:
        p = tf.add_paragraph()
    p.alignment = align
    run = p.add_run()
    run.text = text
    set_run(run, size=size, bold=bold, color=color)
    return p


def add_para(tf, text, size=13, bold=False, color=DARK, space_before=4, space_after=2):
    p = tf.add_paragraph()
    p.alignment = PP_ALIGN.LEFT
    p.space_before = Pt(space_before)
    p.space_after = Pt(space_after)
    run = p.add_run()
    run.text = text
    set_run(run, size=size, bold=bold, color=color)
    return p


def blank_slide(prs):
    layout = prs.slide_layouts[6]  # blank
    return prs.slides.add_slide(layout)


def banner(slide, title, subtitle=None):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.85)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = NAVY
    shape.line.fill.background()
    tf = shape.text_frame
    tf.word_wrap = True
    add_text(tf, title, size=22, bold=True, color=WHITE)
    tf.paragraphs[0].alignment = PP_ALIGN.LEFT
    shape.text_frame.margin_left = Inches(0.35)
    shape.text_frame.margin_top = Inches(0.18)
    if subtitle:
        bar = slide.shapes.add_textbox(Inches(0.35), Inches(0.92), Inches(12.5), Inches(0.35))
        add_text(bar.text_frame, subtitle, size=12, color=GRAY)


def footer(slide, page, total=11):
    box = slide.shapes.add_textbox(Inches(0.35), Inches(7.15), Inches(10), Inches(0.3))
    add_text(
        box.text_frame,
        "西友MD基幹統合｜生鮮IF増分見積（標準案）説明｜2026-09-21 Rev.e",
        size=10,
        color=GRAY,
    )
    num = slide.shapes.add_textbox(Inches(12.2), Inches(7.15), Inches(0.9), Inches(0.3))
    add_text(num.text_frame, f"{page}/{total}", size=10, color=GRAY, align=PP_ALIGN.RIGHT)


def card(slide, left, top, width, height, fill=LIGHT):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    sh.line.color.rgb = RGBColor(0xB0, 0xC4, 0xDE)
    sh.adjustments[0] = 0.08
    return sh


def table_slide_data(slide, left, top, rows, col_widths, font_size=11):
    """rows: list of list of str; first row = header."""
    nrows = len(rows)
    ncols = len(rows[0])
    table_shape = slide.shapes.add_table(
        nrows, ncols, left, top, sum(col_widths), Inches(0.32 * nrows)
    )
    table = table_shape.table
    for i, w in enumerate(col_widths):
        table.columns[i].width = w
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            cell = table.cell(r, c)
            cell.text = str(val)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            for p in cell.text_frame.paragraphs:
                p.alignment = PP_ALIGN.LEFT if c > 0 or r == 0 else PP_ALIGN.CENTER
                for run in p.runs:
                    set_run(
                        run,
                        size=font_size,
                        bold=(r == 0),
                        color=WHITE if r == 0 else DARK,
                    )
            if r == 0:
                cell.fill.solid()
                cell.fill.fore_color.rgb = NAVY
            elif r == nrows - 1 and ("合計" in str(row[0]) or "計" == str(row[0])[-1:]):
                cell.fill.solid()
                cell.fill.fore_color.rgb = GREEN_BG
            elif r % 2 == 0:
                cell.fill.solid()
                cell.fill.fore_color.rgb = RGBColor(0xF5, 0xF8, 0xFB)
    return table


def build():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    total = 11

    # ---- 1 Cover ----
    s = blank_slide(prs)
    bg = s.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5)
    )
    bg.fill.solid()
    bg.fill.fore_color.rgb = NAVY
    bg.line.fill.background()
    accent = s.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0), Inches(5.9), Inches(13.333), Inches(1.6)
    )
    accent.fill.solid()
    accent.fill.fore_color.rgb = BLUE
    accent.line.fill.background()

    t = s.shapes.add_textbox(Inches(0.7), Inches(2.0), Inches(11.5), Inches(1.2))
    add_text(t.text_frame, "生鮮 IF 増分見積　説明資料", size=36, bold=True, color=WHITE)
    t2 = s.shapes.add_textbox(Inches(0.7), Inches(3.3), Inches(11.5), Inches(0.6))
    add_text(
        t2.text_frame,
        "標準案｜調査・Mapping → 詳細設計・開発・単体（＋PJ按分）",
        size=18,
        color=LIGHT,
    )
    t3 = s.shapes.add_textbox(Inches(0.7), Inches(4.1), Inches(11.5), Inches(0.5))
    add_text(
        t3.text_frame,
        "焦点：生鮮の不確実性を見積前提として明示し、提示値の読み方を揃える",
        size=16,
        bold=True,
        color=ORANGE_BG,
    )
    t4 = s.shapes.add_textbox(Inches(0.7), Inches(6.25), Inches(11.5), Inches(0.8))
    tf = t4.text_frame
    add_text(tf, "西友 MD基幹統合　／　Sinops 自動補充 PMI", size=14, color=WHITE)
    add_para(tf, "資料日 2026-09-21（Rev.e）　出典：標準案 Excel・20260918 生鮮IF会議", size=12, color=LIGHT)

    # ---- 2 結論 ----
    s = blank_slide(prs)
    banner(s, "1. 結論（本見積の提示値）", "PMI課題 Sinops非生鮮のPhase2は契約据え置き。本見積は生鮮増分のみ。")
    cards = [
        ("人日合計", "221.5 人日", "直接201.5（P1 59.5＋P2 142）＋PJ20"),
        ("金額（税抜）", "¥7,085,000", "Phase1×40,000＋Phase2×27,500＋PJ×40,000"),
        ("本見積の前提", "マスタ流用が\n概ね成立", "既存波及が限定的な場合\nの工数として算定"),
    ]
    for i, (h, v, n) in enumerate(cards):
        c = card(s, Inches(0.4 + i * 4.2), Inches(1.5), Inches(4.0), Inches(2.0), LIGHT)
        tf = c.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.2)
        tf.margin_top = Inches(0.2)
        add_text(tf, h, size=13, color=BLUE, bold=True)
        add_para(tf, v, size=26, bold=True, color=NAVY, space_before=8)
        add_para(tf, n, size=12, color=GRAY, space_before=6)

    box = card(s, Inches(0.4), Inches(3.8), Inches(12.5), Inches(2.9), ORANGE_BG)
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.25)
    tf.margin_top = Inches(0.15)
    add_text(tf, "説明時に最初に伝えること", size=15, bold=True, color=ORANGE)
    for line in [
        "・提示値は確定請負ではない。生鮮は不確実性が大きく、前提が崩れたら再見積。",
        "・最大の不確実性は生鮮専用マスタの流用可否（B1→B2仮置き）。",
        "・B2（マスタIF本体51人日）は仮置き。B1判定後に確定値へ置換する。",
        "・PMI課題 Sinops非生鮮のPhase2の契約額は据え置き。",
    ]:
        add_para(tf, line, size=14, color=DARK, space_before=6)
    footer(s, 2, total)

    # ---- 3 見積前提（不確実性）----
    s = blank_slide(prs)
    banner(
        s,
        "2. 見積前提条件（生鮮の不確実性・必読）",
        "提示値は「現時点の前提が成り立つ場合」の工数。確定請負額ではない。",
    )
    rows = [
        ["区分", "前提として明示すること", "見積上の扱い"],
        ["マスタ", "専用マスタ多数。非生鮮変換の流用可否・新規本数は未確定", "B1で1本ずつ三択判定。B2は仮置き"],
        ["横断影響", "マスタ前提変更は PMI課題 Sinops非生鮮のPhase2 にも波及し得る", "B3に波及バッファ。非生鮮Phase2契約額は据え置き"],
        ["源・GCS", "生鮮側サーバー→GCS。GCS無ければアップロードが必要（無い見込み）", "確認・アップロードを本見積に含む"],
        ["入荷", "実績＝2系統ソース、予定＝生鮮基幹新規。非生鮮11/12の単純流用不可", "Aで新規IF計上"],
        ["勧告の連携先", "生鮮発注統合直／自動発注経由は未確定", "連携先変更時は再見積"],
        ["棚割・在庫", "SM vs 新棚割未確定。営業在庫の生鮮適用は未検証", "方針確定・検証結果で再見積"],
        ["本案前提", "マスタ流用が概ね成立し、既存波及が限定的", "この条件が崩れたら提示値は無効→再見積"],
    ]
    table_slide_data(
        s,
        Inches(0.25),
        Inches(1.3),
        rows,
        [Inches(1.5), Inches(6.2), Inches(5.2)],
        font_size=11,
    )
    note = s.shapes.add_textbox(Inches(0.35), Inches(6.45), Inches(12.5), Inches(0.5))
    add_text(
        note.text_frame,
        "Excel「見積前提条件・注記」シートに全文あり。説明時は本頁を数字より先に置く。",
        size=13,
        bold=True,
        color=ORANGE,
    )
    footer(s, 3, total)

    # ---- 4 構成 ----
    s = blank_slide(prs)
    banner(s, "3. 見積のブロック構成", "A（先行トラン）は比較的確定。B（マスタ横断）が不確実性の本体。")
    rows = [
        ["#", "ブロック", "人日", "金額（円）", "リスク感"],
        ["A", "先行トラン（新規IF 4本）", "100.5", "3,070,000", "中（P1+P2単価混在）"],
        ["B1", "マスタ棚卸・判定（必須）", "35", "1,400,000", "高（Phase1単価40,000）"],
        ["B2", "マスタIF本体（仮置き）", "51", "1,402,500", "最高（Phase2直接27,500）"],
        ["B3", "非生鮮Phase2（PMI課題）への波及", "15", "412,500", "高（Phase2直接27,500）"],
        ["—", "直接小計", "201.5", "6,285,000", ""],
        ["C", "PJ管理（按分）", "20", "800,000", "既存比率 50/499.5"],
        ["合計", "本見積", "221.5", "7,085,000", ""],
    ]
    table_slide_data(
        s,
        Inches(0.4),
        Inches(1.4),
        rows,
        [Inches(0.7), Inches(3.6), Inches(1.3), Inches(1.8), Inches(5.0)],
        font_size=12,
    )
    note = s.shapes.add_textbox(Inches(0.4), Inches(6.35), Inches(12.5), Inches(0.6))
    add_text(
        note.text_frame,
        "オレンジ相当の B1〜B3 が「見積リスクを数値化した場所」。A だけ見せるとリスクを過小評価する。",
        size=13,
        bold=True,
        color=ORANGE,
    )
    footer(s, 4, total)

    # ---- 5 A明細 ----
    s = blank_slide(prs)
    banner(s, "4. A｜先行トラン4本（比較的見えている範囲）", "会議結論：非生鮮24/25・11/12の流用不可 → 新規IF必須")
    rows = [
        ["IF", "仮No.", "Phase1", "Layer①", "Layer③", "計", "なぜ新規か"],
        ["発注勧告（当日）", "S24", "6.0", "17.0", "4.5", "27.5", "項目・ファイルが非生鮮と違う／連携先＝生鮮発注統合"],
        ["発注勧告（翌日以降）", "S25", "4.5", "14.0", "3.0", "21.5", "未来スロット差分。当日と類似だが別行"],
        ["入荷実績", "S11", "8.0", "17.0", "4.0", "29.0", "2系統（PC/CK＝生鮮基幹、取引先＝MD基幹）"],
        ["入荷予定", "S12", "6.0", "14.0", "2.5", "22.5", "生鮮基幹が新ソース。Layer①別作"],
        ["A小計", "", "24.5", "62.0", "14.0", "100.5", "Layer②＝0（全件）"],
    ]
    table_slide_data(
        s,
        Inches(0.3),
        Inches(1.35),
        rows,
        [Inches(2.4), Inches(1.0), Inches(0.9), Inches(0.9), Inches(0.9), Inches(1.0), Inches(5.5)],
        font_size=11,
    )
    c = card(s, Inches(0.4), Inches(5.5), Inches(12.5), Inches(1.35), LIGHT)
    tf = c.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.2)
    add_text(tf, "A の見積リスク（本案で上振れを大きく載せない理由）", size=13, bold=True, color=NAVY)
    add_para(
        tf,
        "要件の骨格は 9/18 会議で先行確定。残リスクは項目精査・2系統合流設計・連携先経路の細部。"
        "本数確定前の大幅上振れは、B1判定後の再見積で扱う。",
        size=13,
        color=DARK,
        space_before=4,
    )
    footer(s, 5, total)

    # ---- 6 リスク核心 ----
    s = blank_slide(prs)
    banner(s, "5. 最重要リスク｜生鮮専用マスタ（見積の心臓部）", "ここを説明しないと、提示値は「根拠のない楽観」に見える")
    left = card(s, Inches(0.35), Inches(1.35), Inches(6.2), Inches(5.4), ORANGE_BG)
    tf = left.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.22)
    tf.margin_top = Inches(0.15)
    add_text(tf, "何が未確定か", size=16, bold=True, color=ORANGE)
    for line in [
        "・生鮮は専用マスタを大量に持つ",
        "・非生鮮の③変換ルールで足りるか未確認",
        "・流用／差分／新規の本数が未確定",
        "・影響は先行4本に閉じず Sinops 全体",
        "・源：生鮮側サーバー→GCS（無い見込み）",
        "・候補例：TM商品、店別差分、カテゴリ、",
        "　仕入先、停止フラグ、発注スケジュール、",
        "　棚割14／新商品在庫15 など",
    ]:
        add_para(tf, line, size=14, color=DARK, space_before=8)

    right = card(s, Inches(6.8), Inches(1.35), Inches(6.1), Inches(5.4), LIGHT)
    tf = right.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.22)
    tf.margin_top = Inches(0.15)
    add_text(tf, "見積上どう織り込んだか", size=16, bold=True, color=NAVY)
    for line in [
        "B1 35人日＝判定そのもの（必須・並行着手）",
        "　一覧化／1本ずつ三択判定／既存影響／",
        "　GCSアップロード確認／ゲート資料・再見積更新",
        "",
        "B2 51人日＝仮置き（3本×17人日想定）",
        "　→ B1後に確定値へ置換",
        "",
        "B3 15人日＝非生鮮Phase2（PMI課題）への軽微波及バッファ",
        "",
        "※ B を A と並行必須（後追いすると手戻り）",
    ]:
        add_para(tf, line if line else " ", size=14, color=DARK, space_before=5)
    footer(s, 6, total)

    # ---- 7 波及 ----
    s = blank_slide(prs)
    banner(s, "6. 非生鮮Phase2への波及リスク（B3）", "PMI課題 Sinops非生鮮のPhase2。契約額は据え置き。手戻り工数は増分側に計上")
    rows = [
        ["既存IF", "なぜ効くか", "顕在化したときの姿"],
        ["01/02 商品・店別", "TM商品・停止・発注期間が生鮮専用の可能性", "変換流用不可 → 別IF or 分岐"],
        ["03/04 カテゴリ・仕入先", "生鮮階層・仕入先体系が別", "コード体系の再マッピング"],
        ["05 ケースバラ", "計量・パック差", "パターン追加"],
        ["13 発注スケジュール", "締め・便・曜日が別源", "Layer①ソース分岐"],
        ["14/15 棚割・新商品在庫", "店舗調達を使わない可能性", "15個別対応／14ダミー可否"],
        ["06–10 実績・在庫", "営業在庫の生鮮適用未検証", "計算前提の再確認"],
    ]
    table_slide_data(
        s,
        Inches(0.3),
        Inches(1.35),
        rows,
        [Inches(2.8), Inches(5.0), Inches(5.0)],
        font_size=12,
    )
    footer(s, 7, total)

    # ---- 8 ゲート ----
    s = blank_slide(prs)
    banner(s, "7. 判定ゲートと再見積トリガー（リスク管理の運用）", "数字を固定契約にしすぎないための説明スライド")
    steps = [
        ("①", "B1 着手", "A の Phase1 と並行\nマスタ一覧・三択判定"),
        ("②", "ゲート", "流用／差分／新規の\n本数を確定"),
        ("③", "B2 置換", "仮置き51を\n確定値へ書き換え"),
        ("④", "手戻り", "影響は B3 で吸収\n超過なら再見積"),
    ]
    for i, (n, t, d) in enumerate(steps):
        c = card(s, Inches(0.35 + i * 3.25), Inches(1.5), Inches(3.05), Inches(2.4), LIGHT)
        tf = c.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.15)
        tf.margin_top = Inches(0.15)
        add_text(tf, f"{n} {t}", size=16, bold=True, color=NAVY)
        add_para(tf, d, size=13, color=DARK, space_before=10)

    trig = card(s, Inches(0.35), Inches(4.2), Inches(12.6), Inches(2.5), ORANGE_BG)
    tf = trig.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.25)
    tf.margin_top = Inches(0.15)
    add_text(tf, "再見積を上げるトリガー（説明用）", size=15, bold=True, color=ORANGE)
    for line in [
        "・B1 判定で「新規／大幅差分」が本案の仮置き（≒3本）を明確に超える",
        "・既存01–25のうち複数IFで再設計・再単体が必要（B3=15では足りない）",
        "・GCSアップロード・生鮮側接続が想定より重い／連携先経路が会議前提から変わる",
        "・目安：直接工数が本案想定を大きく超える場合は増分再見積",
    ]:
        add_para(tf, line, size=13, color=DARK, space_before=5)
    footer(s, 8, total)

    # ---- 9 前提除外 ----
    s = blank_slide(prs)
    banner(s, "8. 含むもの／含まないもの", "スコープ外を明示しないと、後から「全部入り」に読まれる")
    left = card(s, Inches(0.35), Inches(1.4), Inches(6.2), Inches(5.3), GREEN_BG)
    tf = left.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.22)
    add_text(tf, "含む（本案の範囲）", size=16, bold=True, color=GREEN)
    for line in [
        "・調査用表・Mapping（Phase1）",
        "・詳細設計・開発・単体（Phase2）",
        "・先行トラン新規IF 4本",
        "・マスタ棚卸・判定（B1）",
        "・マスタIF仮置き（B2）と非生鮮Phase2波及（B3）",
        "・増分PJ管理按分（C）",
        "・生鮮マスタのGCSアップロード確認・作業前提",
    ]:
        add_para(tf, line, size=14, color=DARK, space_before=8)

    right = card(s, Inches(6.8), Inches(1.4), Inches(6.1), Inches(5.3), ORANGE_BG)
    tf = right.text_frame
    tf.word_wrap = True
    tf.margin_top = Inches(0.05)
    tf.margin_left = Inches(0.22)
    add_text(tf, "含まない（明示除外）", size=16, bold=True, color=ORANGE)
    for line in [
        "・結合・総合・移行・運用",
        "・生鮮基幹／生鮮発注統合本体改修（先方）",
        "・青果@rms／鮮魚市場EOS（自動補充対象外）",
        "・共通基盤の二重構築",
        "・PMI課題 Sinops非生鮮のPhase2・共通基盤43・既存PJ50の契約変更",
        "・店舗系非生鮮Rの確定契約の読替え",
    ]:
        add_para(tf, line, size=14, color=DARK, space_before=10)
    footer(s, 9, total)

    # ---- 10 話し方 ----
    s = blank_slide(prs)
    banner(s, "9. 説明トーク（前提を先に言う）", "数字の前に「不確実性の置き方」を置く。前提が崩れたら再見積。")
    lines = [
        ("導入", "「生鮮は PMI課題 Sinops非生鮮のPhase2 の外の増分です。本日の提示は221.5人日／7,085,000円です。」"),
        ("前提先出し", "「ただし生鮮は不確実性が大きい。提示値は前提が成り立つ場合の工数で、確定請負ではありません。」"),
        ("最大変数", "「最大の変数は専用マスタです。流用可否・本数が未確定で、既存Sinops全体にも効きます。」"),
        ("本案の前提", "「本案はマスタ流用が概ね成立し、既存波及が限定的な場合の工数です。崩れたら再見積です。」"),
        ("仮置き", "「B2の51人日は仮置き。B1判定後に確定値へ置換します。」"),
        ("進め方", "「B1はAのPhase1と並行必須。ゲート後に確定させる進め方でお願いします。」"),
        ("範囲", "「Phase2は詳細設計・開発・単体まで。結合以降は含みません。」"),
    ]
    y = 1.35
    for title_, body in lines:
        lab = s.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.35), Inches(y), Inches(1.7), Inches(0.72)
        )
        lab.fill.solid()
        lab.fill.fore_color.rgb = NAVY
        lab.line.fill.background()
        lab.adjustments[0] = 0.2
        add_text(lab.text_frame, title_, size=12, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
        lab.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
        bx = s.shapes.add_textbox(Inches(2.2), Inches(y), Inches(10.7), Inches(0.72))
        add_text(bx.text_frame, body, size=13, color=DARK)
        y += 0.85
    footer(s, 10, total)

    # ---- 11 まとめ ----
    s = blank_slide(prs)
    banner(s, "10. まとめ", "本案を説明するとき、必ず一緒に伝える3点")
    items = [
        ("① 提示値", "221.5人日／¥7,085,000（税抜）\n非生鮮Phase2（PMI課題）は据え置き"),
        ("② 見積前提", "生鮮は不確実性が大\n前提崩壊＝再見積"),
        ("③ 運用条件", "B1並行 → 判定ゲート → B2置換\n超過時は再見積"),
    ]
    for i, (h, b) in enumerate(items):
        c = card(s, Inches(0.4 + i * 4.2), Inches(1.5), Inches(4.0), Inches(2.6), LIGHT)
        tf = c.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.2)
        tf.margin_top = Inches(0.2)
        add_text(tf, h, size=16, bold=True, color=NAVY)
        add_para(tf, b, size=14, color=DARK, space_before=12)

    end = card(s, Inches(0.4), Inches(4.4), Inches(12.5), Inches(2.3), ORANGE_BG)
    tf = end.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.25)
    tf.margin_top = Inches(0.2)
    add_text(tf, "添付・参照", size=14, bold=True, color=ORANGE)
    for line in [
        "・【見積】生鮮IF増分_標準案_…_20260921.xlsx",
        "・sources/20260918-自動補充PMI-生鮮IF会議/会議記録.md",
        "・本フォルダ README.md",
    ]:
        add_para(tf, line, size=13, color=DARK, space_before=4)
    footer(s, 11, total)

    prs.save(OUT)
    print(f"saved: {OUT}")
    return OUT


if __name__ == "__main__":
    build()
