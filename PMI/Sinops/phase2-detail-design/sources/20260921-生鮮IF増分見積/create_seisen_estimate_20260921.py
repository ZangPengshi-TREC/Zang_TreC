#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Generate standalone 生鮮 IF 増分見積 — 標準案 / リスク込み案 を別ファイルで出力。

既存 403 workbook は変更しない。
Rev.2026-09-21e: 分位用語を使わず「標準案」「リスク込み案」で分冊。各ファイルに PJ管理按分を含む。
Rev.2026-09-23: 書式整理（表頭・罫線・数値揃え・列幅・シート体裁）。
"""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

OUT = Path(__file__).resolve().parent
UNIT_P1 = 40000   # Phase1（調査・Mapping）
UNIT_P2 = 27500   # Phase2 直接（詳細設計・開発・単体）※PJ管理以外
UNIT_PJ = 40000   # PJ管理
UNIT = UNIT_P2    # 互換用エイリアス（Phase2直接）
PJ_RATIO = 50.0 / 499.5  # 既存: PJ管理50 / 開発スコープ499.5

# --- palette ---
NAVY = "1F4E78"
NAVY_SOFT = "2E75B6"
LIGHT = "D6E3F0"
GREEN = "C6EFCE"
GREEN_TXT = "006100"
ORANGE = "FCE4D6"
ORANGE_TXT = "C65911"
ZEBRA = "F5F8FB"
GRAY_TXT = "595959"
THIN = Side(style="thin", color="B0B0B0")
MED = Side(style="medium", color=NAVY)
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
FONT = "Yu Gothic"
FONT_BODY = Font(name=FONT, size=10, color="2F2F2F")
FONT_BOLD = Font(name=FONT, size=10, bold=True, color="2F2F2F")
FONT_HEAD = Font(name=FONT, size=10, bold=True, color="FFFFFF")
FONT_TITLE = Font(name=FONT, size=14, bold=True, color="FFFFFF")
FONT_SECTION = Font(name=FONT, size=11, bold=True, color=NAVY)
FONT_NOTE = Font(name=FONT, size=9, italic=True, color=GRAY_TXT)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)
LEFT_TOP = Alignment(horizontal="left", vertical="top", wrap_text=True)

P1 = [
    ("生鮮発注勧告(当日分)", "Sinops-S24", 1.5, 2.0, 1.0, 1.5, 6.0,
     "非生鮮24と項目・ファイルが違うため新規。連携先=生鮮発注統合前提"),
    ("生鮮発注勧告(翌日以降分)", "Sinops-S25", 1.0, 1.5, 1.0, 1.0, 4.5,
     "当日とレイアウト類似。未来スロット差分"),
    ("生鮮入荷実績", "Sinops-S11", 2.0, 2.5, 1.5, 2.0, 8.0,
     "2系統ソース：PC/CK=生鮮基幹、取引先=MD基幹"),
    ("生鮮入荷予定", "Sinops-S12", 1.5, 2.0, 1.0, 1.5, 6.0,
     "生鮮基幹新規ソース。Layer①別作"),
]
P1_TOTAL = sum(r[6] for r in P1)

L1 = [
    ("生鮮発注勧告(当日分)", "Sinops-S24", "TBD(≠R6)", 6.0, 6.0, 5.0, 17.0,
     "基準24の14に+3：連携先が生鮮発注統合"),
    ("生鮮発注勧告(翌日以降分)", "Sinops-S25", "TBD(≠R6)", 5.0, 5.0, 4.0, 14.0,
     "基準25の14。見積行は2本"),
    ("生鮮入荷実績", "Sinops-S11", "TBD(~6+合流)", 6.0, 6.0, 5.0, 17.0,
     "基準11の10に+7：2系統合流"),
    ("生鮮入荷予定", "Sinops-S12", "TBD(~8)", 5.0, 5.0, 4.0, 14.0,
     "基準12の10に+4：生鮮基幹新規"),
]
L1_TOTAL = sum(r[6] for r in L1)

L3 = [
    ("生鮮発注勧告(当日分)", "Sinops-S24", "TBD", 1.5, 1.5, 1.5, 4.5,
     "基準24の3に+1.5"),
    ("生鮮発注勧告(翌日以降分)", "Sinops-S25", "TBD", 1.0, 1.0, 1.0, 3.0,
     "基準25の3を維持"),
    ("生鮮入荷実績", "Sinops-S11", "TBD", 1.0, 1.5, 1.5, 4.0,
     "基準11の2.5に+1.5：合流後取込"),
    ("生鮮入荷予定", "Sinops-S12", "TBD", 0.5, 1.0, 1.0, 2.5,
     "基準12の2.5を維持"),
]
L3_TOTAL = sum(r[6] for r in L3)
TRAN_DETAIL = P1_TOTAL + L1_TOTAL + L3_TOTAL
TRAN_RISK = 120.0

MASTER_MUST = [
    ("生鮮マスタ棚卸・一覧化", 6.0,
     "生鮮専用マスタの種類・所在・GCS有無・Sinops参照関係を一覧化"),
    ("変換ルール単位の流用判定（1本ずつ）", 10.0,
     "既存01–05/13–15等の③変換と突合。流用／差分／新規の三択を付ける"),
    ("既存Sinops横断影響調査", 10.0,
     "マスタ前提が変わる場合の01–25設計・マッピング・単体観点への波及洗い出し"),
    ("生鮮側サーバー→GCS 上げ接続確認", 5.0,
     "会議: GCSに既にあれば再開発不要、無い見込み。見積にアップロード作業を含める"),
    ("判定結果のゲート資料・再見積更新", 4.0,
     "判定表→増分IF採番→本見積のマスタIF枠を確定値に更新"),
]
MASTER_MUST_STD = sum(r[1] for r in MASTER_MUST)
MASTER_MUST_RISK = 40.0
MASTER_IF_STD = 51.0
MASTER_IF_RISK = 95.0
IMPACT_403_STD = 15.0
IMPACT_403_RISK = 45.0

SCENARIOS = {
    "standard": {
        "label": "標準案",
        "outfile": "【見積】生鮮IF増分_標準案_調査Mappingから詳細設計開発単体_20260921.xlsx",
        "tran": TRAN_DETAIL,
        "b1": MASTER_MUST_STD,
        "b2": MASTER_IF_STD,
        "b3": IMPACT_403_STD,
        "a_note": "IF別明細の合計",
        "b1_note": "必須5作業の合計",
        "b2_note": "仮置き：新規／大幅差分 3本×17人日",
        "b3_note": "軽微パッチ・注記・共通変換の追い込み",
        "positioning": "マスタ流用が概ね成立し、波及が限定的な場合の標準工数",
        "tab": NAVY,
    },
    "risk": {
        "label": "リスク込み案",
        "outfile": "【見積】生鮮IF増分_リスク込み案_調査Mappingから詳細設計開発単体_20260921.xlsx",
        "tran": TRAN_RISK,
        "b1": MASTER_MUST_RISK,
        "b2": MASTER_IF_RISK,
        "b3": IMPACT_403_RISK,
        "a_note": "IF別明細に上振れを加味",
        "b1_note": "必須作業＋判定遅延バッファ",
        "b2_note": "仮置き：5本×17＋予備",
        "b3_note": "複数IF再設計・再単体",
        "positioning": "生鮮専用マスタの新規本数・既存波及が大きい場合の工数",
        "tab": ORANGE_TXT,
    },
}

OLD_FILES = [
    "【試算見積】生鮮IF増分_調査Mappingから詳細設計開発単体_20260921.xlsx",
    "【見積】生鮮IF増分_P80採用_調査Mappingから詳細設計開発単体_20260921.xlsx",
    "【見積】生鮮IF増分_P50_調査Mappingから詳細設計開発単体_20260921.xlsx",
    "【見積】生鮮IF増分_P80_調査Mappingから詳細設計開発単体_20260921.xlsx",
]


def fill(cell, color):
    cell.fill = PatternFill("solid", fgColor=color)


def set_widths(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def paint_title(ws, text, cols):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=cols)
    cell = ws.cell(1, 1, text)
    cell.font = FONT_TITLE
    fill(cell, NAVY)
    cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[1].height = 28
    for c in range(2, cols + 1):
        fill(ws.cell(1, c), NAVY)


def paint_section(ws, row, text, cols, accent=LIGHT):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=cols)
    cell = ws.cell(row, 1, text)
    cell.font = FONT_SECTION
    fill(cell, accent)
    cell.alignment = LEFT
    ws.row_dimensions[row].height = 22
    for c in range(2, cols + 1):
        fill(ws.cell(row, c), accent)


def paint_note(ws, row, text, cols):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=cols)
    cell = ws.cell(row, 1, text)
    cell.font = FONT_NOTE
    cell.alignment = LEFT_TOP
    ws.row_dimensions[row].height = 18


def paint_header(ws, row, headers):
    ws.row_dimensions[row].height = 32
    for c, h in enumerate(headers, 1):
        cell = ws.cell(row, c, h)
        cell.font = FONT_HEAD
        fill(cell, NAVY)
        cell.alignment = CENTER
        cell.border = BORDER


def style_data_row(ws, row, cols, *, num_cols=(), money_cols=(), zebra=False, total=False, warn=False):
    for c in range(1, cols + 1):
        cell = ws.cell(row, c)
        cell.border = BORDER
        if total:
            fill(cell, GREEN)
            cell.font = Font(name=FONT, size=10, bold=True, color=GREEN_TXT)
        elif warn:
            fill(cell, ORANGE)
            cell.font = FONT_BOLD
        elif zebra:
            fill(cell, ZEBRA)
            cell.font = FONT_BODY
        else:
            cell.font = FONT_BODY

        if c in num_cols:
            cell.alignment = CENTER
            if isinstance(cell.value, (int, float)):
                cell.number_format = "0.0"
        elif c in money_cols:
            cell.alignment = CENTER
            if isinstance(cell.value, (int, float)):
                cell.number_format = "#,##0"
        else:
            cell.alignment = LEFT if c == 1 or c == cols else CENTER
    ws.row_dimensions[row].height = 20 if not total else 22


def apply_sheet_chrome(ws, tab_color=NAVY):
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.page_margins.left = 0.4
    ws.page_margins.right = 0.4
    ws.page_margins.top = 0.5
    ws.page_margins.bottom = 0.5
    ws.sheet_properties.tabColor = tab_color
    ws.freeze_panes = "A2"


def money_breakdown(sc):
    """Phase1=4万／Phase2直接=2.75万／PJ=4万 で金額を分解する。"""
    a_p1 = P1_TOTAL
    a_p2 = sc["tran"] - a_p1
    if a_p2 < 0:
        raise ValueError(f"A Phase2 negative: tran={sc['tran']} p1={a_p1}")
    phase1 = a_p1 + sc["b1"]
    phase2 = a_p2 + sc["b2"] + sc["b3"]
    direct = phase1 + phase2
    pj = round(direct * PJ_RATIO)
    yen_a = int(a_p1 * UNIT_P1 + a_p2 * UNIT_P2)
    yen_b1 = int(sc["b1"] * UNIT_P1)
    yen_b2 = int(sc["b2"] * UNIT_P2)
    yen_b3 = int(sc["b3"] * UNIT_P2)
    yen_p1 = int(phase1 * UNIT_P1)
    yen_p2 = int(phase2 * UNIT_P2)
    yen_direct = yen_p1 + yen_p2
    yen_pj = int(pj * UNIT_PJ)
    return {
        "a_p1": a_p1,
        "a_p2": a_p2,
        "phase1": phase1,
        "phase2": phase2,
        "direct": direct,
        "pj": pj,
        "total_md": direct + pj,
        "yen_a": yen_a,
        "yen_b1": yen_b1,
        "yen_b2": yen_b2,
        "yen_b3": yen_b3,
        "yen_p1": yen_p1,
        "yen_p2": yen_p2,
        "yen_direct": yen_direct,
        "yen_pj": yen_pj,
        "yen_total": yen_direct + yen_pj,
    }


def build_summary(ws, sc, label, m):
    cols = 5
    paint_title(ws, f"西友MD基幹統合　生鮮IF増分 見積（{label}）", cols)
    apply_sheet_chrome(ws, sc["tab"])

    meta = [
        ("見積区分", label),
        ("作成日", "2026/09/21"),
        ("対象範囲", "生鮮増分のみ（Phase1調査・Mapping〜Phase2詳細設計・開発・単体テスト＋PJ管理按分）"),
        ("非生鮮Phase2", "PMI課題 Sinops非生鮮のPhase2。契約額は据え置き（本表は生鮮増分のみ）"),
    ]
    for i, (k, v) in enumerate(meta, 2):
        kcell = ws.cell(i, 1, k)
        kcell.font = FONT_BOLD
        fill(kcell, LIGHT)
        kcell.alignment = CENTER
        kcell.border = BORDER
        ws.merge_cells(start_row=i, start_column=2, end_row=i, end_column=cols)
        vcell = ws.cell(i, 2, v)
        vcell.font = FONT_BODY
        vcell.alignment = LEFT
        vcell.border = BORDER
        for c in range(3, cols + 1):
            ws.cell(i, c).border = BORDER
            fill(ws.cell(i, c), "FFFFFF")
        ws.row_dimensions[i].height = 20

    paint_section(ws, 7, "1. 提示値（このファイルの結論）", cols)
    paint_header(ws, 8, ["項目", "人日", "金額（円・税抜）", "単価", ""])
    headline = [
        ("Phase1直接（調査・Mapping）", m["phase1"], m["yen_p1"], f"{UNIT_P1:,} 円/人日", False),
        ("Phase2直接（設計・開発・単体）", m["phase2"], m["yen_p2"], f"{UNIT_P2:,} 円/人日", False),
        ("直接小計（A〜B3）", m["direct"], m["yen_direct"], "混在（上表）", False),
        ("PJ管理（按分）", m["pj"], m["yen_pj"], f"{UNIT_PJ:,} 円/人日", False),
        (f"合計　【{label}】", m["total_md"], m["yen_total"], "—", True),
    ]
    for i, (name, md, yen, unit, is_total) in enumerate(headline, 9):
        ws.cell(i, 1, name)
        ws.cell(i, 2, md)
        ws.cell(i, 3, yen)
        ws.cell(i, 4, unit)
        ws.cell(i, 5, "")
        style_data_row(ws, i, 4, num_cols={2}, money_cols={3}, total=is_total)
        if name.startswith("直接小計"):
            for c in range(1, 5):
                fill(ws.cell(i, c), LIGHT)
                ws.cell(i, c).font = FONT_BOLD
        if is_total:
            for c in range(1, 5):
                ws.cell(i, c).font = Font(name=FONT, size=12, bold=True, color=GREEN_TXT)

    paint_note(ws, 14, f"　{sc['positioning']}　／　単価：Phase1={UNIT_P1:,}・Phase2直接={UNIT_P2:,}・PJ管理={UNIT_PJ:,}（円/人日・税抜）", cols)

    paint_section(ws, 16, "2. 内訳", cols)
    paint_header(ws, 17, ["#", "ブロック", "人日", "金額（円）", "内容／単価区分"])
    blocks = [
        ("A", "先行トラン（新規IF 4本）", sc["tran"], m["yen_a"],
         f"Phase1 {m['a_p1']}人日×{UNIT_P1:,} ＋ Phase2 {m['a_p2']}人日×{UNIT_P2:,}", False, False),
        ("B1", "マスタ棚卸・判定（必須）", sc["b1"], m["yen_b1"],
         f"Phase1（調査・判定）×{UNIT_P1:,}", False, True),
        ("B2", "マスタIF本体（仮置き）", sc["b2"], m["yen_b2"],
         f"Phase2直接×{UNIT_P2:,}／{sc['b2_note']}", False, True),
        ("B3", "非生鮮Phase2（PMI課題）への波及", sc["b3"], m["yen_b3"],
         f"Phase2直接×{UNIT_P2:,}／{sc['b3_note']}", False, True),
        ("", "直接小計", m["direct"], m["yen_direct"], "A＋B1＋B2＋B3", True, False),
        ("C", "PJ管理", m["pj"], m["yen_pj"],
         f"×{UNIT_PJ:,}／進捗・品質・横断・MT（既存比率按分）", False, True),
        ("", f"合計（{label}）", m["total_md"], m["yen_total"], "＝ 提示値", True, False),
    ]
    for i, (num, name, md, yen, note, is_total, is_warn) in enumerate(blocks, 18):
        ws.cell(i, 1, num)
        ws.cell(i, 2, name)
        ws.cell(i, 3, md)
        ws.cell(i, 4, yen)
        ws.cell(i, 5, note)
        style_data_row(
            ws, i, 5, num_cols={3}, money_cols={4},
            total=is_total, warn=is_warn and not is_total,
            zebra=(i % 2 == 0 and not is_total and not is_warn),
        )
        ws.cell(i, 1).alignment = CENTER
        ws.cell(i, 5).alignment = LEFT

    paint_section(ws, 26, "3. 単価", cols)
    paint_header(ws, 27, ["対象", "単価（円/人日）", "備考", "", ""])
    for c in range(4, 6):
        fill(ws.cell(27, c), NAVY)
        ws.cell(27, c).border = BORDER
    unit_rows = [
        ("Phase1（調査・Mapping）", UNIT_P1, "AのPhase1分＋B1マスタ棚卸・判定"),
        ("Phase2直接（詳細設計・開発・単体）", UNIT_P2, "AのLayer①③＋B2＋B3。PJ管理以外"),
        ("PJ管理", UNIT_PJ, f"直接工数 × (50÷499.5) ≒ {m['pj']}人日"),
    ]
    for i, (a, b, c) in enumerate(unit_rows, 28):
        ws.cell(i, 1, a)
        ws.cell(i, 2, b)
        ws.merge_cells(start_row=i, start_column=3, end_row=i, end_column=5)
        ws.cell(i, 3, c)
        style_data_row(ws, i, 3, money_cols={2})
        for col in range(4, 6):
            ws.cell(i, col).border = BORDER

    paint_section(ws, 32, "4. 見積前提（不確実性含む・必読）", cols)
    notes = [
        "【本件の性質】生鮮は調査未完の領域が多い。本提示値は「現時点の前提が成り立つ場合」の工数であり、確定請負額ではない。",
        "【最大の不確実性】生鮮専用マスタの種類・本数・非生鮮変換流用可否が未確定。影響は先行4本に閉じず PMI課題 Sinops非生鮮のPhase2（既存Sinops）全体に及ぶ。",
        "【本ファイルの置き方】B1＝判定必須（Aと並行）。B2＝マスタIF仮置き。B3＝非生鮮Phase2（PMI課題）への波及バッファ。B1判定後にB2を確定値へ置換し、超過時は再見積。",
        f"【本案が前提とする状況】{sc['positioning']}。",
        "【未確定の設計】勧告の連携先経路、棚割（SM/新）、営業在庫の生鮮適用、TM商品・停止・新商品在庫15等。詳細は「見積前提条件・注記」シート。",
        "【動かさないもの】PMI課題 Sinops非生鮮のPhase2・共通基盤43・既存PJ管理50の契約行。結合以降・生鮮基幹本体は含まない。",
        f"【単価】Phase1={UNIT_P1:,}円／Phase2直接={UNIT_P2:,}円／PJ管理={UNIT_PJ:,}円（いずれも人日・税抜）。",
    ]
    for i, t in enumerate(notes, 33):
        paint_note(ws, i, "・" + t, cols)
        ws.row_dimensions[i].height = 30
        cell = ws.cell(i, 1)
        cell.font = Font(name=FONT, size=9, color="2F2F2F")
        bg = ORANGE if i in (33, 34) else ZEBRA
        fill(cell, bg)
        for c in range(2, cols + 1):
            fill(ws.cell(i, c), bg)

    set_widths(ws, [14, 36, 12, 16, 48])
    ws.row_dimensions[8].height = 22
    ws.row_dimensions[17].height = 22


def write_if_table(ws, start_row, section_title, headers, rows, total_label, total_value, cols=8):
    """Write section banner + header + data + subtotal. Returns next free row."""
    paint_section(ws, start_row, section_title, cols)
    hr = start_row + 1
    paint_header(ws, hr, headers)
    r = hr + 1
    num_cols = {i for i, h in enumerate(headers, 1) if h in (
        "調査用表", "①項目マッピング表", "②GAP分析書", "③変換ルール定義書",
        "詳細設計", "開発", "単体テスト", "小計", "人日", "人日（明細）",
    )}
    for idx, row in enumerate(rows):
        for c, v in enumerate(row, 1):
            ws.cell(r, c, v)
        style_data_row(
            ws, r, cols, num_cols=num_cols,
            zebra=(idx % 2 == 1),
        )
        # 備考列 left
        ws.cell(r, cols).alignment = LEFT
        ws.cell(r, 1).alignment = LEFT
        r += 1
    # subtotal
    ws.cell(r, 1, total_label)
    for c in range(2, cols):
        ws.cell(r, c, "")
    # put total in 小計 column (usually 7) or col 2 if short
    subtotal_col = 7 if cols >= 7 else 2
    ws.cell(r, subtotal_col, total_value)
    style_data_row(ws, r, cols, num_cols={subtotal_col}, total=True)
    return r + 2


def build_sheet_a(ws, sc, label):
    cols = 8
    paint_title(ws, f"【A】先行トラン4本 — Phase1＋Phase2（IF別明細／本ファイルA={sc['tran']}）", cols)
    apply_sheet_chrome(ws, sc["tab"])
    paint_note(
        ws, 2,
        f"単位：人日。本見積ファイルのAブロック値は {sc['tran']}（{label}）。"
        "略語：Phase1＝調査・Mapping、Layer①＝TRIAL↔GCS、Layer②＝GCS→DWH、Layer③＝DWH・GCS↔業務。",
        cols,
    )
    ws.row_dimensions[2].height = 32

    r = 4
    r = write_if_table(
        ws, r,
        "■ Phase1（調査・Mapping：調査用表＋①項目マッピング表／②GAP分析書／③変換ルール定義書）",
        ["IF名称", "仮IF No.", "調査用表", "①項目マッピング表", "②GAP分析書", "③変換ルール定義書", "小計", "備考"],
        [[x[0], x[1], x[2], x[3], x[4], x[5], x[6], x[7]] for x in P1],
        "小計 Phase1", P1_TOTAL, cols,
    )
    r = write_if_table(
        ws, r,
        "■ Layer① TRIAL↔GCS（外部IFプログラム開発）",
        ["IF名称", "仮IF No.", "項目数", "詳細設計", "開発", "単体テスト", "小計", "備考"],
        [list(x) for x in L1],
        "小計 Layer①", L1_TOTAL, cols,
    )
    paint_section(ws, r, "■ Layer② GCS→DWH（DataSpider取込ジョブ）＝ 0（全件）", cols, LIGHT)
    r += 2
    r = write_if_table(
        ws, r,
        "■ Layer③ DWH・GCS↔業務システム（DataSpider連携ジョブ）",
        ["IF名称", "仮IF No.", "項目数", "詳細設計", "開発", "単体テスト", "小計", "備考"],
        [list(x) for x in L3],
        "小計 Layer③", L3_TOTAL, cols,
    )

    ws.cell(r, 1, f"Aブロック（{label}）")
    for c in range(2, 7):
        ws.cell(r, c, "")
    ws.cell(r, 7, sc["tran"])
    ws.cell(r, 8, f"IF別明細合計={TRAN_DETAIL}　／　本ファイル={sc['tran']}")
    style_data_row(ws, r, cols, num_cols={7}, total=True)
    ws.cell(r, 8).alignment = LEFT

    set_widths(ws, [26, 13, 12, 14, 12, 14, 8, 42])


def build_sheet_b(ws, sc, label):
    cols = 4
    paint_title(ws, f"【B】生鮮専用マスタ横断リスク（{label}）", cols)
    apply_sheet_chrome(ws, sc["tab"])
    paint_section(
        ws, 2,
        "生鮮専用マスタは先行トランに閉じない。非生鮮中心のSinops全体に効く。",
        cols, ORANGE,
    )
    ws.cell(2, 1).font = Font(name=FONT, size=10, bold=True, color=ORANGE_TXT)

    paint_section(ws, 4, f"■ B1 必須作業／本ファイル={sc['b1']}", cols)
    paint_header(ws, 5, ["作業", "人日（明細）", "成果物", "備考"])
    deliverables = [
        "マスタ一覧・所在マップ",
        "IF別 流用/差分/新規 判定表",
        "既存01–25影響リスト・改修要否",
        "GCS接続・上げ手順の確認結果",
        "増分IF採番案・見積更新版",
    ]
    for i, ((name, days, note), d) in enumerate(zip(MASTER_MUST, deliverables), 6):
        ws.cell(i, 1, name)
        ws.cell(i, 2, days)
        ws.cell(i, 3, d)
        ws.cell(i, 4, note)
        style_data_row(ws, i, cols, num_cols={2}, zebra=(i % 2 == 0))
        ws.cell(i, 1).alignment = LEFT
        ws.cell(i, 3).alignment = LEFT
        ws.cell(i, 4).alignment = LEFT_TOP
        ws.row_dimensions[i].height = 28
    ws.cell(11, 1, f"B1 本ファイル（{label}）")
    ws.cell(11, 2, sc["b1"])
    ws.cell(11, 3, "")
    ws.cell(11, 4, "")
    style_data_row(ws, 11, cols, num_cols={2}, total=True)

    paint_section(ws, 13, "■ B2 マスタIF本体（仮置き）", cols)
    paint_header(ws, 14, ["想定", "人日", "算定", "置換条件"])
    ws.cell(15, 1, "新規または大幅差分のマスタIF")
    ws.cell(15, 2, sc["b2"])
    ws.cell(15, 3, sc["b2_note"])
    ws.cell(15, 4, "B1判定後に行を分解して置換")
    style_data_row(ws, 15, cols, num_cols={2}, warn=True)
    ws.cell(15, 1).alignment = LEFT
    ws.cell(15, 3).alignment = LEFT
    ws.cell(15, 4).alignment = LEFT
    ws.cell(16, 1, "候補例（未確定）")
    ws.merge_cells("C16:D16")
    ws.cell(16, 3, "TM商品、店舗商品差分、カテゴリ/仕入先、停止フラグ、発注スケジュール、新商品在庫15 等")
    style_data_row(ws, 16, cols, zebra=True)
    ws.cell(16, 1).alignment = LEFT
    ws.cell(16, 3).alignment = LEFT
    ws.row_dimensions[16].height = 28

    paint_section(ws, 18, "■ B3 非生鮮Phase2（PMI課題）への波及", cols)
    paint_header(ws, 19, ["内容", "人日", "説明", ""])
    fill(ws.cell(19, 4), NAVY)
    ws.cell(19, 4).border = BORDER
    ws.cell(20, 1, "既存Sinops-01〜25の手戻り")
    ws.cell(20, 2, sc["b3"])
    ws.merge_cells("C20:D20")
    ws.cell(20, 3, sc["b3_note"] + "。非生鮮Phase2（PMI課題）の契約額は据え置き")
    style_data_row(ws, 20, 3, num_cols={2}, warn=True)
    ws.cell(20, 4).border = BORDER
    ws.cell(20, 1).alignment = LEFT
    ws.cell(20, 3).alignment = LEFT

    paint_section(ws, 22, "■ 影響が及び得る既存IF（例）", cols)
    paint_header(ws, 23, ["既存IF", "なぜ効くか", "リスク", ""])
    fill(ws.cell(23, 4), NAVY)
    ws.cell(23, 4).border = BORDER
    impacts = [
        ("01/02 商品・店別", "TM商品・停止・発注期間が生鮮専用の可能性", "変換流用不可→別IF or 分岐"),
        ("03/04 カテゴリ・仕入先", "生鮮階層・仕入先体系が別", "コード体系の再マッピング"),
        ("05 ケースバラ", "生鮮は計量・パック差", "パターン追加"),
        ("13 発注スケジュール", "生鮮締め・便・曜日が別源", "Layer①ソース分岐"),
        ("14/15 棚割・新商品在庫", "生鮮は店舗調達を使わない可能性", "15個別対応・14ダミー可否"),
        ("06–10 実績・在庫", "営業在庫の生鮮適用未検証", "計算前提の再確認"),
        ("11/12/24/25", "先行トラン自体が新ソース", "本見積Aで別計上済"),
    ]
    for i, row in enumerate(impacts, 24):
        ws.cell(i, 1, row[0])
        ws.cell(i, 2, row[1])
        ws.merge_cells(start_row=i, start_column=3, end_row=i, end_column=4)
        ws.cell(i, 3, row[2])
        style_data_row(ws, i, 3, zebra=(i % 2 == 0))
        ws.cell(i, 4).border = BORDER
        ws.cell(i, 1).alignment = LEFT
        ws.cell(i, 2).alignment = LEFT
        ws.cell(i, 3).alignment = LEFT

    b_sum = sc["b1"] + sc["b2"] + sc["b3"]
    ws.cell(32, 1, f"B合計（{label}）")
    ws.cell(32, 2, b_sum)
    ws.merge_cells("C32:D32")
    ws.cell(32, 3, f"B1 {sc['b1']} + B2 {sc['b2']} + B3 {sc['b3']}")
    style_data_row(ws, 32, 3, num_cols={2}, warn=True)
    ws.cell(32, 4).border = BORDER
    ws.cell(32, 3).alignment = LEFT

    set_widths(ws, [34, 12, 36, 42])


def build_sheet_c(ws, sc, label, direct, pj, yen_pj):
    cols = 5
    paint_title(ws, f"【C】PJ管理（増分按分・{label}）", cols)
    apply_sheet_chrome(ws, sc["tab"])
    paint_note(
        ws, 2,
        "既存PJ管理50は開発スコープ499.5横断分。本増分は同じ比率50/499.5を本ファイル直接工数に適用。",
        cols,
    )
    paint_header(ws, 4, ["項目", "算定", "人日", "単価", "金額(円)"])
    ws.cell(5, 1, f"PJ管理（{label}）")
    ws.cell(5, 2, f"{direct} × (50/499.5) ≈ {pj}")
    ws.cell(5, 3, pj)
    ws.cell(5, 4, UNIT_PJ)
    ws.cell(5, 5, yen_pj)
    style_data_row(ws, 5, cols, num_cols={3}, money_cols={4, 5}, total=True)
    ws.cell(5, 2).alignment = LEFT

    ws.cell(7, 1, "含む作業")
    ws.merge_cells("B7:E7")
    ws.cell(7, 2, "進捗・品質管理、横断課題対応、会議体（MT）対応など")
    style_data_row(ws, 7, 2)
    for c in range(3, 6):
        ws.cell(7, c).border = BORDER
    ws.cell(7, 1).alignment = CENTER
    fill(ws.cell(7, 1), LIGHT)
    ws.cell(7, 2).alignment = LEFT

    ws.cell(8, 1, "既存契約")
    ws.merge_cells("B8:E8")
    ws.cell(8, 2, "既存PJ管理50の契約行は変更しない。本表は生鮮増分の追加PJ管理。")
    style_data_row(ws, 8, 2)
    for c in range(3, 6):
        ws.cell(8, c).border = BORDER
    ws.cell(8, 1).alignment = CENTER
    fill(ws.cell(8, 1), LIGHT)
    ws.cell(8, 2).alignment = LEFT

    set_widths(ws, [18, 42, 10, 12, 14])


def build_sheet_diff(ws, sc):
    cols = 6
    paint_title(ws, "既存IF流用可否と新規判定", cols)
    apply_sheet_chrome(ws, sc["tab"])
    paint_note(ws, 2, "トランは新規IF必須。マスタは判定待ち（B1）。", cols)
    paint_header(ws, 4, ["対象", "既存対照", "レイアウト", "ソース", "連携先", "判定"])
    diffs = [
        ("生鮮勧告 当日/翌日", "24/25", "項目・ファイルが違う", "Sinops復路", "生鮮発注統合", "新規IF必須"),
        ("生鮮入荷実績", "11", "近い可能性（要精査）", "生鮮基幹+MD基幹", "Sinops", "新規IF必須（2系統ソース）"),
        ("生鮮入荷予定", "12", "近い可能性（要精査）", "生鮮基幹", "Sinops", "新規IF必須"),
        ("生鮮専用マスタ群", "01–05/13–15等", "不明（大量）", "生鮮側サーバー", "Sinops", "棚卸後に流用/新規判定"),
    ]
    for i, row in enumerate(diffs, 5):
        for c, v in enumerate(row, 1):
            ws.cell(i, c, v)
        style_data_row(ws, i, cols, warn=(i == 8), zebra=(i % 2 == 0 and i != 8))
        for c in range(1, cols + 1):
            ws.cell(i, c).alignment = LEFT if c in (1, 3, 4, 6) else CENTER
    set_widths(ws, [22, 14, 26, 22, 14, 24])


def build_sheet_notes(ws, sc, label, m):
    paint_title(ws, f"見積前提条件・注記（生鮮IF増分 {label} Rev.e）", 5)
    apply_sheet_chrome(ws, sc["tab"])
    paint_section(
        ws, 2,
        "【必読】生鮮は不確実性が大きい。下表の前提が崩れた場合、提示人日・金額は再見積対象となる。",
        5, ORANGE,
    )
    ws.cell(2, 1).font = Font(name=FONT, size=11, bold=True, color=ORANGE_TXT)

    notes = [
        ("■ 0. 見積の読み方（最重要）", True, True),
        (f"1. 本ファイルは【{label}】の独立見積。調査未完の前提を置いたうえでの提示値であり、確定請負額ではない。", False, False),
        ("2. 生鮮領域は「PMI課題 Sinops非生鮮のPhase2」と同じ確度では見積もることができない。不確実性を前提に織り込んでいる。", False, False),
        ("3. B1（マスタ棚卸・判定）完了前にB2本数は確定しない。B2は仮置き。判定後に置換し、超過時は再見積する。", False, False),
        (f"4. 本案の前提状況：{sc['positioning']}。", False, False),
        ("", False, False),
        ("■ 1. 範囲・契約関係・単価", True, False),
        ("1. 対象工程：Phase1（調査・Mapping）＋Phase2（詳細設計・開発・単体テスト）＋増分PJ管理按分。", False, False),
        (f"2. 単価：Phase1 {UNIT_P1:,}円／Phase2直接 {UNIT_P2:,}円／PJ管理 {UNIT_PJ:,}円（税抜・人日）。", False, False),
        ("3. Phase1単価の対象：AのPhase1（調査用表・①②③）＋B1マスタ棚卸・判定。", False, False),
        ("4. Phase2直接単価の対象：AのLayer①③＋B2マスタIF＋B3波及（PJ管理以外）。", False, False),
        ("5. PMI課題 Sinops非生鮮のPhase2の契約額は据え置き（本表は生鮮増分のみ）。", False, False),
        ("6. 共通基盤43・BO・既存PJ管理50の契約行も変更しない。", False, False),
        ("", False, False),
        ("■ 2. 生鮮の不確実性（見積前提として明示）", True, True),
        ("【マスタ】生鮮は専用マスタを多数持つ。非生鮮の③変換ルールで足りるか／何本新規かは未確定。", False, False),
        ("【横断影響】マスタ前提が変わると PMI課題 Sinops非生鮮のPhase2（Sinops-01〜25）の設計・Mapping・単体にも波及し得る。", False, False),
        ("【源・GCS】マスタ源は生鮮側サーバー→GCS。GCSに既に無ければアップロード作業が必要（会議見込：無い）。本見積に確認・アップロード作業を含む。", False, False),
        ("【判定方法】変換ルールを1本ずつ「流用／差分／新規」の三択で判定する（B1）。一括流用は前提にしない。波多野確認：マスタの違いがあるので各IFのマッピングから始める必要がある。", False, False),
        ("【候補例・未確定】TM商品、店舗商品差分、カテゴリ／仕入先、停止フラグ、発注スケジュール、棚割14、新商品在庫15 等。", False, False),
        ("【棚割】Store Manager継続 vs 新棚割は未確定。ダミー棚の入り方により14の生鮮対応が分岐する。", False, False),
        ("【在庫】営業在庫を生鮮で使えるかは未検証。06–10系の前提再確認が必要になり得る。", False, False),
        ("【勧告の連携先】生鮮発注統合への直送か、自動発注GCS経由かは未確定。連携先変更はLayer①工数に効く。", False, False),
        ("【入荷】入荷実績は2系統ソース（生鮮基幹＋MD基幹）、入荷予定は生鮮基幹の新規ソース。非生鮮11/12の単純流用は不可。", False, False),
        ("【便】sinopsは便を持たず渡すだけ（振り分けは生鮮発注統合）。必須項目なら固定値で足りる前提。", False, False),
        ("", False, False),
        ("■ 3. 本案での数値の置き方", True, False),
        ("1. A（先行トラン4本）は会議で骨格確定した新規IF。残リスクは項目精査・合流・連携先の細部。", False, False),
        ("2. B1＝判定そのもの（必須・AのPhase1と並行）。各IFのマッピングから着手（マスタ差分があるため）。", False, False),
        (f"3. B2＝マスタIF本体の仮置き（{sc['b2_note']}）。B1後に行分解して確定値へ置換。", False, False),
        (f"4. B3＝非生鮮Phase2（PMI課題）への波及バッファ（{sc['b3_note']}）。契約額自体は動かさない。", False, False),
        ("5. 仮置き超過・連携先変更・棚割方針確定などで前提が崩れた場合は増分再見積とする。", False, False),
        ("", False, False),
        ("■ 4. 略語", True, False),
        ("1. Phase1＝調査・Mapping（調査用表／①項目マッピング表／②GAP分析書／③変換ルール定義書）。", False, False),
        ("2. Phase2＝詳細設計・開発・単体テスト（Layer①②③の工数を含む）。", False, False),
        ("3. Layer①＝TRIAL↔GCS（外部IFプログラム）。Layer②＝GCS→DWH。Layer③＝DWH・GCS↔業務。", False, False),
        ("4. PJ管理＝プロジェクト管理。GCS＝Google Cloud Storage。IF＝インターフェース。", False, False),
        ("", False, False),
        ("■ 5. 本ファイル数値", True, False),
        (f"1. Phase1={m['phase1']}人日（{m['yen_p1']:,}円）＋Phase2直接={m['phase2']}人日（{m['yen_p2']:,}円）＝直接{m['direct']}人日（{m['yen_direct']:,}円）。", False, False),
        (f"2. ＋PJ管理={m['pj']}人日（{m['yen_pj']:,}円）＝合計{m['total_md']}人日／{m['yen_total']:,}円。", False, False),
        (f"3. 内訳 A={sc['tran']}（Phase1 {m['a_p1']}+Phase2 {m['a_p2']}）＋B1={sc['b1']}＋B2={sc['b2']}＋B3={sc['b3']}。", False, False),
        ("4. PJ管理按分＝直接×(50/499.5)。", False, False),
        ("", False, False),
        ("■ 6. 除外（本見積に含まない）", True, False),
        ("1. 結合・総合・移行・運用テスト。", False, False),
        ("2. 生鮮基幹／生鮮発注統合本体の改修（先方）。", False, False),
        ("3. 青果@rms／鮮魚市場EOS等（自動補充対象外）。", False, False),
        ("4. 共通基盤の二重構築。", False, False),
        ("5. 店舗系非生鮮R（24/25）確定契約の読替え。", False, False),
    ]
    r = 3
    for text, is_sec, is_warn in notes:
        if not text:
            r += 1
            continue
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=5)
        cell = ws.cell(r, 1, text)
        if is_sec:
            cell.font = FONT_SECTION
            bg = ORANGE if is_warn else LIGHT
            fill(cell, bg)
            for c in range(2, 6):
                fill(ws.cell(r, c), bg)
            ws.row_dimensions[r].height = 22
        else:
            cell.font = FONT_BODY
            cell.alignment = LEFT_TOP
            ws.row_dimensions[r].height = 20 if len(text) > 60 else 18
        r += 1
    set_widths(ws, [110, 12, 12, 12, 12])
    ws.column_dimensions["A"].width = 118
    ws.sheet_view.zoomScale = 90


def build_sheet_basis(ws, sc, label, m):
    cols = 4
    paint_title(ws, f"算定根拠（{label}）", cols)
    apply_sheet_chrome(ws, sc["tab"])
    paint_note(
        ws, 2,
        f"本ファイルは【{label}】の独立見積。単価：Phase1={UNIT_P1:,}／Phase2直接={UNIT_P2:,}／PJ={UNIT_PJ:,}。",
        cols,
    )
    paint_header(ws, 4, ["項目", "基準/仮定", "人日", "金額・根拠"])
    rows = [
        ("A Phase1", "調査・Mapping明細", m["a_p1"], f"{int(m['a_p1']*UNIT_P1):,}円（×{UNIT_P1:,}）", False),
        ("A Phase2", "Layer①③（Layer②=0）", m["a_p2"], f"{int(m['a_p2']*UNIT_P2):,}円（×{UNIT_P2:,}）", False),
        ("A 小計", sc["a_note"], sc["tran"], f"{m['yen_a']:,}円", False),
        ("B1 必須", sc["b1_note"] + "／Phase1単価", sc["b1"], f"{m['yen_b1']:,}円", False),
        ("B2 マスタIF", sc["b2_note"] + "／Phase2直接", sc["b2"], f"{m['yen_b2']:,}円", False),
        ("B3 非生鮮Phase2波及", sc["b3_note"] + "／Phase2直接", sc["b3"], f"{m['yen_b3']:,}円", False),
        ("直接小計", f"Phase1 {m['phase1']}+Phase2 {m['phase2']}", m["direct"], f"{m['yen_direct']:,}円", True),
        ("C PJ管理", "50/499.5 按分", m["pj"], f"{m['yen_pj']:,}円（×{UNIT_PJ:,}）", False),
        (f"合計（{label}）", "", m["total_md"], f"{m['yen_total']:,}円", True),
    ]
    for i, (a, b, c, d, is_total) in enumerate(rows, 5):
        ws.cell(i, 1, a)
        ws.cell(i, 2, b)
        ws.cell(i, 3, c)
        ws.cell(i, 4, d)
        style_data_row(
            ws, i, cols, num_cols={3}, total=is_total,
            warn=(a.startswith("B") and not is_total),
            zebra=(i % 2 == 0 and not is_total),
        )
        ws.cell(i, 1).alignment = LEFT
        ws.cell(i, 2).alignment = LEFT
        ws.cell(i, 4).alignment = LEFT
    paint_note(ws, 15, "出典: 20260918生鮮IF会議、センター認識合わせ、既存見積Sinops明細・サマリー、単価区分は本件指定", cols)
    set_widths(ws, [22, 40, 10, 36])


def build_one(scenario_key: str) -> Path:
    sc = SCENARIOS[scenario_key]
    label = sc["label"]
    m = money_breakdown(sc)

    wb = Workbook()
    ws = wb.active
    ws.title = "サマリー"
    build_summary(ws, sc, label, m)

    ws5 = wb.create_sheet("見積前提条件・注記")
    build_sheet_notes(ws5, sc, label, m)

    ws2 = wb.create_sheet("A_先行トラン明細")
    build_sheet_a(ws2, sc, label)

    ws3 = wb.create_sheet("B_マスタ横断リスク")
    build_sheet_b(ws3, sc, label)

    ws_pj = wb.create_sheet("C_PJ管理")
    build_sheet_c(ws_pj, sc, label, m["direct"], m["pj"], m["yen_pj"])

    ws4 = wb.create_sheet("対既存差分")
    build_sheet_diff(ws4, sc)

    ws6 = wb.create_sheet("算定根拠")
    build_sheet_basis(ws6, sc, label, m)

    path = OUT / sc["outfile"]
    wb.save(path)
    print(
        f"saved {path.name}: direct={m['direct']} "
        f"(P1={m['phase1']}@{UNIT_P1}/P2={m['phase2']}@{UNIT_P2}) "
        f"pj={m['pj']} total={m['total_md']} yen={m['yen_total']}"
    )
    return path



def build():
    paths = [build_one("standard"), build_one("risk")]
    for name in OLD_FILES:
        old = OUT / name
        if old.exists():
            old.unlink()
            print("removed", name)
    return paths


if __name__ == "__main__":
    build()
