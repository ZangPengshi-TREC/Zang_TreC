#!/usr/bin/env python3
"""御見積書（正式版）: 生鮮IF増分 標準案。惣菜DX1.0 見積書テンプレ（4頁）に準拠。

数値は create_seisen_estimate_20260921.py の標準案から取る（二重管理しない）。
出力: PDF（PyMuPDF描画）と Excel（編集用）。
"""
from __future__ import annotations

import sys
from pathlib import Path

import fitz
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
import create_seisen_estimate_20260921 as est  # noqa: E402

# ---------------------------------------------------------------- 設定
DOC_DATE = "2026/10/07"
QUOTE_NO = "（採番待ち）"
SYSTEM = "（採番待ち）"
PROJECT = "（採番待ち）"
BRANCH = "（採番待ち）"
SUBJECT = "生鮮IF増分（自動補充 生鮮IF）"
PJ_FORM = "プロジェクト"
CLIENT = "(株)トライアルテクノロジー"
CONTACT = "波多野 弘一"
PAYMENT = "検収月20日〆翌月払い"
VALID = "1ヶ月間"
CONTRACT = "開発"
TAX_RATE = "0%"
STAMPS = ["李萌", "宮相聡", "臧鵬仕"]
ISSUER = ["TRE-CHINA青島", "中国・山東省青島市市北区龍城路39号",
          "二十二世紀広電産業園10階", "TEL：86-532-85938100"]
PAGE_SUB = "生鮮IF増分"

MONTHS = ["10月", "11月", "12月"]
# 月別（人日）。4人体制の概算日程（10/5着手・12/8完了目安）の月別営業日に合わせた概算配分。
MONTH_ALLOC = {
    "map": [39.5, 0.0, 0.0],
    "p2": [32.5, 64.5, 7.0],
    "pj": [4.0, 7.0, 3.0],
}

SC = est.SCENARIOS["standard"]
mb = est.money_breakdown(SC)
P1_DAYS = float(mb["phase1"])
P2_DAYS = float(mb["phase2"])
PJ_DAYS = float(mb["pj"])
U1, U2, UPJ = est.UNIT_P1, est.UNIT_P2, est.UNIT_PJ

TASKS = [  # (名称, 単価, 工数)  工数 None = 対象外（空欄）
    ("調査・Mapping", U1, P1_DAYS),
    ("詳細設計・開発・単体テスト", U2, P2_DAYS),
    ("プロジェクト管理", UPJ, PJ_DAYS),
    ("基本設計", None, None),
    ("概要設計", None, None),
    ("結合テスト仕様書", None, None),
    ("結合テスト", None, None),
    ("総合テスト仕様書", None, None),
    ("総合テスト", None, None),
    ("操作マニュアル", None, None),
    ("移行準備", None, None),
    ("移行作業", None, None),
    ("本番稼働立会い", None, None),
]
AMT = [round(u * d) if d is not None else None for _, u, d in TASKS]
TOTAL_DAYS = P1_DAYS + P2_DAYS + PJ_DAYS
TOTAL_AMT = sum(a for a in AMT if a)
assert TOTAL_AMT == 5_000_000, TOTAL_AMT

DELIVERY = [
    ("1回目", "生鮮IF増分 調査・Mapping完了", "202610", "202611", AMT[0]),
    ("2回目", "生鮮IF増分 詳細設計・開発・単体テスト完了", "202612", "202701", AMT[1] + AMT[2]),
]
assert sum(d[4] for d in DELIVERY) == TOTAL_AMT

NOTES = [
    "・本見積の範囲は調査・Mapping〜詳細設計・開発・単体テストです。結合テスト以降、移行、本番稼働立会い、生鮮基幹／生鮮発注統合本体、@rms／EOS、倉庫系は含みません。",
    "・店舗系12本の既存IFは、先方協議での流用方針を前提とします。当社は協議に未同席のため、流用可否の突合を本見積（流用確認・残件）に含みます。",
    "・IF番号 S11〜S26 は仮採番です。前提（流用判定・生鮮データ作成時刻・連携先仕様）が変わる場合は、再見積のうえ調整します。",
    "・体制4名、2026/10/05着手、完了目安は2026/12/08です（着手遅れ・前提変更は後ろ倒し）。",
]

# 機能別: (機能, IF名, 区分, 新規/既存, 調査, 詳細設計, 開発, 単体, 内訳なし)
def _by_name(rows):
    return {r[0]: r for r in rows}

P1 = _by_name(est.P1_STD)
L1 = _by_name(est.L1_STD)
L3 = _by_name(est.L3_STD)

IF_GROUPS = [
    ("生鮮発注勧告", [("生鮮発注勧告(当日分)", "新規"),
                      ("生鮮発注勧告(翌日以降分)", "新規"),
                      ("生鮮勧告→発注統合変換", "新規")]),
    ("生鮮入荷・発注", [("生鮮入荷実績", "差分"),
                        ("生鮮入荷予定", "差分"),
                        ("生鮮発注スケジュール", "差分")]),
]
FUNC_ROWS = []  # (機能, 名称, 区分, 新規/既存, map, dd, dev, ut, other)
for grp, items in IF_GROUPS:
    for i, (name, kind) in enumerate(items):
        dd = L1[name][3] + L3[name][3]
        dev = L1[name][4] + L3[name][4]
        ut = L1[name][5] + L3[name][5]
        label = f"{name}（{L1[name][1].split('（')[0]}）"
        FUNC_ROWS.append((grp if i == 0 else "", label, "IF", kind,
                          P1[name][6], dd, dev, ut, 0.0))
for i, (name, days, _out, _note) in enumerate(est.B1_STD):
    short = name.split("（")[0]
    FUNC_ROWS.append(("流用確認・残件" if i == 0 else "", short, "調査", "流用", days, 0.0, 0.0, 0.0, 0.0))
FUNC_ROWS.append(("非生鮮Phase2波及", "伝票区分・振替の注記反映", "改修", "既存", 0.0, 0.0, 0.0, 0.0, SC["b3"]))
_chk = sum(r[4] + r[5] + r[6] + r[7] + r[8] for r in FUNC_ROWS)
assert abs(_chk - (P1_DAYS + P2_DAYS)) < 1e-6, (_chk, P1_DAYS + P2_DAYS)

# ---------------------------------------------------------------- PDF
FONT_PATH = "/System/Library/Fonts/ヒラギノ明朝 ProN.ttc"
NAVY = (31 / 255, 78 / 255, 120 / 255)
LBLUE = (232 / 255, 241 / 255, 249 / 255)
GRAY = (242 / 255, 242 / 255, 242 / 255)
LINE = (0.55, 0.55, 0.55)
BLACK = (0, 0, 0)
WHITE = (1, 1, 1)
K = 595.0 / 654.0  # テンプレ画像px → pt


def fmt_int(n):
    return f"{n:,}"


def fmt_d(n):
    return f"{n:,.1f}"


class Pdf:
    def __init__(self):
        self.doc = fitz.open()
        self.font = fitz.Font(fontfile=FONT_PATH)
        self.page = None

    def new_page(self):
        self.page = self.doc.new_page(width=595, height=842)
        self.page.insert_font(fontname="min", fontfile=FONT_PATH)
        return self.page

    # 座標はすべてテンプレpx（654幅）で受け、ptへ換算
    def text(self, x, y, s, size=9.0, align="l", color=BLACK, w=None):
        """y は文字の縦中央。align: l/c/r。c/r は w（幅）が必要。"""
        if s is None or s == "":
            return
        tl = self.font.text_length(s, fontsize=size)
        xp = x * K
        if align == "c":
            xp += (w * K - tl) / 2
        elif align == "r":
            xp += w * K - tl
        yp = y * K + size * 0.33
        self.page.insert_text((xp, yp), s, fontname="min", fontsize=size, color=color)

    def rect(self, x, y, w, h, fill=None, stroke=LINE, width=0.5):
        r = fitz.Rect(x * K, y * K, (x + w) * K, (y + h) * K)
        self.page.draw_rect(r, color=stroke, fill=fill, width=width if stroke else 0)

    def hline(self, x1, x2, y, width=0.8, color=BLACK):
        self.page.draw_line((x1 * K, y * K), (x2 * K, y * K), color=color, width=width)

    def cell(self, x, y, w, h, s="", size=8.0, align="l", fill=None, color=BLACK, pad=3, stroke=LINE):
        self.rect(x, y, w, h, fill=fill, stroke=stroke)
        if s != "":
            if align == "l":
                self.text(x + pad, y + h / 2, s, size, "l", color)
            else:
                self.text(x + pad, y + h / 2, s, size, align, color, w=w - 2 * pad)


def page1(p: Pdf):
    p.new_page()
    p.text(505, 37, "見積№：", 8.5)
    p.text(545, 37, QUOTE_NO, 8.5, "c", w=90)
    p.hline(543, 637, 44, 0.6)
    p.text(505, 61, "作成日：", 8.5)
    p.text(545, 61, DOC_DATE, 8.5, "c", w=90)
    p.hline(543, 637, 68, 0.6)

    p.text(255, 73, "御見積書", 17, "c", w=135)
    p.hline(255, 390, 88, 0.9)

    p.text(22, 115, CLIENT, 11)
    p.hline(22, 333, 124, 0.8)
    p.text(340, 118, "御中", 8.5)
    p.text(22, 138, CONTACT, 11, "c", w=311)
    p.hline(22, 333, 147, 0.8)
    p.text(340, 140, "様", 8.5)

    for i, s in enumerate(ISSUER):
        p.text(435, 169 + i * 17.7, s, 8.5)

    # 押印欄
    x0, y0, cw, ch = 433, 258, 202 / 3, 62
    for i, s in enumerate(STAMPS):
        p.cell(x0 + i * cw, y0, cw, ch, s, 10, "c", stroke=BLACK)

    def field(y, label, value, right=False, ulx2=410):
        p.text(22, y, label, 8.5)
        if right:
            p.text(112, y, value, 10, "r", w=ulx2 - 112)
        else:
            p.text(112, y, value, 9)
        p.hline(112, ulx2, y + 8, 0.6)

    field(181, "システム名", SYSTEM)
    field(217, "プロジェクト名", PROJECT)
    field(235, "発注枝番", BRANCH)
    field(270, "発注内容", SUBJECT)
    field(287, "PJ形態", PJ_FORM)
    field(303, "見積金額", f"{fmt_int(TOTAL_AMT)} 円（税抜き）", right=True)
    field(321, "支払条件", PAYMENT)
    field(342, "見積有効期限", VALID)

    # タスク表
    cols = [(22, 45, "No."), (67, 366, "タスク"), (433, 67, "単価"), (500, 66, "工数(人日)"), (566, 67, "金額(円)")]
    y = 360
    rh = 17.7
    for x, w, t in cols:
        p.cell(x, y, w, 16, t, 8, "c", fill=NAVY, color=WHITE, stroke=NAVY)
    y += 17
    for i, (name, unit, days) in enumerate(TASKS):
        p.cell(22, y, 45, rh, str(i + 1), 7.5, "r", pad=4)
        p.cell(67, y, 366, rh, name, 7.5)
        p.cell(433, y, 67, rh, fmt_int(unit) if unit else "", 7.5, "r", pad=4)
        p.cell(500, y, 66, rh, fmt_d(days) if days is not None else "", 7.5, "r", pad=4)
        p.cell(566, y, 67, rh, fmt_int(AMT[i]) if AMT[i] else "", 7.5, "r", pad=4)
        y += rh
    # 小計
    p.cell(280, y, 220, 16, "小計", 8.5, "l", fill=NAVY, color=WHITE, stroke=NAVY, pad=4)
    p.cell(500, y, 66, 16, fmt_d(TOTAL_DAYS), 7.5, "r", pad=4)
    p.cell(566, y, 67, 16, fmt_int(TOTAL_AMT), 7.5, "r", pad=4)
    y += 16
    p.cell(280, y, 57, 18, "契約形態", 8, "l", fill=NAVY, color=WHITE, stroke=NAVY, pad=4)
    p.cell(337, y, 53, 18, CONTRACT, 7.5, "c")
    p.cell(390, y, 50, 18, "消費税", 8, "l", fill=NAVY, color=WHITE, stroke=NAVY, pad=4)
    p.cell(440, y, 60, 18, TAX_RATE, 7.5, "c")
    p.cell(500, y, 66, 18, "", 7.5)
    p.cell(566, y, 67, 18, "0", 7.5, "r", pad=4)
    y += 18
    p.cell(280, y, 220, 18, "見積合計", 8.5, "l", fill=NAVY, color=WHITE, stroke=NAVY, pad=4)
    p.cell(500, y, 66, 18, "", 7.5)
    p.cell(566, y, 67, 18, fmt_int(TOTAL_AMT), 7.5, "r", pad=4)
    y += 18 + 14

    # 納品条件
    dcols = [(22, 89, "納品回数"), (111, 322, "納品条件"), (433, 67, "納品月"),
             (500, 66, "検収月"), (566, 67, "金額(税抜)")]
    for x, w, t in dcols:
        p.cell(x, y, w, 18, t, 8, "c", fill=NAVY, color=WHITE, stroke=NAVY)
    y += 19
    for no, cond, dm, am, amt in DELIVERY:
        p.cell(22, y, 89, 17.5, no, 7.5, "c")
        p.cell(111, y, 322, 17.5, cond, 7.5)
        p.cell(433, y, 67, 17.5, dm, 7.5, "c")
        p.cell(500, y, 66, 17.5, am, 7.5, "c")
        p.cell(566, y, 67, 17.5, fmt_int(amt), 7.5, "r", pad=4)
        y += 17.5
    y += 8
    p.text(480, y, "検収金額合計：", 7.5, "r", w=80)
    p.text(566, y, fmt_int(sum(d[4] for d in DELIVERY)), 7.5, "r", w=64)
    y += 16

    # 備考
    p.cell(22, y, 43, 19, "備考", 8.5, "c", fill=NAVY, color=WHITE, stroke=NAVY)
    note_h = 12.6 * len(NOTES) * 1.45 + 6
    p.rect(65, y, 568, note_h, stroke=LINE)
    ny = y + 9
    for n in NOTES:
        # 長文は2行に折り返す（等幅近似でなく実測）
        lines = wrap(p.font, n, 7.0, 556 * K)
        for ln in lines:
            p.text(69, ny, ln, 7.0)
            ny += 9.4
        ny += 1.5
    return y + note_h


def wrap(font, s, size, maxw):
    out, cur = [], ""
    for ch in s:
        if font.text_length(cur + ch, fontsize=size) > maxw:
            out.append(cur)
            cur = "　" + ch if False else ch
        else:
            cur += ch
    if cur:
        out.append(cur)
    return out


def page_title(p, title):
    p.new_page()
    p.text(22, 44, "◆" + title, 15)
    p.text(22, 89, f"{QUOTE_NO}:{PAGE_SUB}", 8)


def page2(p: Pdf):
    page_title(p, "作業タスク別見積")
    x, y = 22, 100
    p.cell(x, y, 188, 42, "タスク", 8, "c", fill=NAVY, color=WHITE, stroke=NAVY)
    p.cell(210, y, 429, 20, "CHN", 8, "c", fill=NAVY, color=WHITE, stroke=(1, 1, 1))
    heads = [("役割", 210, 111), ("単価", 321, 111), ("工数", 432, 111), ("金額", 543, 96)]
    for t, hx, hw in heads:
        p.cell(hx, y + 20, hw, 22, t, 8, "c", fill=NAVY, color=WHITE, stroke=(1, 1, 1))
    y += 43
    rh = 22.2
    for i, (name, unit, days) in enumerate(TASKS):
        p.cell(22, y, 21, rh, str(i + 1), 7.5, "r", fill=GRAY, pad=4)
        p.cell(43, y, 167, rh, name, 7.5, fill=GRAY)
        p.cell(210, y, 111, rh, "SE" if unit else "", 7.5)
        p.cell(321, y, 111, rh, fmt_int(unit) if unit else "", 7.5, "r", pad=4)
        p.cell(432, y, 111, rh, fmt_d(days) if days is not None else "", 7.5, "r", pad=4)
        p.cell(543, y, 96, rh, fmt_int(AMT[i]) if AMT[i] else "", 7.5, "r", pad=4)
        y += rh
    p.cell(22, y, 188, rh, "合計", 7.5, "c", fill=LBLUE)
    p.cell(210, y, 111, rh, "", fill=LBLUE)
    p.cell(321, y, 111, rh, "", fill=LBLUE)
    p.cell(432, y, 111, rh, fmt_d(TOTAL_DAYS), 7.5, "r", fill=LBLUE, pad=4)
    p.cell(543, y, 96, rh, fmt_int(TOTAL_AMT), 7.5, "r", fill=LBLUE, pad=4)


def page3(p: Pdf):
    page_title(p, "機能別見積")
    cols = [("NO", 22, 22), ("機能", 44, 78), ("IF名称（仮採番）", 122, 150), ("区分", 272, 30),
            ("新規/\n既存", 302, 34), ("調査・\nMapping", 336, 42), ("詳細\n設計", 378, 38),
            ("開発", 416, 38), ("単体\nテスト", 454, 38), ("内訳\nなし", 492, 38),
            ("合計", 530, 109)]
    y = 100
    for t, x, w in cols:
        p.cell(x, y, w, 32, "", fill=NAVY, stroke=(1, 1, 1))
        ls = t.split("\n")
        for j, ln in enumerate(ls):
            yy = y + 16 + (j - (len(ls) - 1) / 2) * 11
            p.text(x, yy, ln, 7.5, "c", WHITE, w=w)
    y += 33
    rh = 21
    tot = [0.0] * 6
    for i, r in enumerate(FUNC_ROWS):
        grp, name, kind, new, mp, dd, dev, ut, oth = r
        total = mp + dd + dev + ut + oth
        vals = [mp, dd, dev, ut, oth, total]
        p.cell(22, y, 22, rh, str(i + 1), 7, "r", pad=3)
        p.cell(44, y, 78, rh, grp, 7, fill=GRAY, stroke=LINE)
        p.cell(122, y, 150, rh, name, 7)
        p.cell(272, y, 30, rh, kind, 7, "c")
        p.cell(302, y, 34, rh, new, 7, "c")
        for (t, x, w), v in zip(cols[5:], vals):
            p.cell(x, y, w, rh, fmt_d(v) if v else ("0.0" if t != "合計" else fmt_d(v)), 7, "r",
                   fill=LBLUE if t == "合計" else None, pad=4)
        for k, v in enumerate(vals):
            tot[k] += v
        y += rh
    p.cell(272, y, 64, rh, "小計", 7.5, "c", fill=LBLUE)
    for (t, x, w), v in zip(cols[5:], tot):
        p.cell(x, y, w, rh, fmt_d(v), 7.5, "r", fill=LBLUE, pad=4)
    y += rh + 8
    p.cell(272, y, 64, rh, "PJ管理", 7.5, "c", fill=LBLUE)
    p.cell(336, y, 194, rh, "", fill=None)
    p.cell(530, y, 109, rh, fmt_d(PJ_DAYS), 7.5, "r", fill=LBLUE, pad=4)
    y += rh
    p.cell(272, y, 64, rh, "合計", 7.5, "c", fill=LBLUE)
    p.cell(336, y, 194, rh, "", fill=None)
    p.cell(530, y, 109, rh, fmt_d(tot[5] + PJ_DAYS), 7.5, "r", fill=LBLUE, pad=4)
    y += rh + 14
    notes = [
        "※ 詳細設計・開発・単体テストは、各IFの Layer①（IF本体）と Layer③（取込側）を合算。",
        "※ 内訳なし＝非生鮮Phase2への波及対応（伝票区分・振替の注記反映）。内訳は設計時に確定。",
        "※ 調査・Mapping は高5／中4／低3の難易度で積算。流用確認・残件は調査・Mapping 区分に計上。",
    ]
    for n in notes:
        p.text(22, y, n, 7)
        y += 12


def page4(p: Pdf):
    page_title(p, "工程計画")
    y = 100
    heads = [("タスク", 22, 143), ("工数", 165, 58)] + [(m, 223 + i * 35, 35) for i, m in enumerate(MONTHS)]
    for t, x, w in heads:
        p.cell(x, y, w, 22, t, 8, "c", fill=NAVY, color=WHITE, stroke=(1, 1, 1))
    for j in range(len(MONTHS), 12):
        p.cell(223 + j * 35, y, 35, 22, "", fill=NAVY, stroke=(1, 1, 1))
    y += 23
    rh = 22.2
    alloc = [MONTH_ALLOC["map"], MONTH_ALLOC["p2"], MONTH_ALLOC["pj"]]
    for i, (name, unit, days) in enumerate(TASKS):
        p.cell(22, y, 21, rh, str(i + 1), 7.5, "r", fill=GRAY, pad=4)
        p.cell(43, y, 122, rh, name, 7, fill=GRAY)
        p.cell(165, y, 58, rh, fmt_d(days) if days is not None else "", 7.5, "r", pad=4)
        for j in range(12):
            v = alloc[i][j] if i < 3 and j < len(MONTHS) else None
            p.cell(223 + j * 35, y, 35, rh, fmt_d(v) if v else "", 7.5, "r", pad=3)
        y += rh
    p.cell(22, y, 143, rh, "合計", 7.5, "c", fill=LBLUE)
    p.cell(165, y, 58, rh, fmt_d(TOTAL_DAYS), 7.5, "r", fill=LBLUE, pad=4)
    for j in range(12):
        v = sum(a[j] for a in alloc) if j < len(MONTHS) else 0
        p.cell(223 + j * 35, y, 35, rh, fmt_d(v) if j < len(MONTHS) else "0", 7.5, "r", fill=LBLUE, pad=3)
    y += rh + 14
    for n in ["※ 体制4名、2026/10/05着手、完了目安2026/12/08（概算）。月別は営業日按分の概算で、着手遅れ・前提変更により変動します。",
              "※ 10/12・11/3・11/23の祝日を控除。"]:
        p.text(22, y, n, 7)
        y += 12


def build_pdf(out: Path):
    p = Pdf()
    page1(p)
    page2(p)
    page3(p)
    page4(p)
    p.doc.subset_fonts()
    p.doc.save(out, deflate=True, garbage=4)
    p.doc.close()


# ---------------------------------------------------------------- Excel
def build_xlsx(out: Path):
    wb = Workbook()
    thin = Side(style="thin", color="999999")
    bd = Border(left=thin, right=thin, top=thin, bottom=thin)
    navy = PatternFill("solid", fgColor="1F4E78")
    lblue = PatternFill("solid", fgColor="E8F1F9")
    gray = PatternFill("solid", fgColor="F2F2F2")
    fw = Font(name="MS PMincho", size=10, color="FFFFFF", bold=True)
    fb = Font(name="MS PMincho", size=10)
    ctr = Alignment(horizontal="center", vertical="center", wrap_text=True)
    rgt = Alignment(horizontal="right", vertical="center")
    lft = Alignment(horizontal="left", vertical="center", wrap_text=True)

    def put(ws, ref, v, font=fb, fill=None, al=lft, fmt=None, border=True):
        c = ws[ref]
        c.value = v
        c.font = font
        if fill:
            c.fill = fill
        c.alignment = al
        if border:
            c.border = bd
        if fmt:
            c.number_format = fmt
        return c

    # --- 表紙
    ws = wb.active
    ws.title = "見積書"
    for col, w in zip("ABCDEFGH", [8, 40, 16, 12, 12, 14, 14, 4]):
        ws.column_dimensions[col].width = w
    put(ws, "F1", "見積№：", border=False, al=rgt)
    put(ws, "G1", QUOTE_NO, border=False, al=ctr)
    put(ws, "F2", "作成日：", border=False, al=rgt)
    put(ws, "G2", DOC_DATE, border=False, al=ctr)
    ws.merge_cells("A4:G4")
    put(ws, "A4", "御見積書", Font(name="MS PMincho", size=18, bold=True), al=ctr, border=False)
    put(ws, "A6", f"{CLIENT}　御中　{CONTACT}　様", Font(name="MS PMincho", size=12), border=False)
    for i, s in enumerate(ISSUER):
        put(ws, f"E{6 + i}", s, border=False)
    info = [("システム名", SYSTEM), ("プロジェクト名", PROJECT), ("発注枝番", BRANCH), ("発注内容", SUBJECT),
            ("PJ形態", PJ_FORM), ("見積金額", None), ("支払条件", PAYMENT), ("見積有効期限", VALID)]
    r0 = 11
    for i, (k, v) in enumerate(info):
        put(ws, f"A{r0 + i}", k, border=False)
        ws.merge_cells(f"B{r0 + i}:C{r0 + i}")
        put(ws, f"B{r0 + i}", v)
    amt_row = r0 + 5
    hdr = 21
    for c, t in zip("ABCDEFG", ["No.", "タスク", "単価", "工数(人日)", "", "金額(円)", ""]):
        put(ws, f"{c}{hdr}", t, fw, navy, ctr)
    ws.merge_cells(f"D{hdr}:E{hdr}")
    ws.merge_cells(f"F{hdr}:G{hdr}")
    for i, (name, unit, days) in enumerate(TASKS):
        r = hdr + 1 + i
        put(ws, f"A{r}", i + 1, al=rgt)
        put(ws, f"B{r}", name)
        put(ws, f"C{r}", unit, al=rgt, fmt="#,##0")
        ws.merge_cells(f"D{r}:E{r}")
        put(ws, f"D{r}", days, al=rgt, fmt="#,##0.0")
        ws.merge_cells(f"F{r}:G{r}")
        put(ws, f"F{r}", f"=IF(D{r}=\"\",\"\",ROUND(C{r}*D{r},0))" if days is not None else None, al=rgt, fmt="#,##0")
    last = hdr + len(TASKS)
    sub = last + 1
    put(ws, f"B{sub}", "小計", fw, navy)
    ws.merge_cells(f"D{sub}:E{sub}")
    put(ws, f"D{sub}", f"=SUM(D{hdr + 1}:D{last})", al=rgt, fmt="#,##0.0")
    ws.merge_cells(f"F{sub}:G{sub}")
    put(ws, f"F{sub}", f"=SUM(F{hdr + 1}:F{last})", al=rgt, fmt="#,##0")
    put(ws, f"B{sub + 1}", f"契約形態　{CONTRACT}　　消費税　{TAX_RATE}", fw, navy)
    ws.merge_cells(f"F{sub + 1}:G{sub + 1}")
    put(ws, f"F{sub + 1}", 0, al=rgt, fmt="#,##0")
    put(ws, f"B{sub + 2}", "見積合計", fw, navy)
    ws.merge_cells(f"F{sub + 2}:G{sub + 2}")
    put(ws, f"F{sub + 2}", f"=F{sub}+F{sub + 1}", al=rgt, fmt="#,##0")
    ws.merge_cells(f"B{amt_row}:C{amt_row}")
    put(ws, f"B{amt_row}", f'=TEXT(F{sub + 2},"#,##0")&" 円（税抜き）"', al=rgt)
    dh = sub + 4
    for c, t in zip("ABCDEFG", ["納品回数", "納品条件", "", "納品月", "検収月", "金額(税抜)", ""]):
        put(ws, f"{c}{dh}", t, fw, navy, ctr)
    ws.merge_cells(f"B{dh}:C{dh}")
    ws.merge_cells(f"F{dh}:G{dh}")
    for i, (no, cond, dm, am, _a) in enumerate(DELIVERY):
        r = dh + 1 + i
        put(ws, f"A{r}", no, al=ctr)
        ws.merge_cells(f"B{r}:C{r}")
        put(ws, f"B{r}", cond)
        put(ws, f"D{r}", dm, al=ctr)
        put(ws, f"E{r}", am, al=ctr)
        ws.merge_cells(f"F{r}:G{r}")
        f = f"=F{hdr + 1}" if i == 0 else f"=F{hdr + 2}+F{hdr + 3}"
        put(ws, f"F{r}", f, al=rgt, fmt="#,##0")
    dt = dh + 1 + len(DELIVERY)
    put(ws, f"E{dt}", "検収金額合計", al=rgt, border=False)
    ws.merge_cells(f"F{dt}:G{dt}")
    put(ws, f"F{dt}", f"=SUM(F{dh + 1}:F{dt - 1})", al=rgt, fmt="#,##0")
    put(ws, f"A{dt + 2}", "備考", fw, navy, ctr)
    for i, n in enumerate(NOTES):
        ws.merge_cells(f"B{dt + 2 + i}:G{dt + 2 + i}")
        put(ws, f"B{dt + 2 + i}", n)
        ws.row_dimensions[dt + 2 + i].height = 40
    for i, s in enumerate(STAMPS):
        put(ws, f"{'EFG'[i]}10", s, Font(name="MS PMincho", size=11), al=ctr)
    ws.row_dimensions[10].height = 36
    # --- 作業タスク別
    w2 = wb.create_sheet("作業タスク別見積")
    for col, w in zip("ABCDEF", [6, 28, 8, 12, 12, 16]):
        w2.column_dimensions[col].width = w
    put(w2, "A1", "◆作業タスク別見積", Font(name="MS PMincho", size=14, bold=True), border=False)
    put(w2, "A2", f"{QUOTE_NO}:{PAGE_SUB}", border=False)
    for c, t in zip("ABCDEF", ["No.", "タスク", "役割", "単価", "工数", "金額"]):
        put(w2, f"{c}4", t, fw, navy, ctr)
    for i, (name, unit, days) in enumerate(TASKS):
        r = 5 + i
        put(w2, f"A{r}", i + 1, fill=gray, al=rgt)
        put(w2, f"B{r}", name, fill=gray)
        put(w2, f"C{r}", "SE" if unit else None)
        put(w2, f"D{r}", f"=見積書!C{hdr + 1 + i}" if unit else None, al=rgt, fmt="#,##0")
        put(w2, f"E{r}", f"=見積書!D{hdr + 1 + i}" if days is not None else None, al=rgt, fmt="#,##0.0")
        put(w2, f"F{r}", f"=見積書!F{hdr + 1 + i}" if days is not None else None, al=rgt, fmt="#,##0")
    e = 5 + len(TASKS)
    put(w2, f"B{e}", "合計", fill=lblue, al=ctr)
    put(w2, f"E{e}", f"=SUM(E5:E{e - 1})", fill=lblue, al=rgt, fmt="#,##0.0")
    put(w2, f"F{e}", f"=SUM(F5:F{e - 1})", fill=lblue, al=rgt, fmt="#,##0")
    # --- 機能別
    w3 = wb.create_sheet("機能別見積")
    widths = [5, 16, 36, 7, 8, 11, 9, 9, 9, 9, 10]
    for i, w in enumerate(widths, 1):
        w3.column_dimensions[get_column_letter(i)].width = w
    put(w3, "A1", "◆機能別見積", Font(name="MS PMincho", size=14, bold=True), border=False)
    put(w3, "A2", f"{QUOTE_NO}:{PAGE_SUB}", border=False)
    heads = ["NO", "機能", "IF名称（仮採番）", "区分", "新規/既存", "調査・Mapping", "詳細設計", "開発", "単体テスト", "内訳なし", "合計"]
    for i, t in enumerate(heads, 1):
        put(w3, f"{get_column_letter(i)}4", t, fw, navy, ctr)
    for i, r in enumerate(FUNC_ROWS):
        rr = 5 + i
        grp, name, kind, new, mp, dd, dev, ut, oth = r
        put(w3, f"A{rr}", i + 1, al=rgt)
        put(w3, f"B{rr}", grp, fill=gray)
        put(w3, f"C{rr}", name)
        put(w3, f"D{rr}", kind, al=ctr)
        put(w3, f"E{rr}", new, al=ctr)
        for j, v in enumerate([mp, dd, dev, ut, oth]):
            put(w3, f"{get_column_letter(6 + j)}{rr}", v, al=rgt, fmt="#,##0.0")
        put(w3, f"K{rr}", f"=SUM(F{rr}:J{rr})", fill=lblue, al=rgt, fmt="#,##0.0")
    e3 = 5 + len(FUNC_ROWS)
    put(w3, f"C{e3}", "小計", fill=lblue, al=ctr)
    for j in range(6, 12):
        cl = get_column_letter(j)
        put(w3, f"{cl}{e3}", f"=SUM({cl}5:{cl}{e3 - 1})", fill=lblue, al=rgt, fmt="#,##0.0")
    put(w3, f"C{e3 + 1}", "PJ管理", fill=lblue, al=ctr)
    put(w3, f"K{e3 + 1}", f"=見積書!D{hdr + 3}", fill=lblue, al=rgt, fmt="#,##0.0")
    put(w3, f"C{e3 + 2}", "合計", fill=lblue, al=ctr)
    put(w3, f"K{e3 + 2}", f"=K{e3}+K{e3 + 1}", fill=lblue, al=rgt, fmt="#,##0.0")
    # --- 工程計画
    w4 = wb.create_sheet("工程計画")
    w4.column_dimensions["A"].width = 6
    w4.column_dimensions["B"].width = 28
    w4.column_dimensions["C"].width = 10
    put(w4, "A1", "◆工程計画", Font(name="MS PMincho", size=14, bold=True), border=False)
    put(w4, "A2", f"{QUOTE_NO}:{PAGE_SUB}", border=False)
    for i, t in enumerate(["No.", "タスク", "工数"] + MONTHS, 1):
        put(w4, f"{get_column_letter(i)}4", t, fw, navy, ctr)
    alloc = [MONTH_ALLOC["map"], MONTH_ALLOC["p2"], MONTH_ALLOC["pj"]]
    for i, (name, unit, days) in enumerate(TASKS):
        rr = 5 + i
        put(w4, f"A{rr}", i + 1, fill=gray, al=rgt)
        put(w4, f"B{rr}", name, fill=gray)
        put(w4, f"C{rr}", f"=見積書!D{hdr + 1 + i}" if days is not None else None, al=rgt, fmt="#,##0.0")
        for j in range(len(MONTHS)):
            v = alloc[i][j] if i < 3 and alloc[i][j] else None
            put(w4, f"{get_column_letter(4 + j)}{rr}", v, al=rgt, fmt="#,##0.0")
    e4 = 5 + len(TASKS)
    put(w4, f"B{e4}", "合計", fill=lblue, al=ctr)
    for j in range(3, 4 + len(MONTHS)):
        cl = get_column_letter(j)
        put(w4, f"{cl}{e4}", f"=SUM({cl}5:{cl}{e4 - 1})", fill=lblue, al=rgt, fmt="#,##0.0")
    put(w4, f"A{e4 + 2}", "※ 体制4名、2026/10/05着手、完了目安2026/12/08（概算）。月別は営業日按分の概算。", border=False)
    for ws_ in wb.worksheets:
        ws_.sheet_view.showGridLines = False
    wb.save(out)


if __name__ == "__main__":
    pdf = HERE / "【御見積書】生鮮IF増分_標準案_20261007.pdf"
    xlsx = HERE / "【御見積書】生鮮IF増分_標準案_20261007.xlsx"
    build_pdf(pdf)
    build_xlsx(xlsx)
    print("OK", pdf.name, xlsx.name, TOTAL_DAYS, TOTAL_AMT)
