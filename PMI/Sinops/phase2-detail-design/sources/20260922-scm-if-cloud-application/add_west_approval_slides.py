#!/usr/bin/env python3
"""Append one compact West差戻し page after slide 1 of the submitted deck."""
from __future__ import annotations

import re
import shutil
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

SRC = Path(
    "/Users/treqd/Downloads/SCM連携クラウド機 GCP環境申請 (3).pptx"
)
OUT_DIR = Path(
    "/Users/treqd/Desktop/Cursor/西友PMI/PMI/Sinops/phase2-detail-design/"
    "sources/20260922-scm-if-cloud-application"
)
OUT = OUT_DIR / "SCM連携クラウド機 GCP環境申請_West差戻し補足.pptx"
DL = Path("/Users/treqd/Downloads/SCM連携クラウド機 GCP環境申請_West差戻し補足.pptx")

INK = RGBColor(0x1E, 0x23, 0x56)
MUTED = RGBColor(0x4A, 0x52, 0x63)
CARD = RGBColor(0xF1, 0xF3, 0xF8)
WHITE = RGBColor(0xFC, 0xFC, 0xFD)
NAVY = RGBColor(0x1E, 0x23, 0x56)
ACCENT = RGBColor(0x2F, 0x75, 0xB5)
FONT = "Noto Sans JP"

L = Emu(812800)
TITLE_T = Emu(457200)
TITLE_H = Emu(513060)
BADGE_L = Emu(8385870)
BADGE_T = Emu(553045)
BADGE_W = Emu(2990155)
BADGE_H = Emu(295870)
BADGE_TXT_L = Emu(8516045)
BADGE_TXT_T = Emu(594320)
BADGE_TXT_W = Emu(2811700)
BADGE_TXT_H = Emu(238720)
SUB_T = Emu(995660)
SUB_H = Emu(274241)
FOOT_L = Emu(10058400)
FOOT_T = Emu(6339880)
FOOT_W = Emu(1320800)
FOOT_H = Emu(238720)
W_CONTENT = Emu(10883392)


def no_line(sh) -> None:
    spPr = sh._element.spPr
    ln = spPr.find(qn("a:ln"))
    if ln is None:
        ln = etree.SubElement(spPr, qn("a:ln"))
    if "w" in ln.attrib:
        del ln.attrib["w"]
    for child in list(ln):
        ln.remove(child)
    etree.SubElement(ln, qn("a:noFill"))


def rr(slide, l, t, w, h, fill):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, int(l), int(t), int(w), int(h))
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    no_line(sh)
    sh.adjustments[0] = 0.08
    return sh


def set_tf(sh, text, size, color, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE):
    tf = sh.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    try:
        sh.text_frame._txBody.bodyPr.set("anchor", {MSO_ANCHOR.MIDDLE: "ctr", MSO_ANCHOR.TOP: "t"}.get(anchor, "t"))
    except Exception:
        pass
    p = tf.paragraphs[0]
    p.clear()
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = etree.SubElement(rPr, qn(tag))
        el.set("typeface", FONT)


def add_text(slide, l, t, w, h, text, size, color, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    sh = slide.shapes.add_textbox(int(l), int(t), int(w), int(h))
    tf = sh.text_frame
    tf.word_wrap = True
    lines = text.split("\n")
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(2)
        run = p.add_run()
        run.text = line
        run.font.name = FONT
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color
        rPr = run._r.get_or_add_rPr()
        for tag in ("a:latin", "a:ea", "a:cs"):
            el = rPr.find(qn(tag))
            if el is None:
                el = etree.SubElement(rPr, qn(tag))
            el.set("typeface", FONT)
    return sh


def style_cell(cell, text, size, color, bold, fill, align=PP_ALIGN.LEFT):
    cell.text = ""
    cell.fill.solid()
    cell.fill.fore_color.rgb = fill
    tf = cell.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = etree.SubElement(rPr, qn(tag))
        el.set("typeface", FONT)


def chrome(slide, title: str, subtitle: str) -> None:
    add_text(slide, L, TITLE_T, Emu(6800000), TITLE_H, title, 22, INK, bold=True, anchor=MSO_ANCHOR.MIDDLE)
    badge = rr(slide, BADGE_L, BADGE_T, BADGE_W, BADGE_H, CARD)
    add_text(
        slide,
        BADGE_TXT_L,
        BADGE_TXT_T,
        BADGE_TXT_W,
        BADGE_TXT_H,
        "東京 asia-northeast1 ・ dev / stg / prod",
        10,
        MUTED,
        align=PP_ALIGN.CENTER,
        anchor=MSO_ANCHOR.MIDDLE,
    )
    rr(slide, L, SUB_T, W_CONTENT, SUB_H, CARD)
    add_text(slide, Emu(990600), Emu(SUB_T + 20000), Emu(10500000), Emu(230000), subtitle, 12, MUTED, bold=False)


def add_footer_placeholder(slide) -> None:
    sh = rr(slide, FOOT_L, FOOT_T, FOOT_W, FOOT_H, WHITE)
    sh.name = "page-num"
    set_tf(sh, "", 11, MUTED, align=PP_ALIGN.RIGHT)


def add_west_slide(prs) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    chrome(
        slide,
        "実施条件とテスト環境の同時進行",
        "実施条件が成立し、dev／stg／prod の接続先を整理できた場合に構築します。本番だけを先行する申請ではありません。",
    )
    # Left: why this environment is required.
    rr(slide, L, Emu(1360000), Emu(5150000), Emu(3740000), CARD)
    add_text(slide, Emu(1016000), Emu(1490000), Emu(4750000), Emu(280000),
             "① 実施が必要となる条件", 15, INK, bold=True)
    add_text(
        slide,
        Emu(1016000), Emu(1840000), Emu(4750000), Emu(2880000),
        "A　大阪の集計データを日次で西友へ渡す\n"
        "B　型変換・共通化は TRIAL 側 Layer①で行う\n"
        "C　最終 TXT のみ exchange へ一方向納品する\n"
        "D　元データは本PJ GCS（landing）に置く\n"
        "E　SCM 連携面は 1 プロジェクトに集約する\n"
        "F　Layer②③ は西友の既存 DataSpider を使う\n\n"
        "この前提なら既存 DSS では Layer①を代替できず、TRIAL GCP の日次バッチ機が必要。未合意なら申請を進めない。",
        12,
        INK,
    )
    # Right: environment alignment and cost behavior.
    rows = [
        ("環境", "TRIAL 本申請", "接続先"),
        ("dev", "md-data-integration-dev", "exchange dev／DSS 開発"),
        ("stg", "md-data-integration-stg", "exchange stg／DSS 検証"),
        ("prod", "md-data-integration-prod", "exchange prod／本番 DSS"),
    ]
    add_text(slide, Emu(6260000), Emu(1490000), Emu(5200000), Emu(280000),
             "② テスト環境と同時進行", 15, INK, bold=True)
    tbl_shape = slide.shapes.add_table(
        len(rows), 3, Emu(6162675), Emu(1840000), Emu(5213350), Emu(1760000)
    )
    tbl = tbl_shape.table
    widths = [Emu(800000), Emu(2400000), Emu(2013350)]
    for i, w in enumerate(widths):
        tbl.columns[i].width = w
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            cell = tbl.cell(i, j)
            if i == 0:
                style_cell(cell, val, 11, WHITE, True, NAVY)
            else:
                fill = WHITE if i % 2 else CARD
                style_cell(cell, val, 10, INK, j == 0, fill)
    add_text(
        slide, Emu(6340000), Emu(3740000), Emu(4850000), Emu(920000),
        "進め方：接続先・パス・権限を揃え、dev → stg で起動〜_SUCCESS を確認。prod の定期実行はテスト経路整理後。",
        11, MUTED,
    )
    rr(slide, Emu(6162675), Emu(4720000), Emu(5213350), Emu(1380000), CARD)
    add_text(
        slide, Emu(6340000), Emu(4840000), Emu(4850000), Emu(1160000),
        "費用の考え方\n"
        "dev／stg の Cloud Run Jobs は常時起動せず、テスト実施時だけ計算費が発生します。GCS 保管・データ転送は、保持量・転送量に応じて発生します。",
        11, INK, bold=False,
    )
    add_footer_placeholder(slide)


def reorder_after_first(prs) -> None:
    sldIdLst = prs.slides._sldIdLst
    children = list(sldIdLst)
    new_slide = children[-1]
    sldIdLst.remove(new_slide)
    first = list(sldIdLst)[0]
    first.addnext(new_slide)


def annotate_cost_slide(prs) -> None:
    """Clarify that dev/stg compute is not always-on; storage/transfer remains usage based."""
    for slide in prs.slides:
        for sh in slide.shapes:
            if not sh.has_text_frame:
                continue
            text = sh.text_frame.text
            if "稼働：prod 3h/日" not in text:
                continue
            replacement = text.replace(
                "稼働：prod 3h/日（90h/月）、stg 20h/月、dev 2.5h/月。",
                "稼働：prod 90h/月、stg 20h/月、dev 2.5h/月。\n"
                "dev／stg は常時起動せず、テスト実施時のみ計算費が発生。",
            )
            sh.text_frame.text = replacement
            for p in sh.text_frame.paragraphs:
                for run in p.runs:
                    run.font.name = FONT
                    run.font.size = Pt(9.5)
                    run.font.color.rgb = MUTED
            return


def renumber(prs) -> None:
    n = len(prs.slides)
    pat = re.compile(r"^\s*\d+\s*/\s*\d+\s*$")
    for i, slide in enumerate(prs.slides, 1):
        for sh in slide.shapes:
            if not sh.has_text_frame:
                continue
            raw = sh.text_frame.text.strip()
            if sh.top > Emu(6000000) and pat.match(raw):
                tf = sh.text_frame
                p = tf.paragraphs[0]
                if p.runs:
                    p.runs[0].text = f"{i} / {n}"
                    for r in p.runs[1:]:
                        r.text = ""
                else:
                    set_tf(sh, f"{i} / {n}", 11, MUTED, align=PP_ALIGN.RIGHT)
            elif getattr(sh, "name", "") == "page-num":
                set_tf(sh, f"{i} / {n}", 11, MUTED, align=PP_ALIGN.RIGHT)


def main() -> None:
    if not SRC.exists():
        raise SystemExit(f"missing {SRC}")
    shutil.copy2(SRC, OUT)
    prs = Presentation(str(OUT))
    annotate_cost_slide(prs)
    add_west_slide(prs)
    reorder_after_first(prs)
    renumber(prs)
    prs.save(str(OUT))
    shutil.copy2(OUT, DL)
    print("saved", OUT, "slides", len(Presentation(str(OUT)).slides))
    print("copied", DL)


if __name__ == "__main__":
    main()
