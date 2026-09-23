#!/usr/bin/env python3
"""Rebuild slide 3 to match the rest of the deck's visual language."""
from pathlib import Path
import shutil

from pptx import Presentation
from pptx.util import Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN

PPTX = Path(
    "/Users/treqd/Desktop/Cursor/西友PMI/PMI/Sinops/phase2-detail-design/"
    "sources/20260922-scm-if-cloud-application/SCM連携クラウド機_GCP環境申請.pptx"
)

# Template palette (same as other slides)
BLUE = RGBColor(0x2F, 0x75, 0xB5)
BLUE_BG = RGBColor(0xEA, 0xF2, 0xF8)
BLUE_BG2 = RGBColor(0xE2, 0xF0, 0xF8)
BLUE_AREA = RGBColor(0xF5, 0xF9, 0xFD)
GOLD = RGBColor(0xC9, 0xA2, 0x27)
GOLD_BG = RGBColor(0xFF, 0xF2, 0xCC)
ORANGE = RGBColor(0xC6, 0x59, 0x11)
ORANGE_BG = RGBColor(0xFC, 0xE4, 0xD6)
INK = RGBColor(0x17, 0x36, 0x5D)
MUTED = RGBColor(0x5B, 0x65, 0x73)
RED = RGBColor(0xC0, 0x00, 0x00)
RED_BG = RGBColor(0xFF, 0xF7, 0xF7)
PURPLE = RGBColor(0x70, 0x30, 0xA0)
PURPLE_BG = RGBColor(0xE4, 0xDF, 0xEC)
RED_LINE = RGBColor(0xD9, 0x96, 0x94)
LW = Pt(1.2)
FONT = "Arial"

KEEP = {
    "Shape 0",
    "Text 1",
    "Text 2",
    "Text 3",
    "Shape 4",
    "Text 5",
    "Shape 6",
    "Text 7",
    "Shape 8",
    "Text 9",
    "Shape 12",
    "Text 13",
    "Shape 51",
    "Text 52",
}


def main() -> None:
    prs = Presentation(str(PPTX))
    slide = prs.slides[2]
    sp_tree = slide.shapes._spTree

    for shape in list(slide.shapes):
        if shape.name not in KEEP:
            sp_tree.remove(shape._element)

    def rr(l, t, w, h, fill, line=None, lw=None):
        """Filled rounded rect with NO stroke — matches slides 1/2/5–7."""
        sh = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, int(l), int(t), int(w), int(h)
        )
        sh.fill.solid()
        sh.fill.fore_color.rgb = fill
        # no visible border (deck standard)
        from pptx.oxml.ns import qn
        from lxml import etree

        spPr = sh._element.spPr
        ln = spPr.find(qn("a:ln"))
        if ln is None:
            ln = etree.SubElement(spPr, qn("a:ln"))
        if "w" in ln.attrib:
            del ln.attrib["w"]
        for child in list(ln):
            ln.remove(child)
        etree.SubElement(ln, qn("a:noFill"))
        try:
            sh.adjustments[0] = 0.05
        except Exception:
            pass
        return sh

    def textbox(l, t, w, h, text, size=10, bold=False, color=MUTED, align=PP_ALIGN.LEFT):
        box = slide.shapes.add_textbox(int(l), int(t), int(w), int(h))
        p = box.text_frame.paragraphs[0]
        p.alignment = align
        r = p.add_run()
        r.text = text
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = color
        r.font.name = FONT
        return box

    def fill_text(shape, lines, align=PP_ALIGN.CENTER, anchor="ctr"):
        tf = shape.text_frame
        tf.clear()
        tf.word_wrap = True
        tf.margin_left = Pt(8)
        tf.margin_right = Pt(8)
        tf.margin_top = Pt(6)
        tf.margin_bottom = Pt(6)
        try:
            tf._txBody.bodyPr.set("anchor", anchor)
        except Exception:
            pass
        for i, (text, size, bold, color) in enumerate(lines):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = align
            p.space_before = Pt(0)
            p.space_after = Pt(1)
            r = p.add_run()
            r.text = text
            r.font.size = Pt(size)
            r.font.bold = bold
            r.font.color.rgb = color
            r.font.name = FONT

    # --- same structural grid, template look ---
    col_y = 1420000
    col_h = 3000000
    gap = 100000
    arrow_w = 360000

    left_x = 400000
    left_w = 2100000
    center_x = left_x + left_w + gap + arrow_w
    center_w = 6100000
    right_x = center_x + center_w + gap + arrow_w
    right_w = 2350000

    # Section labels — plain, like other slides (no ①②③)
    textbox(left_x, 1240000, left_w + 200000, 150000, "既存・外部（今回申請しない）", 10, True, MUTED)
    textbox(center_x, 1240000, center_w, 150000, "GCP　東京（今回申請）", 10, True, BLUE)
    textbox(right_x, 1240000, right_w, 150000, "既存・相手先（バケット利用）", 10, True, MUTED)

    left = rr(left_x, col_y, left_w, col_h, ORANGE_BG, ORANGE)
    fill_text(
        left,
        [
            ("大阪　集計 GCS", 13, True, INK),
            ("asia-northeast2", 10, False, MUTED),
            ("", 6, False, MUTED),
            ("日次データ + _READY", 11, False, INK),
        ],
    )

    rr(center_x, col_y, center_w, col_h, BLUE_AREA, BLUE)

    pad = 140000
    ig = 120000
    avail_w = center_w - 2 * pad
    avail_h = col_h - 2 * pad
    card_w = (avail_w - ig) // 2
    card_h = (avail_h - ig) // 2
    ix0 = center_x + pad
    iy0 = col_y + pad

    # PJ banner inside area
    banner = rr(ix0, iy0 - 20000, avail_w, 200000, BLUE_BG2, BLUE, Pt(1.0))
    fill_text(
        banner,
        [("本PJ GCP　asia-northeast1　／　暫定PJ ID：md-data-integration-prod", 10, True, BLUE)],
    )
    iy0 = iy0 + 220000
    card_h = (col_y + col_h - pad - iy0 - ig) // 2

    c11 = rr(ix0, iy0, card_w, card_h, BLUE_BG, BLUE)
    fill_text(
        c11,
        [
            ("Storage Transfer Service", 12, True, INK),
            ("manifest / prefix", 10, False, MUTED),
        ],
    )
    c12 = rr(ix0 + card_w + ig, iy0, card_w, card_h, BLUE_BG2, BLUE)
    fill_text(
        c12,
        [
            ("本PJ Cloud Storage", 12, True, INK),
            ("source-landing/YYYYMMDD/", 10, False, MUTED),
            ("元データ（7日保管）", 10, False, MUTED),
        ],
    )
    c21 = rr(ix0, iy0 + card_h + ig, card_w, card_h, BLUE_BG, BLUE)
    fill_text(
        c21,
        [
            ("Cloud Run Job", 12, True, INK),
            ("SMART 日次処理", 10, False, MUTED),
            ("4 vCPU / 16 GiB", 10, False, MUTED),
        ],
    )
    c22 = rr(ix0 + card_w + ig, iy0 + card_h + ig, card_w, card_h, BLUE_BG2, BLUE)
    fill_text(
        c22,
        [
            ("大容量：resumable", 11, True, INK),
            ("通常：直接アップロード", 10, False, MUTED),
        ],
    )

    right = rr(right_x, col_y, right_w, col_h, ORANGE_BG, ORANGE)
    fill_text(
        right,
        [
            ("seiyu-trial-data-exchange", 11, True, INK),
            ("", 6, False, MUTED),
            ("東京 asia-northeast1", 10, False, MUTED),
            ("result/YYYYMMDD/final.txt", 10, False, MUTED),
            ("_SUCCESS", 11, False, INK),
        ],
    )

    # Arrows — Arial, template ink colors
    ay = col_y + col_h // 2 - 80000
    textbox(left_x + left_w, ay - 110000, arrow_w + gap, 90000, "読み取り", 8, True, ORANGE, PP_ALIGN.CENTER)
    textbox(left_x + left_w, ay, arrow_w + gap, 180000, "→", 18, True, ORANGE, PP_ALIGN.CENTER)
    textbox(center_x + center_w, ay - 110000, arrow_w + gap, 90000, "納品", 8, True, ORANGE, PP_ALIGN.CENTER)
    textbox(center_x + center_w, ay, arrow_w + gap, 180000, "→", 18, True, ORANGE, PP_ALIGN.CENTER)

    textbox(ix0 + card_w, iy0 + card_h // 2 - 70000, ig, 140000, "→", 14, True, BLUE, PP_ALIGN.CENTER)
    textbox(ix0 + card_w + card_w // 2 - 50000, iy0 + card_h - 10000, 140000, 130000, "↓", 12, True, BLUE, PP_ALIGN.CENTER)
    textbox(ix0 + card_w, iy0 + card_h + ig + card_h // 2 - 70000, ig, 140000, "→", 14, True, BLUE, PP_ALIGN.CENTER)

    # Control — gold like other slides
    ctl_y = col_y + col_h + 110000
    textbox(left_x, ctl_y + 50000, 900000, 130000, "制御フロー", 10, True, GOLD)
    sch = rr(left_x + 1000000, ctl_y, 2300000, 430000, GOLD_BG, GOLD)
    fill_text(
        sch,
        [
            ("Cloud Scheduler", 12, True, INK),
            ("毎日 指定時刻", 10, False, MUTED),
        ],
    )
    textbox(left_x + 3350000, ctl_y + 70000, 320000, 220000, "→", 16, True, GOLD, PP_ALIGN.CENTER)
    wf_x = left_x + 3700000
    wf_w = right_x + right_w - wf_x
    wf = rr(wf_x, ctl_y, wf_w, 430000, GOLD_BG, GOLD)
    fill_text(
        wf,
        [
            ("Workflows", 12, True, INK),
            ("READY → STS → Job → 検証", 10, False, MUTED),
        ],
    )

    # Notes — same note colors as S4
    note_y = ctl_y + 510000
    total_w = right_x + right_w - left_x
    note_gap = 120000
    note_w = (total_w - note_gap) // 2
    note_h = 470000
    n1 = rr(left_x, note_y, note_w, note_h, RED_BG, RED_LINE)
    fill_text(
        n1,
        [
            ("安全境界", 11, True, RED),
            ("元データは seiyu-trial-data-exchange に入らない。", 10, False, INK),
            ("西友側は objectCreator のみ。", 10, False, INK),
        ],
        align=PP_ALIGN.LEFT,
        anchor="t",
    )
    n2 = rr(left_x + note_w + note_gap, note_y, note_w, note_h, PURPLE_BG, PURPLE)
    fill_text(
        n2,
        [
            ("アップロード方針", 11, True, PURPLE),
            ("大容量：ストリーム + resumable upload。", 10, False, INK),
            ("通常：Cloud Run 処理後、seiyu-trial-data-exchange へ直接アップロード。", 10, False, INK),
        ],
        align=PP_ALIGN.LEFT,
        anchor="t",
    )

    assert note_y + note_h < 6446520 - 40000
    prs.save(str(PPTX))
    print("slide3 rebuilt, shapes=", len(list(prs.slides[2].shapes)))


def unify_greens(prs: Presentation) -> None:
    """Replace leftover 予測フロー green with legend-aligned colors."""
    BLUE = RGBColor(0x2F, 0x75, 0xB5)
    BLUE_BG = RGBColor(0xEA, 0xF2, 0xF8)
    ORANGE = RGBColor(0xC6, 0x59, 0x11)
    ORANGE_BG = RGBColor(0xFC, 0xE4, 0xD6)
    GOLD_BG = RGBColor(0xFF, 0xF2, 0xCC)
    GOLD = RGBColor(0xC9, 0xA2, 0x27)

    # (slide_idx, match_nearby_text_substr) -> (fill, line)
    # Default green content boxes → blue (今回申請), seiyu/partner → orange
    for si, slide in enumerate(prs.slides):
        for shape in slide.shapes:
            try:
                fr = str(shape.fill.fore_color.rgb)
            except Exception:
                continue
            if fr != "E2F0D9":
                continue
            # find nearby label
            label = ""
            for s in slide.shapes:
                if not s.has_text_frame or not s.text_frame.text.strip():
                    continue
                if s.left is None or shape.left is None:
                    continue
                if abs(s.left - shape.left) < 400000 and abs((s.top or 0) - (shape.top or 0)) < 400000:
                    label = s.text_frame.text
                    break
            if not label and shape.has_text_frame:
                label = shape.text_frame.text

            # seiyu / partner / external → orange
            if "seiyu-trial" in label or "相手先" in label or "西友" in label:
                fill, line = ORANGE_BG, ORANGE
                why = "orange/seiyu"
            # table zebra / conclusion / highlight on content pages → blue
            else:
                fill, line = BLUE_BG, BLUE
                why = "blue/申請"
            shape.fill.solid()
            shape.fill.fore_color.rgb = fill
            try:
                shape.line.color.rgb = line
                shape.line.width = Pt(1.2)
            except Exception:
                pass
            print(f"  S{si+1} green→{why}: {label.split(chr(10))[0][:50]!r}")


if __name__ == "__main__":
    main()
    prs = Presentation(str(PPTX))
    print("unify greens:")
    unify_greens(prs)
    prs.save(str(PPTX))
    shutil.copy2(PPTX, "/Users/treqd/Downloads/SCM連携クラウド機_GCP環境申請.pptx")
    # verify no green
    prs = Presentation(str(PPTX))
    left = 0
    for si, slide in enumerate(prs.slides, 1):
        for shape in slide.shapes:
            try:
                if str(shape.fill.fore_color.rgb) == "E2F0D9":
                    left += 1
                    print("LEFT", si, shape.name)
            except Exception:
                pass
    print("green leftovers", left)

    # font check S3
    from collections import Counter
    fonts = Counter()
    for shape in prs.slides[2].shapes:
        if not shape.has_text_frame:
            continue
        # only new content (skip chrome which may be Mincho)
        if shape.name in KEEP:
            continue
        for p in shape.text_frame.paragraphs:
            for r in p.runs:
                fonts[r.font.name] += 1
    print("S3 content fonts", fonts)
