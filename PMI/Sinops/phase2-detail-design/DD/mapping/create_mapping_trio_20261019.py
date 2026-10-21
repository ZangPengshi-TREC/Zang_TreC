# -*- coding: utf-8 -*-
"""DD の IF詳細設計書（2026-10-17）から、IF ごとに
①項目マッピング表 / ②GAP分析書 / ③変換ルール定義書 を作る。
"""
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from pathlib import Path

OUT = Path(__file__).resolve().parent
DATE = "2026-10-19"
SRC = "根拠: phase2-detail-design/DD の IF詳細設計書（2026-10-17 まで）"

NAVY = "1F4E79"
THIN = Border(
    left=Side(style="thin", color="BFBFBF"),
    right=Side(style="thin", color="BFBFBF"),
    top=Side(style="thin", color="BFBFBF"),
    bottom=Side(style="thin", color="BFBFBF"),
)
FILL_HEAD = PatternFill("solid", fgColor=NAVY)
FILL_TITLE = PatternFill("solid", fgColor="DDEBF7")
FILL_OPEN = PatternFill("solid", fgColor="FFF2CC")
FILL_OK = PatternFill("solid", fgColor="E2EFDA")
FILL_PART = PatternFill("solid", fgColor="FCE4D6")
FILL_NOTE = PatternFill("solid", fgColor="F2F2F2")

FONT_TITLE = Font(name="Arial", size=14, bold=True, color=NAVY)
FONT_SUB = Font(name="Arial", size=9, color="595959")
FONT_HEAD = Font(name="Arial", size=9, bold=True, color="FFFFFF")
FONT_CELL = Font(name="Arial", size=9)
FONT_NOTE = Font(name="Arial", size=9, color="595959")

WRAP = Alignment(wrap_text=True, vertical="center")
LEFT = Alignment(wrap_text=True, vertical="center", horizontal="left")


def status_fill(text):
    t = str(text)
    if any(k in t for k in ("未", "確認中", "要確認")):
        return FILL_OPEN
    if any(k in t for k in ("一部", "方針")):
        return FILL_PART
    if any(k in t for k in ("確定", "完了")):
        return FILL_OK
    return None


def write_sheet(ws, title, notes, headers, rows, widths, footers=None):
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.oddHeader.left.text = title
    ws.oddFooter.right.text = DATE
    ws.freeze_panes = "A5"
    last = 4 + len(rows)
    ws.auto_filter.ref = f"A4:{get_column_letter(len(headers))}{last}"
    ws.print_title_rows = "1:4"
    ws.page_setup.horizontalCentered = True
    ws.sheet_view.zoomScale = 110

    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
    c = ws.cell(1, 1, title)
    c.font = FONT_TITLE
    c.alignment = LEFT
    ws.row_dimensions[1].height = 24

    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(headers))
    c = ws.cell(2, 1, "  /  ".join(notes))
    c.font = FONT_SUB
    c.alignment = LEFT
    ws.row_dimensions[2].height = 32

    legend = "判定:  ○ 直連携    △ 変換・編集    ― 固定値または送らない    ✕ 未確定（GAP）"
    ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=len(headers))
    c = ws.cell(3, 1, legend)
    c.font = FONT_NOTE
    c.alignment = LEFT
    ws.row_dimensions[3].height = 18

    for col, h in enumerate(headers, 1):
        cell = ws.cell(4, col, h)
        cell.font = FONT_HEAD
        cell.fill = FILL_HEAD
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        cell.border = THIN
    ws.row_dimensions[4].height = 30

    for r_i, row in enumerate(rows, 5):
        ws.row_dimensions[r_i].height = 36
        for c_i, val in enumerate(row, 1):
            cell = ws.cell(r_i, c_i, val)
            cell.font = FONT_CELL
            cell.alignment = WRAP
            cell.border = THIN
            if c_i == len(headers) or (headers[c_i - 1] in ("確定", "ステータス", "確定度", "判定")):
                fill = status_fill(val)
                if fill:
                    cell.fill = fill
        # long rule rows
        longest = max((len(str(v)) for v in row), default=0)
        if longest > 80:
            ws.row_dimensions[r_i].height = 52
        if longest > 140:
            ws.row_dimensions[r_i].height = 68

    if footers:
        start = 6 + len(rows)
        for i, line in enumerate(footers):
            rr = start + i
            ws.merge_cells(start_row=rr, start_column=1, end_row=rr, end_column=len(headers))
            cell = ws.cell(rr, 1, line)
            cell.font = FONT_NOTE
            cell.alignment = LEFT
            cell.fill = FILL_NOTE
            ws.row_dimensions[rr].height = 18

    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[2].height = 36
    ws.sheet_view.view = "pageBreakPreview"
    ws.page_margins.left = 0.4
    ws.page_margins.right = 0.4
    ws.page_margins.top = 0.6
    ws.page_margins.bottom = 0.5
    ws.oddHeader.center.text = ""
    ws.print_options.horizontalCentered = True


def save(folder, filename, title, notes, headers, rows, widths, footers):
    path = OUT / folder
    path.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = filename.replace(".xlsx", "")[:31]
    write_sheet(ws, title, notes, headers, rows, widths, footers)
    wb.save(path / filename)
    return path / filename


MAP_H = [
    "No", "必須", "受信項目", "受信 型/桁", "送信項目", "送信 型/桁",
    "判定", "変換ルール", "確定", "根拠",
]
MAP_W = [6, 8, 18, 16, 28, 16, 8, 55, 12, 28]

GAP_H = ["No", "優先度", "対象", "現状", "要件", "対応方針", "ステータス", "参照"]
GAP_W = [12, 10, 18, 36, 32, 48, 14, 22]

RULE_H = ["ルールID", "種別", "対象", "変換ロジック", "例", "依存", "確定度"]
RULE_W = [12, 10, 18, 62, 28, 16, 14]


def build():
    made = []

    # ---------- 11 入荷実績 ----------
    notes11 = [
        f"改訂 {DATE}",
        "Sinops-11 入荷実績 nyuka.txt",
        "往路 TRIAL → Sinops",
        "Layer① 変換 / Layer③ 転送のみ",
        SRC,
    ]
    map11 = [
        ["1", "○", "店舗コード", "文字/100", "店CD／拠点コード", "ALLSIRE_MEISAI", "△",
         "統合マスタ後の店CD。本IFで 6→4 等の切捨てはしない", "確定", "R11-02 / G11-02"],
        ["2", "○", "商品コード", "文字/100", "JAN／商品コード", "ALLSIRE_MEISAI", "△",
         "統合マスタ後の商品CD。独自の仕入JAN変換はしない", "確定", "R11-03 / G11-03"],
        ["3", "○", "入荷日", "日付/8", "納品日", "日付", "○",
         "YYYYMMDD のまま", "確定", "R11-04"],
        ["4", "○", "伝票No", "文字/100", "伝票NO／仕入伝票番号", "—", "✕",
         "文字列化。入荷予定とのキー揃え、空白時の扱いは未確定", "確認中", "R11-05 / TBD-NY-05"],
        ["5", "○", "便区分", "整数/2", "（ソースは見ない）", "—", "―",
         "固定値 1。ソース便コードの trim や 0→9 は使わない", "確定", "R11-06 / TBD-NY-BIN"],
        ["6", "○", "入荷数量", "整数/8", "納品数", "数値", "△",
         "整数。メーカー欠品・キャンセルは 0。店間移動は移動元マイナス／移動先プラス", "一部確定", "R11-07 / G11-05"],
    ]
    gap11 = [
        ["G11-01", "—", "便区分", "Sinops に便設定なし。ソース便コードはある可能性",
         "必須。0～99", "固定値 1。ソースは参照しない（2026-09-18）", "完了", "TBD-NY-BIN"],
        ["G11-02", "★高", "店舗コード", "TRIAL 店CD桁", "文字100",
         "マスタ統合後のコード。本IF独自変換禁止", "確定", "TBD-01 相当"],
        ["G11-03", "★高", "商品コード", "JAN 等", "文字100",
         "マスタ統合後のコード。独自仕入JAN変換禁止", "確定", "マスタ統合"],
        ["G11-04", "★高", "伝票No", "仕入伝票番号", "予定と対応する番号",
         "文字列化まで記載。予定↔実績のキー、空白時は未決", "確認中", "TBD-NY-05"],
        ["G11-05", "★中", "欠品・店間移動", "取消・伝票区分", "数量 0 または符号",
         "欠品・キャンセルは 0。店間移動は符号付き。対象伝票区分の一覧は未固定", "一部確定", "R11-07"],
        ["G11-06", "★高", "抽出元", "—", "仕入明細",
         "md-data-integration-prod/LV3/ALLSIRE_MEISAI/YYYYMMDD/", "完了", "TBD-NY-01"],
        ["G11-07", "★中", "連携パス", "—", "西友連携GCS",
         "ods-seiyu-prod/IF0079_LM0026_101(連番)/YYYYMMDD/nyuka_YYYYMMDDhhmmss.gz。完了印は空の nyuka.end", "完了", "TBD-NY-03"],
        ["G11-08", "★中", "Sinops着地", "Layer③ 受信先が【要確認】", "Sinops 取込先",
         "パス未記入。差分取込・56日削除は旧記載。DD は差分とだけ書く", "要確認", "05 Layer③"],
        ["G11-09", "★低", "件数", "未記載", "想定件数・データ量", "未着手", "未着手", "TBD-NY-04"],
        ["G11-10", "★中", "変換エラー", "必須欠落・非数値・日付不正", "中断かレコード退避か",
         "暫定はファイル中断。便区分は固定 1 なので欠損扱いにしない", "確認中", "TBD-NY-11 / ER-NY-002"],
        ["G11-11", "★中", "0件", "対象日 0 件", "空ファイルを出すかスキップか", "未決", "未着手", "TBD-NY-12 / ER-NY-006"],
        ["G11-12", "★低", "文字コード", "—", "UTF-8（BOMなし）・LF・カンマ", "確定", "完了", "TBD-NY-10"],
        ["G11-13", "★低", "再送", "—", "世代・上書き", "DD 上は完了。方針本文は空", "完了", "TBD-NY-09"],
    ]
    rule11 = [
        ["R11-01", "抽出", "対象日", "勧告計算前日分の入荷実績。既定は日次 02:00 前後", "—", "TBD-NY-02", "確定"],
        ["R11-02", "コード", "店舗コード", "統合マスタ後の店CD。切捨て禁止", "—", "G11-02", "確定"],
        ["R11-03", "コード", "商品コード", "統合マスタ後の商品CD。独自仕入JAN変換禁止", "—", "G11-03", "確定"],
        ["R11-04", "日付", "入荷日", "納品日を YYYYMMDD", "20261018", "—", "確定"],
        ["R11-05", "コード", "伝票No", "仕入伝票番号を文字列化。予定側と同一キーになるよう整える", "—", "G11-04", "確認中"],
        ["R11-06", "固定", "便区分", "常に 1。ソース便コードは読まない", "1", "G11-01", "確定"],
        ["R11-07", "数量", "入荷数量", "整数。欠品・キャンセルは 0。店間移動は符号付き", "欠品→0", "G11-05", "一部確定"],
        ["R11-08", "形式", "ファイル", "CSV（カンマ、囲みなし、ヘッダなし）を gz。論理名 nyuka.txt。物理名 nyuka_YYYYMMDDhhmmss.gz。続けて空の nyuka.end",
         "nyuka_20261019020000.gz", "G11-07", "確定"],
        ["R11-09", "配置", "GCS", "gs://ods-seiyu-prod/IF0079_LM0026_101(連番)/YYYYMMDD/", "—", "G11-07", "確定"],
        ["R11-10", "転送", "Layer③", "nyuka.end の後に gz を取得し Sinops へ。型変換なし", "—", "G11-08", "要確認"],
        ["R11-11", "例外", "必須", "必須欠落・非数値・日付不正はファイル中断（レコード退避は未決）。取得失敗・配置失敗は中断してリラン",
         "ER-NY-001〜003", "G11-10", "一部確定"],
    ]
    foot11 = [
        "共通: 抽出元 md-data-integration-prod/LV3/ALLSIRE_MEISAI/YYYYMMDD/",
        "Layer② 対象外。変換は Layer① のみ。",
    ]
    made.append(save("11_入荷実績", "①項目マッピング表(nyuka.txt).xlsx",
                     "項目マッピング表 — 入荷実績（nyuka.txt）6項目", notes11, MAP_H, map11, MAP_W, foot11))
    made.append(save("11_入荷実績", "②GAP分析書(nyuka.txt).xlsx",
                     "GAP分析書 — 入荷実績（nyuka.txt）", notes11, GAP_H, gap11, GAP_W, foot11))
    made.append(save("11_入荷実績", "③変換ルール定義書(nyuka.txt).xlsx",
                     "変換ルール定義書 — 入荷実績（nyuka.txt）", notes11, RULE_H, rule11, RULE_W, foot11))

    # ---------- 12 入荷予定 ----------
    notes12 = [
        f"改訂 {DATE}",
        "Sinops-12 入荷予定 nyuka_yotei.txt",
        "往路 TRIAL → Sinops",
        "抽出元 HACHU_JISEKI",
        SRC,
    ]
    map12 = [
        ["1", "○", "店舗コード", "文字/100", "拠点コード／店CD", "HACHU_JISEKI", "△",
         "統合マスタ後の店CD。切捨てなし", "確定", "R12-02 / G12-04"],
        ["2", "○", "伝票No", "文字/100", "伝票NO／伝票管理番号", "—", "✕",
         "文字列化。店×商品でユニークにする番号。空白時は未決", "確認中", "R12-08 / TBD-NY-05"],
        ["3", "○", "商品コード", "文字/100", "商品コード／JAN", "HACHU_JISEKI", "△",
         "統合マスタ後の商品CD", "確定", "R12-02"],
        ["4", "○", "入荷予定数量", "整数/8", "発注数／発注数量・単位数", "数値", "✕",
         "必ずバラ。源が発注単位数なら ×発注単位。源がどちらに当たるか未決", "確認中", "R12-03 / TBD-NY-06"],
        ["5", "○", "入荷予定日", "日付/8", "納品日付／納品予定日付", "日付", "△",
         "納品日付。空白なら納品予定日付。YYYYMMDD", "一部確定", "R12-04"],
        ["6", "○", "発注日", "日付/8", "発注日付／発注予定日付", "日付", "△",
         "発注日付。空白なら発注予定日付。YYYYMMDD", "一部確定", "R12-05"],
        ["7", "○", "便区分", "整数/2", "（ソースは見ない）", "—", "―",
         "固定値 1。0→9 は採用しない", "確定", "R12-06 / TBD-NY-BIN"],
        ["8", "－", "コメント", "文字/100", "発注データ区分等", "コード", "✕",
         "コードを文言にする（任意）。値表が無い間は空。勧告計算には影響しない", "確認中", "R12-07 / TBD-NY-07"],
    ]
    gap12 = [
        ["G12-01", "—", "便区分", "ソース便コード", "必須",
         "固定値 1。0→9・空白除去は使わない", "完了", "TBD-NY-BIN"],
        ["G12-02", "★最高", "入荷予定数量", "発注数または単位数", "バラ数",
         "×発注単位が要るかは HACHU_JISEKI の数量単位が未確認", "確認中", "TBD-NY-06"],
        ["G12-03", "★中", "コメント", "発注データ区分", "任意の文言",
         "値表未確定。空可。特売・本部・手発注を含む方針は概要に記載", "確認中", "TBD-NY-07"],
        ["G12-04", "★高", "店・商品", "拠点コード／JAN", "文字100",
         "マスタ統合後。切捨てなし", "確定", "R12-02"],
        ["G12-05", "★高", "伝票No", "伝票NO／伝票管理番号", "店×商品でユニーク",
         "空白時、および実績とのキー揃えは未決", "確認中", "TBD-NY-05"],
        ["G12-06", "★中", "日付", "納品日付・発注日付が空の行がある", "必須日付",
         "空白時は予定日付へ落とす。両方空の扱いは未記載", "一部確定", "R12-04 / R12-05"],
        ["G12-07", "★高", "抽出範囲", "発注残", "本日入荷分を含む未来の未入荷",
         "抽出条件の文言はある。対象業務日の確定状態は確認中", "確認中", "TBD-NY-02"],
        ["G12-08", "★中", "取込方式", "—", "差分か全件か", "確認中", "確認中", "TBD-NY-08"],
        ["G12-09", "★中", "連携パス", "—", "西友連携GCS",
         "ods-seiyu-prod/IF0079_LM0026_102(連番)/YYYYMMDD/nyuka_yotei_YYYYMMDDhhmmss.gz。完了印 nyuka_yotei.end", "完了", "TBD-NY-03"],
        ["G12-10", "★高", "抽出元", "旧: 発注伝票明細", "HACHU_JISEKI",
         "md-data-integration-prod/LV3/HACHU_JISEKI/YYYYMMDD/（0.6 で変更）", "完了", "TBD-NY-01"],
        ["G12-11", "★中", "Sinops着地", "Layer③ 受信先が【要確認】", "Sinops 取込先パス", "未記入", "要確認", "05 Layer③"],
        ["G12-12", "★低", "件数", "未記載", "想定件数", "確認中", "確認中", "TBD-NY-04"],
        ["G12-13", "★中", "0件・エラー単位", "—", "空ファイル、中断か退避か", "確認中", "確認中", "TBD-NY-11 / TBD-NY-12"],
        ["G12-14", "★低", "文字コード", "—", "UTF-8（BOMなし）・LF・カンマ", "確定", "完了", "TBD-NY-10"],
    ]
    rule12 = [
        ["R12-01", "抽出", "対象", "HACHU_JISEKI から、本日入荷分を含む未来の未入荷発注残。特売・本部発注・手発注を含む",
         "—", "G12-07", "確認中"],
        ["R12-02", "コード", "店・商品", "統合マスタ後の店CD・商品CD。切捨てなし", "—", "G12-04", "確定"],
        ["R12-03", "数量", "入荷予定数量", "バラで出す。源が発注単位数なら ×発注単位。源がバラならそのまま。丸めの記載なし",
         "単位数 2 × 入数 6 → 12", "G12-02", "確認中"],
        ["R12-04", "日付", "入荷予定日", "納品日付。空白なら納品予定日付。YYYYMMDD", "—", "G12-06", "一部確定"],
        ["R12-05", "日付", "発注日", "発注日付。空白なら発注予定日付。YYYYMMDD", "—", "G12-06", "一部確定"],
        ["R12-06", "固定", "便区分", "常に 1。ソース便コードは無視", "1", "G12-01", "確定"],
        ["R12-07", "コード", "コメント", "発注データ区分を文言化。値表が無い間は空。任意項目", "空", "G12-03", "確認中"],
        ["R12-08", "コード", "伝票No", "伝票番号を文字列化。空白時は未決", "—", "G12-05", "確認中"],
        ["R12-09", "形式", "ファイル", "CSV を gz。論理 nyuka_yotei.txt。物理 nyuka_yotei_YYYYMMDDhhmmss.gz。その後に空の nyuka_yotei.end",
         "—", "G12-09", "確定"],
        ["R12-10", "配置", "GCS", "gs://ods-seiyu-prod/IF0079_LM0026_102(連番)/YYYYMMDD/", "—", "G12-09", "確定"],
        ["R12-11", "転送", "Layer③", "end の後に gz を取得。型変換なし", "—", "G12-11", "要確認"],
        ["R12-12", "例外", "STEP", "取得失敗 ER-NY-001 中断。必須欠落・非数値・日付不正 ER-NY-002。0件は ER-NY-006（継続かスキップは未決）",
         "—", "G12-13", "一部確定"],
    ]
    foot12 = [
        "処理順: 取得 → 未入荷抽出 → 0件判定 → 必須確認 → 店商品 → 伝票No → 数量 → 日付2つ → 便=1 → コメント → CSV → gz → GCS → end",
        "Layer② 対象外。",
    ]
    made.append(save("12_入荷予定", "①項目マッピング表(nyuka_yotei.txt).xlsx",
                     "項目マッピング表 — 入荷予定（nyuka_yotei.txt）8項目", notes12, MAP_H, map12, MAP_W, foot12))
    made.append(save("12_入荷予定", "②GAP分析書(nyuka_yotei.txt).xlsx",
                     "GAP分析書 — 入荷予定（nyuka_yotei.txt）", notes12, GAP_H, gap12, GAP_W, foot12))
    made.append(save("12_入荷予定", "③変換ルール定義書(nyuka_yotei.txt).xlsx",
                     "変換ルール定義書 — 入荷予定（nyuka_yotei.txt）", notes12, RULE_H, rule12, RULE_W, foot12))

    # ---------- 24 / 25 shared item rows ----------
    def kankoku_maps(landing, src_file, src_path):
        return [
            ["1", "○", "店舗CD", "固定長/4  BranchCD", "店舗コード", "固定長/6", "△",
             "統合マスタ後の店CD。6→4 の独自切捨てなし。桁は統合後", "一部確定", "TBD-01"],
            ["2", "○", "仕入発注JAN", "固定長/20  SireJAN", "商品コード", "固定長/13", "△",
             "統合後の商品コード。独自の仕入JAN変換なし。先頭ゼロ保持", "一部確定", "TBD-11"],
            ["3", "○", "推奨発注数量", "可変長  SireJAN_rcmd_end", "発注バラ数", "数値", "○",
             "×1000 しない。発注単位の整数倍なら可。倍数へ丸めない", "確定", "R04"],
            ["4", "○", "ベンダーCD", "固定長/8", "仕入先コード", "固定長/9", "△",
             "統合マスタ後の仕入先。9→8 の独自切捨てなし", "一部確定", "TBD-08"],
            ["5", "○", "納品日", "固定長/8  ShippingDate", "納品日", "固定長/8", "○",
             "yyyymmdd のまま", "確定", "R06"],
            ["6", "○", "Ordertype", "固定長/1", "西友中分類（ライン＝部門）", "勧告ファイル内", "△",
             "Sinops に Ordertype は無い。部門CD → JIHA.BUMON_ZONE → ZONECD → JIHA.ZONE_LOGIC。01→0、02→1、03→1。商品マスタは見ない",
             "確定", "R07 / TBD-09 / TBD-18"],
            ["7", "○", "更新日", "固定長/8", "発注日（候補）", "固定長/8", "✕",
             "業務日か DB 更新日か未決。09 は完了だが回答が空。マッピングは未確定", "未確定", "TBD-10"],
            ["—", "－", "便区分", "送らない", "便区分", "文字/2", "―", "着地 7 項目に含めない", "確定", "選別"],
            ["—", "－", "分類・単価・単位", "送らない", "大/中/小/細目、発注単位数、標準仕入単価、標準売価、発注単位",
             "Sinops 15項目側", "―", "着地に含めない。中分類だけ Ordertype の部門CDに使う", "確定", "選別"],
        ]

    gap_k = [
        ["K-01", "★高", "店CD桁", "Sinops 6桁", "着地 4桁",
         "統合後コード。独自切捨てなし。実桁は統合結果待ち", "一部確定", "TBD-01"],
        ["K-02", "★高", "商品コード", "Sinops 13桁", "SireJAN 20桁",
         "統合後。独自変換なし。先頭ゼロ保持", "一部確定", "TBD-11"],
        ["K-03", "★高", "仕入先", "Sinops 9桁", "ベンダー 8桁",
         "統合後。独自切捨てなし", "一部確定", "TBD-08"],
        ["K-04", "—", "数量", "発注バラ数", "推奨発注数量",
         "×1000 しない。整数倍なら可。丸めない", "完了", "R04 / G04"],
        ["K-05", "—", "Ordertype", "Sinops に項目なし", "0=非日配、1=日配または生鮮",
         "部門は勧告の西友中分類。01→0、02→1、03→1。商品マスタは見ない", "完了", "R07 / TBD-18"],
        ["K-06", "★最高", "更新日", "発注日が候補", "固定長8・必須",
         "業務日か DB 更新日か。09 の状態は完了、回答は空、マッピングは未確定", "未確定", "TBD-10"],
        ["K-07", "★高", "採用判定", "当日と昨日の翌日以降が両方あり得る", "1本だけ採用",
         "Layer① 共通 1 ジョブ。締切までに当日があれば当日、無ければ昨日の翌日以降。2本は重ねない。締切時刻は例 6:00 のまま",
         "一部確定", "TBD-13 / TBD-19"],
        ["K-08", "★高", "同一キー", "BranchCD+SireJAN", "最終 1 件の決め方", "未着手", "未着手", "TBD-12"],
        ["K-09", "★高", "ファイル形式", "固定長混在（数量だけ可変）", "文字コード・改行・区切り",
         "UTF-8 / LF は候補。区切りは未決", "未着手", "TBD-07 / TBD-15"],
        ["K-10", "★高", "時刻", "勧告計算後", "Layer③ 配置時刻と Layer① プル窓", "未着手", "未着手", "TBD-02"],
        ["K-11", "★中", "再送", "同一業務日の再実行", "上書きか新規か", "未着手。同名上書きだと履歴は残らない", "未着手", "TBD-03"],
        ["K-12", "★中", "0件", "勧告 0 件", "空ファイルを出すか", "未着手", "未着手", "TBD-16"],
        ["K-13", "★中", "変換エラー", "コード不一致・整数倍でない・部門が引けない", "ファイル中断かレコード退避か",
         "未着手", "未着手", "TBD-17 / ER-005 / ER-007"],
        ["K-14", "★低", "件数", "未記載", "想定件数・データ量", "未着手", "未着手", "TBD-14"],
        ["K-15", "★高", "Sinops送信先", "Layer③ 送信元が【要確認】", "Sinops が置く場所",
         "受信側（連携用GCS）は確定。Sinops 側の置き場は未記入", "要確認", "05 Layer③"],
        ["K-16", "—", "西友GCS", "—", "gz と空の .end",
         "当日 kankoku_YYYYMMDDhhmmss.gz / kankoku.end。翌日以降は subkankoku", "完了", "TBD-04 / TBD-05"],
        ["K-17", "—", "着地", "—", "LV3 の SIREJAN",
         "当日採用: .../SIREJAN_RCMDORDER_LAST/YYYYMMDD/。翌日以降採用: .../SIREJAN_RCMDORDER_LAST_SUB/YYYYMMDD/。着地は gz ではない",
         "完了", "TBD-06"],
    ]
    rule_common_tail = [
        ["R04", "数量", "推奨発注数量", "発注バラ数をそのまま。×1000 しない。発注単位の整数倍でなければエラー。丸めない",
         "バラ 12 → 12", "K-04", "確定"],
        ["R06", "日付", "納品日", "yyyymmdd のまま ShippingDate へ", "20261020", "—", "確定"],
        ["R07", "編集", "Ordertype",
         "部門CD＝勧告ファイルの西友中分類（ライン、階層4）。JIHA.BUMON_ZONE で ZONECD、JIHA.ZONE_LOGIC で論理区分。01→0、02→1、03→1。TRIAL商品マスタは見ない。参照のみで書き戻さない",
         "01→0 / 02→1 / 03→1", "K-05", "確定"],
        ["R08", "日付", "更新日", "発注日を候補にするが、業務日か DB 更新日かは未決", "—", "K-06", "未確定"],
        ["R09", "コード", "店舗CD", "統合マスタ後。6→4 切捨て禁止。先頭の扱いも統合後", "—", "K-01", "一部確定"],
        ["R10", "コード", "仕入発注JAN", "統合後の商品コード。独自変換なし。先頭ゼロ保持", "—", "K-02", "一部確定"],
        ["R11", "コード", "ベンダーCD", "統合マスタ後。9→8 切捨て禁止", "—", "K-03", "一部確定"],
        ["R12", "選別", "非着地項目", "便区分、分類（部門CDの参照を除く）、発注単位数、単価、売価、発注単位は着地に出さない",
         "—", "—", "確定"],
    ]

    notes24 = [
        f"改訂 {DATE}",
        "Sinops-24 発注勧告 当日 kankoku.txt",
        "復路 Sinops → TRIAL",
        "Layer① は 25 と共通 1 ジョブ",
        SRC,
    ]
    map24 = kankoku_maps(
        "SIREJAN_RCMDORDER_LAST",
        "kankoku",
        "LM0026_IF0079_101",
    )
    rule24 = [
        ["R24-01", "受信", "入力",
         "空の kankoku.end / subkankoku.end を見てから gz を取る。当日 kankoku_*.gz と、昨日の subkankoku_*.gz。最大 2 ファイル",
         "—", "K-16", "確定"],
        ["R24-02", "採用", "1本化",
         "締切（TBD-19、例 6:00）までに当日が作成済なら当日。障害または未作成なら昨日の翌日以降。2本は重ねない。両方無しは中断",
         "当日あり → LAST", "K-07", "一部確定"],
        ["R24-03", "配置", "着地",
         "当日採用の着地は md-data-integration-prod/LV3/SIREJAN_RCMDORDER_LAST/YYYYMMDD/。gz ではない。OrdreSV への転送は本IFの外",
         "—", "K-17", "確定"],
    ] + rule_common_tail + [
        ["R24-04", "形式", "西友側",
         "連携用GCS: ods-seiyu-prod/LM0026_IF0079_101(連番)/YYYYMMDD/kankoku_YYYYMMDDhhmmss.gz。完了印 kankoku.end",
         "—", "K-16", "確定"],
        ["R24-05", "例外", "変換",
         "店・商品・仕入先の不一致、数量が整数倍でない、部門CDが引けない、論理区分が 01/02/03 以外はエラー。中断か退避かは未決",
         "ER-005 / ER-007", "K-13", "未着手"],
    ]
    foot24 = [
        "受信: ods-seiyu-prod/LM0026_IF0079_101(連番)/YYYYMMDD/kankoku_YYYYMMDDhhmmss.gz",
        "着地（当日採用時）: md-data-integration-prod/LV3/SIREJAN_RCMDORDER_LAST/YYYYMMDD/",
        "Layer③ は変換しない。Layer② 対象外。",
    ]
    made.append(save("24_発注勧告", "①項目マッピング表(kankoku.txt).xlsx",
                     "項目マッピング表 — 発注勧告当日（kankoku.txt）着地7項目", notes24, MAP_H, map24, MAP_W, foot24))
    made.append(save("24_発注勧告", "②GAP分析書(kankoku.txt).xlsx",
                     "GAP分析書 — 発注勧告当日（kankoku.txt）", notes24, GAP_H, gap_k, GAP_W, foot24))
    made.append(save("24_発注勧告", "③変換ルール定義書(kankoku.txt).xlsx",
                     "変換ルール定義書 — 発注勧告当日（kankoku.txt）", notes24, RULE_H, rule24, RULE_W, foot24))

    notes25 = [
        f"改訂 {DATE}",
        "Sinops-25 発注勧告 翌日以降 subkankoku.txt",
        "項目変換は 24 と同一",
        "採用されたときだけ LAST_SUB へ着地",
        SRC,
    ]
    map25 = kankoku_maps(
        "SIREJAN_RCMDORDER_LAST_SUB",
        "subkankoku",
        "LM0026_IF0079_102",
    )
    rule25 = [
        ["R25-01", "受信", "入力",
         "Layer③ は毎日 subkankoku_YYYYMMDDhhmmss.gz と空の subkankoku.end を出す。OrdreSV へはここでは出さない",
         "—", "K-16", "確定"],
        ["R25-02", "採用", "代替",
         "Layer① は 24 と共通。当日が締切まで未作成、または障害のとき、昨日の本ファイルを採用。後勝ちしない。当日があれば本ファイルは着地しない",
         "当日なし → LAST_SUB", "K-07", "一部確定"],
        ["R25-03", "配置", "着地",
         "採用時の着地は md-data-integration-prod/LV3/SIREJAN_RCMDORDER_LAST_SUB/YYYYMMDD/。gz ではない",
         "—", "K-17", "確定"],
    ] + rule_common_tail + [
        ["R25-04", "形式", "西友側",
         "ods-seiyu-prod/LM0026_IF0079_102(連番)/YYYYMMDD/subkankoku_YYYYMMDDhhmmss.gz。完了印 subkankoku.end",
         "—", "K-16", "確定"],
        ["R25-05", "例外", "変換",
         "項目変換の例外は 24 と同一（ER-005 / ER-007）。中断か退避かは未決",
         "—", "K-13", "未着手"],
    ]
    foot25 = [
        "受信: ods-seiyu-prod/LM0026_IF0079_102(連番)/YYYYMMDD/subkankoku_YYYYMMDDhhmmss.gz",
        "着地（代替採用時）: md-data-integration-prod/LV3/SIREJAN_RCMDORDER_LAST_SUB/YYYYMMDD/",
        "項目の変換内容は当日分と同じ。違うのはファイル名、連番 102、着地フォルダ。",
    ]
    made.append(save("25_発注勧告翌日以降", "①項目マッピング表(subkankoku.txt).xlsx",
                     "項目マッピング表 — 発注勧告翌日以降（subkankoku.txt）着地7項目", notes25, MAP_H, map25, MAP_W, foot25))
    made.append(save("25_発注勧告翌日以降", "②GAP分析書(subkankoku.txt).xlsx",
                     "GAP分析書 — 発注勧告翌日以降（subkankoku.txt）", notes25, GAP_H, gap_k, GAP_W, foot25))
    made.append(save("25_発注勧告翌日以降", "③変換ルール定義書(subkankoku.txt).xlsx",
                     "変換ルール定義書 — 発注勧告翌日以降（subkankoku.txt）", notes25, RULE_H, rule25, RULE_W, foot25))

    for p in made:
        print(p)


if __name__ == "__main__":
    build()
