#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Remake Phase1 ①②③ for Sinops-11 / 12 / 21 (2026-09-21).

Incorporates Seiyu IF layout upgrades:
  - 入荷実績: 便コード空白除去（空→0、trim）
  - 入荷予定: バラ換算・日付fallback・便0→9・コメント
  - 受払明細: 出庫発行日＝検収相当（非納品日付）
"""

from __future__ import annotations

import shutil
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).parent
TODAY = date(2026, 9, 21)
REF = (
    "参照：sinops-R6_IFレイアウト(入荷実績/入荷予定)、"
    "sinops-W_IFレイアウト(受払明細)（Downloads 2026-09-21 精查）"
)

BLUE = "1F4E78"
LIGHT_BLUE = "D9EAF7"
LIGHT_YELLOW = "FFF2CC"
LIGHT_RED = "FCE4D6"
LIGHT_GREEN = "E2F0D9"
GRAY = "E7E6E6"
WHITE = "FFFFFF"
THIN = Side(style="thin", color="808080")
LEGEND = "判定凡例：  ○ 直接マッピング可  △ 変換・加工が必要  ✕ GAPあり（要対応）  ― 固定値またはNULL設定"


def fill(hex_color: str) -> PatternFill:
    return PatternFill("solid", fgColor=hex_color)


def apply_grid(ws, widths, header_row=4, verdict_col=None, data_start=5):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
    for cell in ws[1]:
        cell.fill = fill(BLUE)
        cell.font = Font(color=WHITE, bold=True, size=12)
    for cell in ws[header_row]:
        cell.fill = fill(LIGHT_BLUE)
        cell.font = Font(bold=True)
    if verdict_col:
        for r in range(data_start, ws.max_row + 1):
            raw = str(ws.cell(r, verdict_col).value or "")
            first = raw[:1]
            if first == "○":
                c = LIGHT_GREEN
            elif first == "△":
                c = LIGHT_YELLOW
            elif first in ("✕", "×"):
                c = LIGHT_RED
            elif first in ("―", "—"):
                c = GRAY
            else:
                c = WHITE
            if str(ws.cell(r, 1).value or "").startswith("■"):
                c = LIGHT_BLUE
            for cell in ws[r]:
                cell.fill = fill(c)
    ws.freeze_panes = f"A{header_row + 1}"


def intro(ws, title: str, subtitle: str, note: str, ncols: int):
    ws.append([title])
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)
    ws.append([subtitle])
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=ncols)
    ws.append([note])
    ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=ncols)


def archive_old(paths: list[Path], bk_dir: Path):
    bk_dir.mkdir(parents=True, exist_ok=True)
    for p in paths:
        if p.exists() and p.is_file():
            dest = bk_dir / p.name
            if dest.exists():
                dest = bk_dir / f"{p.stem}_before_{TODAY.isoformat()}{p.suffix}"
            shutil.move(str(p), str(dest))


# ---------------------------------------------------------------------------
# Sinops-11 入荷実績
# ---------------------------------------------------------------------------

def build_11_mapping(out: Path):
    wb = Workbook()
    ws = wb.active
    ws.title = "入荷実績(nyuka.txt)"
    cols = [
        "No", "必須", "sinops項目名", "sinops型/桁", "FMT・範囲",
        "TRIAL項目名（候補）", "TRIAL型/桁", "判定", "変換ルール・留意事項", "根拠・参照",
    ]
    intro(
        ws,
        "項目マッピング表 — 入荷実績（nyuka.txt）6項目 ／ Sinops-11",
        f"改訂 {TODAY.isoformat()} ／ TRIAL→sinops ／ {REF}",
        LEGEND + " ／ 店・商品はマスタ統合に追随（本IFは独自桁切捨てしない）",
        len(cols),
    )
    ws.append(cols)
    rows = [
        ["1", "○", "店舗コード", "文字/100", "店舗識別",
         "店CD／拠点コード", "4桁固定等", "△",
         "マスタ統合後の店舗コードを設定。本IFは独自の6→4切捨てを持たない。桁は統合作業に追随",
         "マスタ統合方針／TBD相当"],
        ["2", "○", "商品コード", "文字/100", "商品マスタの商品コード",
         "JAN／商品コード", "20桁等", "△",
         "マスタ統合後の商品コードを設定。本IFは独自の仕入JAN変換を持たない",
         "マスタ統合方針"],
        ["3", "○", "入荷日", "日付/8", "yyyymmdd",
         "納品日", "8/固定長", "○",
         "YYYYMMDD 直接設定",
         "sinops IF060"],
        ["4", "○", "伝票No", "文字/100", "入荷予定の伝票Noと対応",
         "伝票NO／仕入伝票番号", "10〜13桁", "△",
         "入荷予定と紐付く番号を設定。現状西友注記では予定/実績紐づけに伝票番号未使用の記載あり→TRIALでも要確認",
         "西友マッピング注記"],
        ["5", "○", "便区分", "整数値/2", "0～99",
         "便コード／物流タイプ等（未確定）", "—", "✕",
         "G-03/G-06 未確定。ソースがある場合：半角空白 trim、空白は 0 に置換（西友 2024-01-17 升级）",
         "西友更新履歴 便コード空白除去"],
        ["6", "○", "入荷数量", "整数値/8", "-99999999～99999999",
         "納品数", "可変", "△",
         "整数で設定。メーカー欠品・キャンセルは数量 0。店間移動は移動元マイナス／移動先プラス",
         "sinops注意点"],
    ]
    for r in rows:
        ws.append(r)
    ws.append([])
    ws.append(["■共通"])
    ws.append(["出力タイミング", "勧告計算前日分。障害時は未送信分を次回に混ぜ込む"])
    ws.append(["sinops取込", "差分更新。取込データは56日後自動削除"])
    ws.append(["ファイル形式", "TSV（タブ区切り）可変長"])
    apply_grid(ws, [5, 6, 14, 14, 18, 18, 12, 6, 42, 22], verdict_col=8)
    wb.save(out)


def build_11_gap(out: Path):
    wb = Workbook()
    ws = wb.active
    ws.title = "GAP分析一覧"
    cols = ["No", "優先度", "対象項目", "TRIAL側", "sinops要件", "対応方針", "ステータス", "参照"]
    intro(
        ws,
        "GAP分析書 — 入荷実績（nyuka.txt）／ Sinops-11",
        f"改訂 {TODAY.isoformat()} ／ {REF}",
        "Phase1旧版＋西友レイアウト升级を統合",
        len(cols),
    )
    ws.append(cols)
    gaps = [
        ["G11-01", "★最高", "便区分", "該当項目なし／物流タイプのみの可能性",
         "必須キー・0～99", "G-03/G-06でキー・値域・付与方式を確定。暫定：trim＋空→0",
         "未完了", "traceability G-03/G-06"],
        ["G11-02", "★高", "店舗コード桁", "4桁等", "文字100",
         "マスタ統合に追随。本IF独自変換禁止", "方針確認済", "マスタ統合"],
        ["G11-03", "★高", "商品コード桁", "20桁等", "文字100",
         "マスタ統合に追随。本IF独自変換禁止", "方針確認済", "マスタ統合"],
        ["G11-04", "★中", "伝票No紐付け", "伝票NO", "予定と対応必須（定義上）",
         "西友注記では現行未使用の記載あり。TRIALで予定↔実績の紐付け要否を確認",
         "未完了", "西友マッピング"],
        ["G11-05", "★中", "欠品・キャンセル", "取消区分等", "数量0で連携",
         "欠品時は入荷実績を数量0で出すルールを固定", "一部確定", "sinops注意点"],
        ["G11-06", "★中", "店間移動・返品", "伝票区分", "入荷数量に±反映",
         "どの伝票区分を入荷実績に含めるか合意", "未完了", "旧GAP#17"],
        ["G11-07", "★低", "取込方式", "—", "差分・56日削除",
         "差分更新で合意（西友レイアウト記載）", "方針確認済", "西友マッピング"],
    ]
    for g in gaps:
        ws.append(g)
    apply_grid(ws, [10, 8, 14, 22, 22, 40, 12, 18], verdict_col=None)
    wb.save(out)


def build_11_rule(out: Path):
    wb = Workbook()
    ws = wb.active
    ws.title = "変換ルール"
    cols = ["ルールID", "種別", "対象", "変換ロジック", "例", "依存GAP", "確定度"]
    intro(
        ws,
        "変換ルール定義書 — 入荷実績（nyuka.txt）／ Sinops-11",
        f"改訂 {TODAY.isoformat()}",
        "Layer①/③ 実装時の STEP 根拠",
        len(cols),
    )
    ws.append(cols)
    rules = [
        ["R11-01", "抽出", "対象日", "勧告計算前日分の入荷実績を出力。障害時は未送信分を次回に合算可",
         "—", "—", "方針確認済"],
        ["R11-02", "コード", "店舗コード", "統合マスタ後の店CDを設定。切捨て禁止",
         "—", "G11-02", "方針確認済"],
        ["R11-03", "コード", "商品コード", "統合マスタ後の商品CDを設定。独自仕入JAN変換禁止",
         "—", "G11-03", "方針確認済"],
        ["R11-04", "日付", "入荷日", "納品日を YYYYMMDD で設定",
         "20260920", "—", "確定"],
        ["R11-05", "コード", "伝票No", "仕入伝票番号等を文字列化。予定側と同一キーになるよう整える（要確認）",
         "—", "G11-04", "未確定"],
        ["R11-06", "編集", "便区分", "ソース値を trim。空／空白のみ → 0。値域外はエラー",
         "'1 '→1 / ''→0", "G11-01", "一部確定（編集のみ）"],
        ["R11-07", "数量", "入荷数量", "整数。欠品・キャンセルは 0。店間移動は符号付き",
         "欠品→0", "G11-05", "一部確定"],
        ["R11-08", "形式", "ファイル", "TSV・可変長・差分更新",
         "nyuka.txt", "G11-07", "確定"],
    ]
    for r in rules:
        ws.append(r)
    apply_grid(ws, [10, 8, 12, 48, 16, 10, 14], verdict_col=None)
    wb.save(out)


# ---------------------------------------------------------------------------
# Sinops-12 入荷予定
# ---------------------------------------------------------------------------

def build_12_mapping(out: Path):
    wb = Workbook()
    ws = wb.active
    ws.title = "入荷予定(nyuka_yotei.txt)"
    cols = [
        "No", "必須", "sinops項目名", "sinops型/桁", "FMT・範囲",
        "TRIAL項目名（候補）", "TRIAL型/桁", "判定", "変換ルール・留意事項", "根拠・参照",
    ]
    intro(
        ws,
        "項目マッピング表 — 入荷予定（nyuka_yotei.txt）8項目 ／ Sinops-12",
        f"改訂 {TODAY.isoformat()} ／ TRIAL→sinops ／ {REF}",
        LEGEND + " ／ 西友基幹切替升级（2023-10〜2024-03）の編集仕様を反映",
        len(cols),
    )
    ws.append(cols)
    rows = [
        ["1", "○", "店舗コード", "文字/100", "店舗識別",
         "拠点コード／店CD", "4桁等", "△",
         "マスタ統合後の店舗コード。独自桁切捨て禁止",
         "マスタ統合"],
        ["2", "○", "伝票No", "文字/100", "店×商品でユニーク。実績と連動",
         "伝票NO／伝票管理番号", "9〜15桁", "△",
         "未設定時の扱い要確認。西友AS-ISは空白時スペース挿入等あり。"
         "注：西友現行は予定/実績紐づけに伝票番号未使用との記載→TRIAL要確認",
         "西友更新 2024-01/03"],
        ["3", "○", "商品コード", "文字/100", "商品マスタの商品コード",
         "商品コード／JAN", "20桁等", "△",
         "マスタ統合後の商品コード。独自変換禁止",
         "マスタ統合"],
        ["4", "○", "入荷予定数量", "整数値/8", "-99999999～99999999（バラ）",
         "発注数／発注数量・発注単位数", "可変", "△",
         "必ずバラ数。TRIAL源が発注単位数の場合は 発注数量×発注単位数。"
         "源が既にバラなら直接設定（×しない）",
         "西友 2023-12 基幹切替"],
        ["5", "○", "入荷予定日", "日付/8", "yyyymmdd。本日以降も含む",
         "納品日付／納品予定日付", "8", "△",
         "納品日付が空白なら納品予定日付をセット（西友編集仕様）",
         "西友 2023-10"],
        ["6", "○", "発注日", "日付/8", "yyyymmdd",
         "発注日付／発注予定日付", "8", "△",
         "発注日付が空白なら発注予定日付をセット",
         "西友 2023-10"],
        ["7", "○", "便区分", "整数値/2", "0～99。同一日複数入荷時必須",
         "便コード／物流タイプ等", "—", "✕",
         "G-03/G-06未確定。西友AS-IS：便コード0→9。TRIAL適用可否は要合意",
         "西友編集仕様／G-03"],
        ["8", "", "コメント", "文字/100", "特売送込・本部発注等。勧告計算には非影響",
         "発注データ区分／売上区分", "1〜2桁", "△",
         "コード→文言変換。案：1通常／4日配／10本部／11店舗特売／12新店特売。"
         "西友売上区分参考：0通常・1返品・2特売・7本部",
         "Phase1変換＋西友"],
    ]
    for r in rows:
        ws.append(r)
    ws.append([])
    ws.append(["■共通"])
    ws.append(["範囲", "本日入荷分を含む未来の入荷予定。仕入先へ実際に発注した実績を全て連携"])
    ws.append(["本部・手発注", "特売・本部発注・手発注も本ファイルで連携"])
    ws.append(["取込", "差分または全件を事前合意。キー重複は後優先。56日後削除"])
    apply_grid(ws, [5, 6, 14, 14, 22, 20, 12, 6, 44, 18], verdict_col=8)
    wb.save(out)


def build_12_gap(out: Path):
    wb = Workbook()
    ws = wb.active
    ws.title = "GAP分析一覧"
    cols = ["No", "優先度", "対象項目", "TRIAL側", "sinops要件", "対応方針", "ステータス", "参照"]
    intro(
        ws,
        "GAP分析書 — 入荷予定（nyuka_yotei.txt）／ Sinops-12",
        f"改訂 {TODAY.isoformat()} ／ {REF}",
        "西友基幹切替升级を反映",
        len(cols),
    )
    ws.append(cols)
    gaps = [
        ["G12-01", "★最高", "便区分", "該当なしの可能性",
         "キー必須・同一日複数入荷で必須", "G-03/G-06確定。0→9ルールの採否を合意",
         "未完了", "G-03/G-06"],
        ["G12-02", "★最高", "入荷予定数量単位", "発注数 or 単位数",
         "バラ数必須", "ソースが単位数なら×発注単位。既にバラなら直接。サンプルで確認",
         "未完了", "西友2023-12"],
        ["G12-03", "★高", "コメント／発注区分", "発注データ区分",
         "文言コメント", "対照表確定（Phase1 Q&A-49／G-06）",
         "未完了", "G-06"],
        ["G12-04", "★高", "店舗・商品コード", "4桁／20桁等",
         "文字100", "マスタ統合追随。独自変換禁止", "方針確認済", "マスタ統合"],
        ["G12-05", "★中", "伝票No", "伝票NO",
         "予定↔実績連動キー（定義上）", "西友現行は紐づけ未使用記載。TRIALで要否確認",
         "未完了", "西友注記"],
        ["G12-06", "★中", "日付fallback", "納品日／予定日が空のケース",
         "必須日付", "空白時の代替項目を固定（西友編集仕様を候補）",
         "一部確定", "西友2023-10"],
        ["G12-07", "★中", "差分／全件", "取消区分",
         "差分or全件選択可", "稼働前に選択合意。取消は出さない or 数量0",
         "未完了", "sinops注意点"],
        ["G12-08", "★中", "未来日範囲", "納品予定日",
         "本日以降も含む全発注残", "抽出条件（何日先まで）を定義",
         "未完了", "sinops概要"],
    ]
    for g in gaps:
        ws.append(g)
    apply_grid(ws, [10, 8, 16, 20, 22, 40, 12, 14], verdict_col=None)
    wb.save(out)


def build_12_rule(out: Path):
    wb = Workbook()
    ws = wb.active
    ws.title = "変換ルール"
    cols = ["ルールID", "種別", "対象", "変換ロジック", "例", "依存GAP", "確定度"]
    intro(
        ws,
        "変換ルール定義書 — 入荷予定（nyuka_yotei.txt）／ Sinops-12",
        f"改訂 {TODAY.isoformat()}",
        "Layer①/③ 実装時の STEP 根拠",
        len(cols),
    )
    ws.append(cols)
    rules = [
        ["R12-01", "抽出", "対象", "本日入荷分を含む未来の未入荷発注残。勧告計算前日出力",
         "—", "G12-08", "方針確認済"],
        ["R12-02", "コード", "店舗・商品", "統合マスタ後コード。切捨て禁止",
         "—", "G12-04", "方針確認済"],
        ["R12-03", "数量", "入荷予定数量",
         "IF ソース.qty_unit == バラ: 直接設定。"
         "IF ソースが発注単位数: 発注数量×発注単位数。丸めない",
         "3×10→30", "G12-02", "未確定（源依存）"],
        ["R12-04", "日付", "入荷予定日", "納品日付。空白なら納品予定日付",
         "空→予定日", "G12-06", "一部確定"],
        ["R12-05", "日付", "発注日", "発注日付。空白なら発注予定日付",
         "空→予定日", "G12-06", "一部確定"],
        ["R12-06", "編集", "便区分", "trim。西友AS-IS候補：0→9。TRIAL適用はG-03確定後",
         "0→9", "G12-01", "未確定"],
        ["R12-07", "変換", "コメント",
         "発注データ区分→文言（1通常発注/4日配発注/10本部発注/11特売送込/12新店特売…）。"
         "未定義コードは原文または空",
         "11→特売送込", "G12-03", "未確定"],
        ["R12-08", "キー", "伝票No", "文字列化。空白時の扱い・実績紐付け要否は確認後固定",
         "—", "G12-05", "未確定"],
        ["R12-09", "形式", "ファイル", "TSV。差分or全件は合意値。キー重複は後優先",
         "nyuka_yotei.txt", "G12-07", "一部確定"],
    ]
    for r in rules:
        ws.append(r)
    apply_grid(ws, [10, 8, 14, 52, 14, 10, 14], verdict_col=None)
    wb.save(out)


# ---------------------------------------------------------------------------
# Sinops-21 受払明細
# ---------------------------------------------------------------------------

def build_21_mapping(out: Path):
    wb = Workbook()

    # 概要
    ws = wb.active
    ws.title = "対象IF概要"
    intro(
        ws,
        "受払明細IF 対象IF概要（2026-09-21 改訂）／ Sinops-21",
        "TRIAL 標準外部IF → sinops 【倉庫CD】uke.txt（TSV・10項目）",
        REF + " ／ Phase1訂正版（振替30／仕入10）を継承し、出庫発行日を検収相当へ更新",
        2,
    )
    ws.append(["項目", "内容"])
    overview = [
        ["出力先", "sinops 倉庫系 受払明細 【倉庫CD】uke.txt"],
        ["方向", "TRIAL → sinops"],
        ["出庫参照先", "振替伝票：bill_kindid=30 かつ btype=0 かつ f_del≠1"],
        ["入庫参照先", "仕入伝票：bill_kindid=10 かつ btype=0 かつ f_del≠1"],
        ["使わないソース", "ALLSIRE_MEISAI 単一ソースでの逆引きはしない"],
        ["対象外", "廃棄・伝票区分11/12/19/20/31/40/41/50・物流タイプ≠0（要再確認）"],
        ["未使用項目", "デポ／注文№／入荷数／伝票№／得意先コード → 固定値 0"],
        ["数量", "TRIAL qy は1000倍格納。sinops は実数。qy÷1000。符号は正値（×(-1)しない）"],
        ["出庫・発行日", "【升级反映】検収相当日付を優先（西友：納品日付→検収日付）。TRIAL候補：in_date 等"],
        ["入庫・発行日", "納品日（dlv_ymd）同値案。count_date 案は残確認"],
        ["西友AS-IS注意", "西友伝票種別(10振替/02/04/06)はMD基幹。TRIALには流用しない"],
        ["倉庫抽出", "西友AS-IS：新店判定=2(DC/TC) かつ センター種別=1(WMS在庫)。TRIAL側の倉庫判定はGAP"],
    ]
    for row in overview:
        ws.append(row)
    apply_grid(ws, [18, 70], header_row=4, data_start=5)

    # 抽出条件
    ws2 = wb.create_sheet("抽出条件")
    intro(ws2, "受払明細IF 抽出条件", f"改訂 {TODAY.isoformat()}", "フィルタはTRIAL伝票側", 8)
    ws2.append(["パターン", "sinops出力", "TRIAL伝票", "伝票区分", "物流タイプ", "削除", "倉庫CD候補", "備考"])
    ws2.append(["出庫", "売上数へ数量・入庫数=0", "振替伝票", "30", "0", "f_del≠1",
                "out_org_id（案）", "発行日＝検収相当（GAP-06）"])
    ws2.append(["入庫", "入庫数へ数量・売上数=0", "仕入伝票", "10", "0", "f_del≠1",
                "to_org_id / org_id（案）", "発行日＝納品日案"])
    ws2.append(["同一キー合算", "1行両方 or 2行", "両伝票", "30と10", "0", "—",
                "倉庫CD+品番+処理日", "GAP-08 未確定"])
    ws2.append(["対象外", "出力しない", "—", "11/12/19/20/31等", "≠0含む", "—", "—", "廃棄含む"])
    apply_grid(ws2, [12, 22, 12, 12, 10, 10, 18, 28], header_row=4)

    # 出庫
    ws3 = wb.create_sheet("マッピング_出庫")
    cols = ["No", "必須", "sinops項目", "sinops型/桁", "FMT", "TRIAL項目名", "TRIAL項目ID",
            "判定", "変換・設定案", "備考"]
    intro(
        ws3,
        "受払明細 出庫 — 振替伝票(30,btype=0) → uke.txt",
        "数量は売上数。入庫数=0。発行日は検収相当を優先",
        LEGEND,
        len(cols),
    )
    ws3.append(cols)
    out_rows = [
        ["1", "○", "処理日", "整数8", "在庫変動日 yyyymmdd", "振替日", "date", "△",
         "date を YYYYMMDD", "変動日でよいか要確認（GAP-06）"],
        ["2", "○", "デポ", "文字100", "未使用", "—", "—", "―", '固定 "0"', "定義どおり"],
        ["3", "○", "品番", "文字100", "商品コード", "商品コード", "item_id", "△",
         "文字列化。マスタ統合追随。int32オーバーフロー注意（GAP-07）", "—"],
        ["4", "", "注文№", "文字100", "未使用", "—", "—", "―", '固定 "0"', "—"],
        ["5", "○", "発行日", "整数8", "変動日/出力日", "入庫日／検収日", "in_date（優先）", "△",
         "【升级】検収相当を優先。in_date 不在時の代替は要合意。納品日へ戻さない",
         "西友2023-12"],
        ["6", "○", "入荷数", "整数8", "未使用", "—", "—", "―", '固定 "0"', "—"],
        ["7", "○", "入庫数", "整数8", "出庫行未使用", "—", "—", "―", "0 固定", "合算時のみ値"],
        ["8", "", "伝票№", "文字100", "未使用", "伝票番号", "code", "―", '固定 "0"', "ukeに出さない"],
        ["9", "○", "売上数", "整数8", "出庫数量", "数量(1000倍)", "qy", "△",
         "qy÷1000・正値。×(-1)しない", "GAP-05"],
        ["10", "", "得意先コード", "整数100", "未使用", "—", "—", "―", '固定 "0"', "—"],
        ["F1", "○", "（抽出）", "—", "伝票区分=30", "伝票区分", "bill_kindid", "○", "30のみ", "—"],
        ["F2", "○", "（抽出）", "—", "物流タイプ=0", "物流タイプ", "btype", "○", "0のみ", "GAP-12"],
        ["F3", "○", "（抽出）", "—", "削除以外", "削除フラグ", "f_del", "○", "1は出さない", "—"],
        ["F4", "△", "ファイル名", "—", "【倉庫CD】", "振替元店舗", "out_org_id", "△",
         "倉庫CDへ変換して分割（案）", "GAP-04"],
    ]
    for r in out_rows:
        ws3.append(r)
    apply_grid(ws3, [6, 6, 12, 10, 16, 14, 12, 6, 40, 16], verdict_col=8)

    # 入庫
    ws4 = wb.create_sheet("マッピング_入庫")
    intro(
        ws4,
        "受払明細 入庫 — 仕入伝票(10,btype=0) → uke.txt",
        "数量は入庫数。売上数=0",
        LEGEND,
        len(cols),
    )
    ws4.append(cols)
    in_rows = [
        ["1", "○", "処理日", "整数8", "在庫変動日 yyyymmdd", "納品日", "dlv_ymd", "△",
         "dlv_ymd を YYYYMMDD。count_dateとの使い分けはGAP-06", "—"],
        ["2", "○", "デポ", "文字100", "未使用", "—", "—", "―", '固定 "0"', "—"],
        ["3", "○", "品番", "文字100", "商品コード", "JAN", "item_id", "△",
         "トリムして品番へ。マスタ統合追随", "—"],
        ["4", "", "注文№", "文字100", "未使用", "—", "—", "―", '固定 "0"', "—"],
        ["5", "○", "発行日", "整数8", "変動日/出力日", "納品日", "dlv_ymd", "△",
         "案：処理日と同値（dlv_ymd）。西友入庫も納品日付", "GAP-06"],
        ["6", "○", "入荷数", "整数8", "未使用", "—", "—", "―", '固定 "0"', "入荷実績IFとは別"],
        ["7", "○", "入庫数", "整数8", "入庫数量", "数量(1000倍)", "qy", "△",
         "qy÷1000・正値", "GAP-05"],
        ["8", "", "伝票№", "文字100", "未使用", "—", "—", "―", '固定 "0"', "—"],
        ["9", "○", "売上数", "整数8", "入庫行未使用", "—", "—", "―", "0 固定", "—"],
        ["10", "", "得意先コード", "整数100", "未使用", "—", "—", "―", '固定 "0"', "—"],
        ["F1", "○", "（抽出）", "—", "伝票区分=10", "伝票区分", "bill_kindid", "○",
         "10のみ。11/12/31は入れない", "—"],
        ["F2", "○", "（抽出）", "—", "物流タイプ=0", "物流タイプ", "btype", "○", "0のみ", "—"],
        ["F3", "○", "（抽出）", "—", "削除以外", "削除フラグ", "f_del", "○", "1は出さない", "—"],
        ["F4", "△", "ファイル名", "—", "【倉庫CD】", "納品先/店舗", "to_org_id/org_id", "△",
         "倉庫キー未確定", "GAP-04"],
    ]
    for r in in_rows:
        ws4.append(r)
    apply_grid(ws4, [6, 6, 12, 10, 16, 14, 14, 6, 40, 16], verdict_col=8)

    # 差分
    ws5 = wb.create_sheet("旧整理との差分")
    intro(ws5, "旧整理・西友升级との差分", f"改訂 {TODAY.isoformat()}", "参照用", 5)
    ws5.append(["箇所", "旧", "今回", "判定", "残件"])
    diffs = [
        ["データソース", "ALLSIRE単一", "出庫=振替／入庫=仕入", "維持（訂正版）", "—"],
        ["出庫発行日", "納品日候補を含む", "検収相当（in_date）優先", "升级反映", "代替項目合意"],
        ["数量", "納品数そのまま", "qy÷1000", "維持", "端数"],
        ["符号", "×(-1)案", "正値", "維持", "—"],
        ["西友伝票種別", "10/02/04/06", "TRIALは30/10を使用", "流用しない", "—"],
        ["倉庫判定", "サンプルファイル名", "西友：DC/TC×WMS在庫提示", "要TRIAL定義", "GAP-04"],
    ]
    for d in diffs:
        ws5.append(d)
    apply_grid(ws5, [14, 22, 28, 14, 18], header_row=4)
    wb.save(out)


def build_21_gap(out: Path):
    wb = Workbook()
    ws = wb.active
    ws.title = "GAP分析書"
    cols = ["No", "優先度", "対象", "内容", "影響", "対応方針", "ステータス"]
    intro(
        ws,
        "GAP分析書 — 受払明細（【倉庫CD】uke.txt）／ Sinops-21",
        f"改訂 {TODAY.isoformat()} ／ {REF}",
        "Phase1訂正版GAPを継承。G06を検収日付升级で更新",
        len(cols),
    )
    ws.append(cols)
    gaps = [
        ["G21-01", "★最高", "参照先取り違え", "ALLSIRE単一は誤り。振替30/仕入10が正",
         "抽出全体", "訂正版を正とする。旧Excelは参照しない", "方針確認済"],
        ["G21-02", "★最高", "出庫抽出", "30+btype0。itemid_typeid・ギフト混入の切り分け未確定",
         "出庫件数", "追加絞り要否を業務確認", "未完了"],
        ["G21-03", "★最高", "入庫抽出", "10以外(11/12/31)除外でよいか未確認",
         "入庫件数", "本資料は10のみ。業務確認", "未完了"],
        ["G21-04", "★高", "倉庫CD", "out_org_id / to_org_id / org_id のどれが倉庫か未定",
         "ファイル分割", "倉庫判定ルール確定", "未完了"],
        ["G21-05", "★高", "数量スケール", "qy÷1000・端数・0件",
         "数量", "÷1000固定。端数処理合意", "一部確定"],
        ["G21-06", "★高", "処理日／発行日",
         "出庫発行日は検収相当へ升级。処理日=変動日、発行日=検収/出力日の割当未最終化。"
         "入庫は納品日同値案",
         "需要予測日付", "出庫：in_date優先。代替と入庫側を最終合意", "一部確定（升级反映）"],
        ["G21-07", "★高", "品番型", "振替item_idがint32で13桁JAN溢れる可能性",
         "品番", "型・コード体系確認", "未完了"],
        ["G21-08", "★中", "行粒度", "出庫行と入庫行を分けるか合算か",
         "レコード件数", "案A別行／案B合算", "未完了"],
        ["G21-09", "★中", "ギフト品振", "振替30に贈答が混ざるか",
         "出庫件数", "除外キー有無を確認", "未完了"],
        ["G21-10", "★中", "伝票区分対応", "出庫/入庫とTRIAL区分の最終表",
         "G-06", "Phase1 Q&A-151", "未完了"],
        ["G21-11", "★低", "qy=0", "出すか捨てるか",
         "件数", "方針合意", "未完了"],
        ["G21-12", "★低", "btype=0意味", "物流タイプ0の業務意味",
         "抽出", "定義確認", "未完了"],
    ]
    for g in gaps:
        ws.append(g)
    apply_grid(ws, [10, 8, 14, 48, 12, 36, 16], verdict_col=None)
    wb.save(out)


def build_21_rule(out: Path):
    wb = Workbook()
    ws = wb.active
    ws.title = "変更ルール定義書"
    cols = ["ルールID", "種別", "対象", "変換ロジック", "例", "依存GAP", "確定度"]
    intro(
        ws,
        "変更ルール定義書 — 受払明細（uke.txt）／ Sinops-21",
        f"改訂 {TODAY.isoformat()}",
        "Layer①/③ STEP 根拠。名称は旧『変更ルール』を継承",
        len(cols),
    )
    ws.append(cols)
    rules = [
        ["R21-01", "抽出", "出庫/入庫",
         "出庫：振替 WHERE bill_kindid=30 AND btype=0 AND f_del<>1 → 売上数行。"
         "入庫：仕入 WHERE bill_kindid=10 AND btype=0 AND f_del<>1 → 入庫数行",
         "—", "G21-01/02/03", "方針確認済"],
        ["R21-02", "ファイル", "【倉庫CD】",
         "案A出庫 out_org_id、案B入庫 to_org_id を倉庫CDへ。未変換は出さない/エラーは未定",
         "007452uke.txt", "G21-04", "未確定"],
        ["R21-03", "日付", "処理日・発行日",
         "出庫処理日：date（振替日）。"
         "出庫発行日：in_date（検収相当）優先。不在時の代替は合意後固定。"
         "入庫処理日・発行日：dlv_ymd 同値（案）",
         "発行日≠納品日（出庫）", "G21-06", "一部確定"],
        ["R21-04", "数量", "売上数・入庫数",
         "出庫：売上数=qy/1000・入庫数=0。入庫：入庫数=qy/1000・売上数=0。正値。×(-1)しない",
         "qy=3000→3", "G21-05", "方針確認済"],
        ["R21-05", "固定値", "未使用5項目",
         'デポ/注文№/入荷数/伝票№/得意先コード = "0"',
         "0", "—", "確定"],
        ["R21-06", "粒度", "行キー",
         "案A：出庫行と入庫行を別レコード。案B：倉庫CD+品番+処理日で合算",
         "—", "G21-08", "未確定"],
        ["R21-07", "品番", "商品コード",
         "文字列化＋trim。マスタ統合追随。型溢れはエラー",
         "—", "G21-07", "一部確定"],
        ["R21-08", "境界", "ギフト",
         "除外条件未定義の間は振替30+物流0を対象とする案（要確認）",
         "—", "G21-09", "未確定"],
        ["R21-09", "形式", "ファイル",
         "TSV・差分更新・勧告計算前日出力。障害時は未送信分を次回合算可",
         "uke.txt", "—", "確定"],
    ]
    for r in rules:
        ws.append(r)

    ws2 = wb.create_sheet("ukeレイアウト")
    intro(ws2, "sinops uke.txt 10項目", "西友レイアウト1.00＋升级注記", "", 5)
    ws2.append(["No", "項目", "必須", "出庫での値", "入庫での値"])
    layout = [
        ["1", "処理日", "○", "振替日", "納品日"],
        ["2", "デポ", "○", "0", "0"],
        ["3", "品番", "○", "商品コード", "JAN"],
        ["4", "注文№", "", "0", "0"],
        ["5", "発行日", "○", "検収相当（優先）", "納品日同値（案）"],
        ["6", "入荷数", "○", "0", "0"],
        ["7", "入庫数", "○", "0（合算時は入庫qy/1000）", "qy/1000"],
        ["8", "伝票№", "", "0", "0"],
        ["9", "売上数", "○", "qy/1000", "0"],
        ["10", "得意先コード", "", "0", "0"],
    ]
    for row in layout:
        ws2.append(row)
    apply_grid(ws2, [6, 14, 6, 28, 28], header_row=4)
    apply_grid(ws, [10, 8, 14, 56, 18, 12, 12], verdict_col=None)
    wb.save(out)


def write_readme(path: Path, if_id: str, name: str, files: list[str], notes: list[str]):
    lines = [
        f"# Sinops-{if_id} {name} — Phase1 ①②③ 改訂（{TODAY.isoformat()}）",
        "",
        "## ファイル",
        "",
    ]
    for f in files:
        lines.append(f"- `{f}`")
    lines += ["", "## 改訂要点", ""]
    for n in notes:
        lines.append(f"- {n}")
    lines += [
        "",
        "## 参照",
        "",
        f"- {REF}",
        "- `phase2-detail-design/sources/20260921-西友IFレイアウト版升级精查/精查結果.md`",
        "",
        "旧版は同フォルダ `BK/` へ退避。",
        "",
    ]
    path.write_text("\n".join(lines), encoding="utf-8")


def main():
    # --- 11 ---
    d11 = ROOT / "11．入荷実績（nyuka.txt）"
    d11.mkdir(parents=True, exist_ok=True)
    old_11_12 = ROOT / "11-12.発注仕入"
    if old_11_12.exists():
        bk = old_11_12 / "BK_20260921"
        archive_old(list(old_11_12.glob("*.xlsx")), bk)
    f11 = [
        "①項目マッピング表(nyuka.txt).xlsx",
        "②GAP分析書(nyuka.txt).xlsx",
        "③変換ルール定義書(nyuka.txt).xlsx",
    ]
    build_11_mapping(d11 / f11[0])
    build_11_gap(d11 / f11[1])
    build_11_rule(d11 / f11[2])
    write_readme(
        d11 / "README.md",
        "11",
        "入荷実績",
        f11,
        [
            "便区分：trim＋空白→0（西友2024-01升级）を変換ルールに明記",
            "店・商品はマスタ統合追随",
            "欠品は数量0",
        ],
    )

    # --- 12 ---
    d12 = ROOT / "12．入荷予定（nyuka_yotei.txt）"
    d12.mkdir(parents=True, exist_ok=True)
    f12 = [
        "①項目マッピング表(nyuka_yotei.txt).xlsx",
        "②GAP分析書(nyuka_yotei.txt).xlsx",
        "③変換ルール定義書(nyuka_yotei.txt).xlsx",
    ]
    build_12_mapping(d12 / f12[0])
    build_12_gap(d12 / f12[1])
    build_12_rule(d12 / f12[2])
    write_readme(
        d12 / "README.md",
        "12",
        "入荷予定",
        f12,
        [
            "数量はバラ必須。源が単位数なら×発注単位（西友基幹切替）",
            "日付空白時のfallback、便0→9候補、コメントコード変換を記載",
            "便区分本体はG-03/G-06待ち",
        ],
    )

    # pointer in old combined folder
    if old_11_12.exists():
        (old_11_12 / "README.md").write_text(
            "\n".join(
                [
                    "# 11-12.発注仕入（旧フォルダ）",
                    "",
                    f"{TODAY.isoformat()} 以降、資料は分割先を正とする：",
                    "",
                    "- `../11．入荷実績（nyuka.txt）/`",
                    "- `../12．入荷予定（nyuka_yotei.txt）/`",
                    "",
                    "旧xlsxは `BK_20260921/`。",
                    "",
                ]
            ),
            encoding="utf-8",
        )

    # --- 21 ---
    d21 = ROOT / "21．受払明細（【倉庫】uke.txt）"
    d21.mkdir(parents=True, exist_ok=True)
    bk21 = d21 / "BK_20260921"
    archive_old(
        [
            d21 / "受払明細_①_項目マッピング表.xlsx",
            d21 / "受払明細_②_GAP分析書.xlsx",
            d21 / "受払明細_③_変更ルール定義書.xlsx",
            # NFC/NFD variants
            *list(d21.glob("受払明細_①*.xlsx")),
            *list(d21.glob("受払明細_②*.xlsx")),
            *list(d21.glob("受払明細_③*.xlsx")),
        ],
        bk21,
    )
    # also archive decomposed name if still present
    for p in list(d21.glob("*.xlsx")):
        if p.parent == d21 and "BK" not in str(p):
            archive_old([p], bk21)

    f21 = [
        "①項目マッピング表(【倉庫CD】uke.txt).xlsx",
        "②GAP分析書(【倉庫CD】uke.txt).xlsx",
        "③変換ルール定義書(【倉庫CD】uke.txt).xlsx",
    ]
    build_21_mapping(d21 / f21[0])
    build_21_gap(d21 / f21[1])
    build_21_rule(d21 / f21[2])
    write_readme(
        d21 / "README.md",
        "21",
        "受払明細【倉庫】",
        f21,
        [
            "出庫発行日：検収相当（in_date）優先 — 西友2023-12升级反映",
            "抽出は振替30／仕入10＋btype0（訂正版維持）",
            "qy÷1000・正値・未使用5項目は0",
        ],
    )

    # sync 21 into phase1-pack-20260825
    pack825 = ROOT.parent / "phase1-pack-20260825" / "21．受払明細（【倉庫】uke.txt）"
    if pack825.exists():
        bk825 = pack825 / "BK_20260921"
        archive_old(list(pack825.glob("*.xlsx")), bk825)
        for name in f21:
            shutil.copy2(d21 / name, pack825 / name)
        shutil.copy2(d21 / "README.md", pack825 / "README.md")

    # also place 11/12 into 20260825 if parent exists
    pack825_root = ROOT.parent / "phase1-pack-20260825"
    if pack825_root.exists():
        for src, name in ((d11, "11．入荷実績（nyuka.txt）"), (d12, "12．入荷予定（nyuka_yotei.txt）")):
            dst = pack825_root / name
            dst.mkdir(parents=True, exist_ok=True)
            for f in src.glob("*"):
                if f.is_file():
                    shutil.copy2(f, dst / f.name)

    print("OK 11 →", d11)
    print("OK 12 →", d12)
    print("OK 21 →", d21)


if __name__ == "__main__":
    main()
