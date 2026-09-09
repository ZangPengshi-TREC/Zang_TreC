#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""新WMS 調整コード一覧（スクショ2枚を日文で全件整理）"""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

OUT = "/Users/treqd/Desktop/Cursor/西友PMI/在庫IF_訂正版/新WMS_調整コード一覧.xlsx"

NAVY = "0D2B4E"
WHITE = "FFFFFF"
BG_HEAD = "1A5C9E"
BG_ALT = "F1F5F9"
BG_SKIP = "E5E7EB"
BG_WARN = "FEF3C7"
BG_NOTE = "EEF2FF"
BG_BANNER = "FFF7ED"

thin = Border(
    left=Side(style="thin", color="CBD5E1"),
    right=Side(style="thin", color="CBD5E1"),
    top=Side(style="thin", color="CBD5E1"),
    bottom=Side(style="thin", color="CBD5E1"),
)
wrap = Alignment(wrap_text=True, vertical="center")
center = Alignment(wrap_text=True, vertical="center", horizontal="center")


def font(size=10, bold=False, color=NAVY):
    return Font(name="Meiryo UI", size=size, bold=bold, color=color)


def fill(h):
    return PatternFill("solid", fgColor=h)


# スクショ「新調整コード」シート。調整理由は一覧の短名、内容は詳細。
# RTV03「油澤」は読み取り誤りと判断し、ラベル貼り作業の記載に合わせた。
# DMG09「被損」は「破損」とする。
ROWS = [
    # code, 調整理由, 項目, 内容, WMS画面, MD_I/F, 連携コード, 増減, 備考, tag
    ["ADJ01", "サイクルカウント差異調整", "サイクル",
     "サイクルカウントの在庫差異調整",
     "DC在庫調整", "在庫調整実績", "ADJ01", "", "新コードをそのまま連携", ""],
    ["ADJ02", "問題商品解決の為", "サイクル",
     "原因不明の問題商品解決の際に発生した在庫差異調整",
     "DC在庫調整", "在庫調整実績", "ADJ02", "", "新コードをそのまま連携", ""],
    ["DMG01", "破損処理の為", "破損処理",
     "破損の問題商品解決の際に発生した在庫差異調整",
     "DC在庫調整", "在庫調整実績", "DMG01", "", "一覧上、理由が赤字", ""],
    ["OUT01", "在庫なしの為", "OUT",
     "在庫0の為、欠品処理",
     "DC在庫調整", "在庫調整実績", "OUT01", "", "", ""],
    ["RTV01", "商品返品", "RTV",
     "本部指示で取引先に商品を返品した場合",
     "DC在庫調整", "在庫調整実績", "RTV01", "", "", ""],
    ["RTV02", "コンテナ入荷時破損/不足", "RTV",
     "コンテナ入荷時発生した破損商品のデータ修正",
     "DC在庫調整", "在庫調整実績", "RTV02", "", "", ""],
    ["RTV03", "ラベル貼り/検品作業時発見破損", "RTV",
     "ラベル貼り作業中に発見された商品のデータ修正",
     "DC在庫調整", "在庫調整実績", "RTV03", "", "", ""],
    ["RTV04", "赤羽本部へ出荷", "RTV",
     "本部指示で赤羽本部に商品出荷した場合",
     "—", "—", "—", "", "9/20時点 MD基幹・会計 不要。行グレー／？", "SKIP"],
    ["DMG02", "小分け作業破損", "破損処理",
     "小分けで発生した作業破損の処理",
     "DC在庫調整", "在庫調整実績", "DMG02", "", "", ""],
    ["DMG03", "ケース作業破損", "破損処理",
     "ケースで発生した作業破損の処理",
     "DC在庫調整", "在庫調整実績", "DMG03", "−", "", ""],
    ["DMG04", "格納作業破損", "破損処理",
     "格納で発生した作業破損の処理",
     "DC在庫調整", "在庫調整実績", "DMG04", "−", "", ""],
    ["DMG05", "補充1（昼）作業破損", "破損処理",
     "補充1（昼）作業で発生した作業破損の処理",
     "DC在庫調整", "在庫調整実績", "DMG05", "−", "IFSYLMS320Rサンプルに出現", ""],
    ["DMG06", "冷凍作業破損", "破損処理",
     "冷凍で発生した作業破損の処理",
     "DC在庫調整", "在庫調整実績", "DMG06", "−", "", ""],
    ["DMG07", "DC入荷作業破損", "破損処理",
     "DC入荷で発生した作業破損の処理",
     "DC在庫調整", "在庫調整実績", "DMG07", "−", "", ""],
    ["DMG08", "QA作業破損", "破損処理",
     "QA作業で発生した作業破損の処理",
     "DC在庫調整", "在庫調整実績", "DMG08", "−", "", ""],
    ["DMG09", "輸送作業破損", "破損処理",
     "輸送中に発生した破損の処理",
     "DC在庫調整", "在庫調整実績", "DMG09", "−", "", ""],
    ["DMG10", "補充2（夜）作業破損", "破損処理",
     "補充2（夜）で発生した作業破損の処理",
     "DC在庫調整", "在庫調整実績", "DMG10", "−", "", ""],
    ["DMG11", "シュート下・乗越破損", "破損処理",
     "シュート下で発生した作業破損の処理",
     "DC在庫調整", "在庫調整実績", "DMG11", "−", "一覧上、行が黄", "WARN"],
    ["DMG12", "通過作業破損", "破損処理",
     "通過作業中に発生した作業破損の処理",
     "TC入出荷実績修正", "（TC入出荷実績）", "DMG12", "増減なし",
     "TC入出荷実績修正画面で入力", ""],
    ["TTD01", "店舗からの在庫移動", "移動",
     "店舗からの在庫返品時に使用",
     "在庫調整", "連携しない", "ZZZ05", "＋",
     "在庫調整画面で入力。MDに連携しない。連携コードはZZZ05", "SKIP"],
    ["TTD02", "他DCからの在庫移動", "移動",
     "DC間の在庫移動に使用",
     "在庫調整", "連携しない", "ZZZ06", "＋/−",
     "在庫調整画面で入力。MDに連携しない。連携コードはZZZ06", "SKIP"],
    ["SUP01", "用度使用", "用度",
     "在庫品を用度使用する際に使用",
     "在庫調整", "在庫調整実績", "SUP01", "−",
     "在庫調整画面で入力", ""],
    ["ZZZ01", "入荷実績数修正（格納後）", "訂正",
     "入庫後の入荷実績修正をした場合の在庫調整。入庫後に入荷数を修正し在庫数は変更されない場合に使用",
     "在庫調整", "連携しない", "ZZZ01", "＋/−",
     "在庫調整画面で入力。MDに連携しない", "SKIP"],
]


def main():
    wb = Workbook()
    ws = wb.active
    ws.title = "新調整コード"

    cols = 10
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=cols)
    ws.cell(1, 1, "新WMS 調整コード一覧（全件整理）").font = font(16, True, WHITE)
    ws.cell(1, 1).fill = fill(NAVY)
    ws.cell(1, 1).alignment = Alignment(vertical="center", horizontal="left", indent=1)
    ws.row_dimensions[1].height = 28

    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=cols)
    ws.cell(
        2, 1,
        "出典：Excel『新WMS 調整コード一覧』シート「新調整コード」のスクリーンショット2枚（2026-08-20受領）。"
        "2022/8/3打合せ「新在庫調整コードをそのまま連携すること」。"
        "対象外シート（調整理由コード登録／新調整コード_0525）は未転記。"
        "IFSYLMS320R 在庫調整コード（5桁）の定義に用いる。①TRIAL別テーブル要否は未決。",
    ).font = font(9, False, "334155")
    ws.cell(2, 1).fill = fill(BG_BANNER)
    ws.cell(2, 1).alignment = wrap
    ws.row_dimensions[2].height = 48

    headers = [
        "No", "新在庫調整コード", "調整理由", "項目", "内容",
        "次期WMS入力画面", "MD基幹I/F", "MD基幹へ連携するコード", "在庫増減", "備考",
    ]
    for c, h in enumerate(headers, 1):
        cell = ws.cell(4, c, h)
        cell.font = font(10, True, WHITE)
        cell.fill = fill(BG_HEAD)
        cell.alignment = center
        cell.border = thin
    ws.row_dimensions[4].height = 24

    for i, row in enumerate(ROWS):
        r = 5 + i
        tag = row[-1]
        vals = [i + 1] + row[:-1]
        bg = BG_SKIP if tag == "SKIP" else BG_WARN if tag == "WARN" else (BG_ALT if i % 2 else WHITE)
        for c, v in enumerate(vals, 1):
            cell = ws.cell(r, c, v)
            cell.font = font(9)
            cell.alignment = wrap if c in (3, 5, 10) else center
            cell.border = thin
            cell.fill = fill(bg)
        ws.row_dimensions[r].height = 36

    widths = [6, 18, 28, 12, 48, 20, 16, 22, 12, 42]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A5"
    ws.auto_filter.ref = f"A4:J{4 + len(ROWS)}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.print_title_rows = "1:4"

    # 凡例
    ws2 = wb.create_sheet("凡例")
    ws2.merge_cells("A1:B1")
    ws2.cell(1, 1, "凡例・読み取り注記").font = font(14, True, WHITE)
    ws2.cell(1, 1).fill = fill(NAVY)
    notes = [
        ["項目", "内容"],
        ["件数", f"新在庫調整コード {len(ROWS)} 件（スクショ2枚の範囲）"],
        ["系統", "ADJ＝サイクル差異／DMG＝作業破損／RTV＝返品／OUT＝欠品／TTD＝移動／SUP＝用度／ZZZ＝訂正"],
        ["連携方針", "2022/8/3：新在庫調整コードをそのままMD基幹へ連携（当時の西友MD基幹）。TRIAL格納先は未決"],
        ["グレー行", "MDに連携しない、または不要（RTV04／TTD01／TTD02／ZZZ01）"],
        ["黄行", "DMG11 は一覧上ハイライト"],
        ["符号01/02", "本一覧のコードとは別項目（数量の正負）"],
        ["未転記", "シート「調整理由コード登録」「新調整コード_0525」はスクショに無いため未収録"],
        ["読み取り", "RTV03の「油澤」は誤読と判断し未採用。DMG09の「被損」は破損として記載"],
    ]
    for i, row in enumerate(notes):
        for c, v in enumerate(row, 1):
            cell = ws2.cell(2 + i, c, v)
            cell.font = font(10, bold=(i == 0 or c == 1))
            cell.alignment = wrap
            cell.border = thin
            if i == 0:
                cell.fill = fill(BG_HEAD)
                cell.font = font(10, True, WHITE)
            elif c == 1:
                cell.fill = fill(BG_NOTE)
        ws2.row_dimensions[2 + i].height = 28
    ws2.column_dimensions["A"].width = 16
    ws2.column_dimensions["B"].width = 88

    wb.save(OUT)
    print("wrote", OUT)


if __name__ == "__main__":
    main()
